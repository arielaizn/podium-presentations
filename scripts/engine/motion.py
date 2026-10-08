"""Motion for Podium: slide transitions + entrance animations, written as native PowerPoint XML.

Content is grouped automatically from geometry (no per-archetype code):
  static  - backgrounds, colour fields, big full-height photos, hairlines, header strip, footer
  units   - connected clusters of the remaining shapes (a card and its text, a row, the title block, a chart)
Units are revealed in reading order (top-to-bottom, then left-to-right, or right-to-left in RTL decks).

modes: none | subtle (auto cascade after the slide appears) | presenter (title auto, then one unit per click)
"""
from lxml import etree
from pptx.oxml.ns import qn  # noqa: F401  (also used for XPath tags above)

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
EMU = 914400.0


def _box(shape):
    try:
        return (shape.left / EMU, shape.top / EMU, shape.width / EMU, shape.height / EMU)
    except Exception:
        return None


def _is_video(sh):
    return sh._element.find(".//" + qn("a:videoFile")) is not None


def _units(slide, W, H, rtl):
    """Cluster animatable shapes into units. Videos, and anything that frames a video (a device body, an offset
    block), stay still: the video plays from the first frame and its frame must already be there around it."""
    cand = []
    vids = [_box(sh) for sh in slide.shapes if _is_video(sh)]

    def frames_video(b):
        x, y, w, h = b
        return any(x - 0.05 <= vx and y - 0.05 <= vy and vx + vw <= x + w + 0.05 and vy + vh <= y + h + 0.05
                   for vx, vy, vw, vh in vids if vx is not None)
    for sh in slide.shapes:
        b = _box(sh)
        if b is None or _is_video(sh):
            continue
        x, y, w, h = b
        if vids and not (sh.has_text_frame and sh.text_frame.text.strip()) and frames_video(b):
            continue
        area = w * h / (W * H)
        thin = w < 0.04 or h < 0.04
        chrome = y > H - 0.75 or (y + h) < 0.78
        has_text = sh.has_text_frame and sh.text_frame.text.strip() != ""
        big_field = area > 0.22 and not has_text and not getattr(sh, "has_chart", False)
        outlined = has_text and sh._element.find(".//" + qn("a:rPr") + "/" + qn("a:noFill")) is not None
        if thin or chrome or big_field or outlined:  # ghost numerals / outlined words are decoration: static
            continue
        cand.append((sh, (x, y, w, h)))
    n = len(cand)
    parent = list(range(n))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def is_text(sh):
        return sh.has_text_frame and sh.text_frame.text.strip() != ""

    def inside(outer, inner):
        ox, oy, ow, oh = outer
        cx, cy = inner[0] + inner[2] / 2, inner[1] + inner[3] / 2
        return ox - 0.05 <= cx <= ox + ow + 0.05 and oy - 0.05 <= cy <= oy + oh + 0.05 and ow * oh > inner[2] * inner[3]

    def same_row(a, b):
        ov = min(a[1] + a[3], b[1] + b[3]) - max(a[1], b[1])
        hgap = max(a[0], b[0]) - min(a[0] + a[2], b[0] + b[2])
        return ov >= 0.4 * min(a[3], b[3]) and hgap < 0.4

    def label_above(a, b):  # small label directly above a title, overlapping in x
        ax, ay, aw, ah = a
        bx, by, bw, bh = b
        return ah < 0.42 and 0 <= by - (ay + ah) < 0.25 and min(ax + aw, bx + bw) - max(ax, bx) > 0.3 * min(aw, bw)

    def stacked(a, b, sa, sb):
        """Number over its label (or label over value): short text directly above other text in one column."""
        ax, ay, aw, ah = a
        bx, by, bw, bh = b
        gap = by - (ay + ah)
        xo = min(ax + aw, bx + bw) - max(ax, bx)
        short = len(sa.text_frame.text.strip()) <= 12
        return short and -0.1 <= gap < 0.45 and xo > 0.5 * min(aw, bw)

    for i in range(n):
        si, bi = cand[i]
        for j in range(n):
            if i == j:
                continue
            sj, bj = cand[j]
            join = inside(bi, bj) or (is_text(si) and is_text(sj) and (label_above(bi, bj) or stacked(bi, bj, si, sj)))
            tiny = lambda b: b[2] * b[3] < 0.04  # bullet squares, ticks, dots travel with their text
            if not join and j > i and same_row(bi, bj) and (
                    (is_text(si) and is_text(sj)) or (is_text(si) and tiny(bj)) or (is_text(sj) and tiny(bi))):
                join = True
            if join:
                parent[find(i)] = find(j)
    groups = {}
    for i, (sh, b) in enumerate(cand):
        groups.setdefault(find(i), []).append((sh, b))

    def span(g):
        return min(b[1] for _, b in g), max(b[1] + b[3] for _, b in g)

    def xkey(g):
        return min(b[0] for _, b in g) if not rtl else -max(b[0] + b[2] for _, b in g)

    # reading order: bands of units that share vertical space, top to bottom; inside a band, across
    units = sorted(groups.values(), key=lambda g: span(g)[0])
    bands = []
    for g in units:
        top, bot = span(g)
        if bands:
            btop, bbot = bands[-1]["span"]
            overlap = min(bot, bbot) - max(top, btop)
            if overlap > 0.5 * min(bot - top, bbot - btop):
                bands[-1]["units"].append(g)
                bands[-1]["span"] = (min(top, btop), max(bot, bbot))
                continue
        bands.append({"span": (top, bot), "units": [g]})
    ordered = [u for b in bands for u in sorted(b["units"], key=xkey)]
    return [[sh for sh, _ in u] for u in ordered]


class _Ids:
    def __init__(self):
        self.n = 2

    def __call__(self):
        self.n += 1
        return str(self.n)


def _effect(ids, spid, delay_ms, node_type, is_chart=False, dur=450, rise=False):
    """One entrance effect for one shape: fade; text units 'rise' (fade + a short upward drift that decelerates,
    PowerPoint's Float In); charts wipe up."""
    if is_chart:
        preset, filt, subtype = "22", "wipe(up)", "4"
    elif rise:
        preset, filt, subtype = "42", "fade", "0"
    else:
        preset, filt, subtype = "10", "fade", "0"
    par = etree.Element(qn("p:par"))
    ctn = etree.SubElement(par, qn("p:cTn"), id=ids(), presetID=preset, presetClass="entr", presetSubtype=subtype,
                           fill="hold", nodeType=node_type)
    st = etree.SubElement(ctn, qn("p:stCondLst"))
    etree.SubElement(st, qn("p:cond"), delay=str(int(delay_ms)))
    kids = etree.SubElement(ctn, qn("p:childTnLst"))
    s = etree.SubElement(kids, qn("p:set"))
    cb = etree.SubElement(s, qn("p:cBhvr"))
    c2 = etree.SubElement(cb, qn("p:cTn"), id=ids(), dur="1", fill="hold")
    etree.SubElement(etree.SubElement(c2, qn("p:stCondLst")), qn("p:cond"), delay="0")
    etree.SubElement(etree.SubElement(cb, qn("p:tgtEl")), qn("p:spTgt"), spid=str(spid))
    an = etree.SubElement(etree.SubElement(cb, qn("p:attrNameLst")), qn("p:attrName"))
    an.text = "style.visibility"
    etree.SubElement(etree.SubElement(s, qn("p:to")), qn("p:strVal"), val="visible")
    ae = etree.SubElement(kids, qn("p:animEffect"), transition="in", filter=filt)
    cb2 = etree.SubElement(ae, qn("p:cBhvr"))
    etree.SubElement(cb2, qn("p:cTn"), id=ids(), dur=str(dur))
    etree.SubElement(etree.SubElement(cb2, qn("p:tgtEl")), qn("p:spTgt"), spid=str(spid))
    if rise and not is_chart:
        for attr, frm in (("ppt_x", "#ppt_x"), ("ppt_y", "#ppt_y+0.035")):
            a = etree.SubElement(kids, qn("p:anim"), calcmode="lin", valueType="num")
            ab = etree.SubElement(a, qn("p:cBhvr"))
            etree.SubElement(ab, qn("p:cTn"), id=ids(), dur=str(dur), decel="100000", fill="hold")
            etree.SubElement(etree.SubElement(ab, qn("p:tgtEl")), qn("p:spTgt"), spid=str(spid))
            etree.SubElement(etree.SubElement(ab, qn("p:attrNameLst")), qn("p:attrName")).text = attr
            tl = etree.SubElement(a, qn("p:tavLst"))
            for tm, v in (("0", frm), ("100000", "#" + attr)):
                tav = etree.SubElement(tl, qn("p:tav"), tm=tm)
                etree.SubElement(etree.SubElement(tav, qn("p:val")), qn("p:strVal"), val=v)
    return par


def _click_par(ids, first_auto):
    par = etree.Element(qn("p:par"))
    ctn = etree.SubElement(par, qn("p:cTn"), id=ids(), fill="hold")
    st = etree.SubElement(ctn, qn("p:stCondLst"))
    etree.SubElement(st, qn("p:cond"), delay="indefinite")
    if first_auto:
        c = etree.SubElement(st, qn("p:cond"), evt="onBegin", delay="0")
        etree.SubElement(c, qn("p:tn"), val="2")
    kids = etree.SubElement(ctn, qn("p:childTnLst"))
    inner = etree.SubElement(kids, qn("p:par"))
    ictn = etree.SubElement(inner, qn("p:cTn"), id=ids(), fill="hold")
    etree.SubElement(etree.SubElement(ictn, qn("p:stCondLst")), qn("p:cond"), delay="0")
    return par, etree.SubElement(ictn, qn("p:childTnLst"))


def _media_call(ids, spid, dur_ms, node_type):
    """Start a video (PowerPoint's 'Play' media effect)."""
    par = etree.Element(qn("p:par"))
    ctn = etree.SubElement(par, qn("p:cTn"), id=ids(), presetID="1", presetClass="mediacall", presetSubtype="0",
                           fill="hold", nodeType=node_type)
    etree.SubElement(etree.SubElement(ctn, qn("p:stCondLst")), qn("p:cond"), delay="0")
    cmd = etree.SubElement(etree.SubElement(ctn, qn("p:childTnLst")), qn("p:cmd"), type="call", cmd="playFrom(0.0)")
    cb = etree.SubElement(cmd, qn("p:cBhvr"))
    etree.SubElement(cb, qn("p:cTn"), id=ids(), dur=str(int(dur_ms)), fill="hold")
    etree.SubElement(etree.SubElement(cb, qn("p:tgtEl")), qn("p:spTgt"), spid=str(spid))
    return par


def _video_node(ids, spid, loop, muted):
    """The media node of a video: volume, mute, loop until the slide ends."""
    v = etree.Element(qn("p:video"))
    mn = etree.SubElement(v, qn("p:cMediaNode"), vol="80000")
    if muted:
        mn.set("mute", "1")
    ctn = etree.SubElement(mn, qn("p:cTn"), id=ids(), fill="hold", display="0")
    if loop:
        ctn.set("repeatCount", "indefinite")
    etree.SubElement(etree.SubElement(ctn, qn("p:stCondLst")), qn("p:cond"), delay="indefinite")
    etree.SubElement(etree.SubElement(mn, qn("p:tgtEl")), qn("p:spTgt"), spid=str(spid))
    return v


def animate(slide, W, H, rtl, mode="subtle", stagger=160, media=None):
    """media: [(shape id, duration ms, loop, muted, autoplay)] for the slide's videos. Autoplay videos start with
    the slide (alongside the first entrance effect) and loop; the entrance cascade runs over them."""
    media = media or []
    play = [m for m in media if m[4]]
    units = [] if mode in (None, "none") else _units(slide, W, H, rtl)
    if not units and not media:
        return 0
    sld = slide._element
    old = sld.find(qn("p:timing"))
    if old is not None:
        sld.remove(old)
    ids = _Ids()
    timing = etree.Element(qn("p:timing"))
    tn = etree.SubElement(etree.SubElement(timing, qn("p:tnLst")), qn("p:par"))
    root = etree.SubElement(tn, qn("p:cTn"), id="1", dur="indefinite", restart="never", nodeType="tmRoot")
    seq = etree.SubElement(etree.SubElement(root, qn("p:childTnLst")), qn("p:seq"), concurrent="1", nextAc="seek")
    main = etree.SubElement(seq, qn("p:cTn"), id="2", dur="indefinite", nodeType="mainSeq")
    main_kids = etree.SubElement(main, qn("p:childTnLst"))
    pc = etree.SubElement(seq, qn("p:prevCondLst"))
    c = etree.SubElement(pc, qn("p:cond"), evt="onPrev", delay="0")
    etree.SubElement(etree.SubElement(c, qn("p:tgtEl")), qn("p:sldTgt"))
    nc = etree.SubElement(seq, qn("p:nextCondLst"))
    c = etree.SubElement(nc, qn("p:cond"), evt="onNext", delay="0")
    etree.SubElement(etree.SubElement(c, qn("p:tgtEl")), qn("p:sldTgt"))

    animated = []
    if mode == "presenter":
        batches = [units[:1]] + [[u] for u in units[1:]]
    else:
        batches = [units]
    if not units:
        batches = []
        if play:  # no entrance effects: the videos alone start with the slide
            par, kids = _click_par(ids, first_auto=True)
            for k, m in enumerate(play):
                kids.append(_media_call(ids, m[0], m[1], "afterEffect" if k == 0 else "withEffect"))
            main_kids.append(par)
    for bi, batch in enumerate(batches):
        par, kids = _click_par(ids, first_auto=(bi == 0))
        t = 0
        for ui, unit in enumerate(batch):
            for k, sh in enumerate(unit):
                first = ui == 0 and k == 0
                node = ("afterEffect" if bi == 0 else "clickEffect") if first else "withEffect"
                rise = any(x.has_text_frame and x.text_frame.text.strip() for x in unit)
                kids.append(_effect(ids, sh.shape_id, t, node, is_chart=getattr(sh, "has_chart", False),
                                    dur=600 if rise else 500, rise=rise))
                animated.append(sh)
            t += stagger
        if bi == 0:  # videos start together with the first entrance effect, not after the cascade
            for m in play:
                kids.append(_media_call(ids, m[0], m[1], "withEffect"))
        main_kids.append(par)
    for m in media:
        seq.getparent().append(_video_node(ids, m[0], m[2], m[3]))
    if not len(main_kids):  # nothing in the main sequence (a video without autoplay): no empty sequence
        seq.getparent().remove(seq)
    texts = [sh for sh in animated if sh.has_text_frame]
    if texts:
        bld = etree.SubElement(timing, qn("p:bldLst"))
        for sh in texts:
            etree.SubElement(bld, qn("p:bldP"), spid=str(sh.shape_id), grpId="0", animBg="1")
    # p:timing must come after clrMapOvr / transition
    anchor = sld.find(qn("p:transition"))
    if anchor is None:
        anchor = sld.find(qn("p:clrMapOvr"))
    anchor.addnext(timing) if anchor is not None else sld.append(timing)
    return len(animated)


def transition(slide, kind="fade"):
    """Slide transition: fade (default) | morph (PowerPoint 2019+, falls back to fade) | none."""
    if kind in (None, "none"):
        return
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    if kind == "morph":
        MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
        P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
        P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
        ac = etree.Element("{%s}AlternateContent" % MC, nsmap={"mc": MC})
        ch = etree.SubElement(ac, "{%s}Choice" % MC, nsmap={"p159": P159, "p14": P14}, Requires="p159")
        tr = etree.SubElement(ch, qn("p:transition"), spd="slow")
        tr.set("{%s}dur" % P14, "1200")
        etree.SubElement(tr, "{%s}morph" % P159, option="byObject")
        fb = etree.SubElement(ac, "{%s}Fallback" % MC)
        etree.SubElement(etree.SubElement(fb, qn("p:transition"), spd="slow"), qn("p:fade"))
        el = ac
    else:
        el = etree.Element(qn("p:transition"), spd="med")
        etree.SubElement(el, qn("p:fade"))
    anchor = sld.find(qn("p:clrMapOvr"))
    anchor.addnext(el) if anchor is not None else sld.append(el)

