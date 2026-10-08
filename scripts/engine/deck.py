"""Podium engine: brief-ready JSON -> on-brand, editable PowerPoint.

Every layout here is original code built on a 12-column grid and a fixed type scale.
No template files are used; the brand comes from a brand.json (colours, fonts, logo).

usage:
  python deck.py build <deck.json> <out.pptx>
  python deck.py schema                 # print the content schema of every archetype
"""
import copy, hashlib, json, math, os, re, shutil, subprocess, sys, tempfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).parent))
import motion  # noqa: E402

# ------------------------------------------------------------------ canvas & grid
W, H = 13.333, 7.5            # inches, 16:9
M = 0.66                      # outer margin
G = 0.24                      # gutter
COLS = 12
COLW = (W - 2 * M - (COLS - 1) * G) / COLS
TOP = 1.05                    # content top (below header strip)
BOTTOM = H - 0.72             # content bottom (above footer strip)
HEB = re.compile(r"[\u0590-\u05FF]")
# a left-to-right run inside a Hebrew line: Latin words, figures, currency amounts, joined by a middle dot, - / : or spaces
# (e.g. "A · 2026", "$1.2M", "Q3 2026"). Written as separate en-US runs so PowerPoint keeps them in order.
LTR_RUN = re.compile(r"[$\u20AC\u00A3]?[A-Za-z0-9][A-Za-z0-9.,%$\u20AC\u00A3+]*(?:(?:\s*[\u00B7/:\-\u2013]\s*|\s+)[$\u20AC\u00A3]?[A-Za-z0-9][A-Za-z0-9.,%$\u20AC\u00A3+]*)*")


def ltr_split(seg):
    """Split a Hebrew line into (text, is_ltr) pieces; only runs that hold Latin letters or currency count as LTR."""
    out, i = [], 0
    for m in LTR_RUN.finditer(seg):
        if not re.search(r"[A-Za-z$€£]", m.group(0)):
            continue
        tok, end = m.group(0), m.end()
        trail = len(tok) - len(tok.rstrip(".,"))      # sentence punctuation after a Latin word belongs to the Hebrew line
        if trail:
            tok, end = tok[:-trail], end - trail
        if m.start() > i:
            out.append((seg[i:m.start()], False))
        out.append((tok, True))
        i = end
    if i < len(seg):
        out.append((seg[i:], False))
    return out


SCALE = {  # pt: (max, min, line spacing, bold, font role)
    "display": (120, 54, 0.92, True, "heading"),
    "h1": (50, 30, 1.0, True, "heading"),
    "h2": (30, 20, 1.05, True, "heading"),
    "h3": (22, 16, 1.1, True, "heading"),
    "lead": (21, 15, 1.3, False, "body"),
    "body": (16, 12, 1.35, False, "body"),
    "small": (13, 10.5, 1.3, False, "body"),
    "label": (11, 9, 1.2, True, "body"),
    "number": (96, 40, 0.9, True, "heading"),
}


def col(i):            # x of column i (0-based)
    return M + i * (COLW + G)


def span(n):           # width of n columns
    return n * COLW + (n - 1) * G


def E(inches):
    return Emu(int(inches * 914400))


# ------------------------------------------------------------------ fonts & measuring
FONT_DIRS = [Path(r"C:\Windows\Fonts"), Path.home() / "AppData/Local/Microsoft/Windows/Fonts",
             Path.home() / "Library/Fonts", Path("/Library/Fonts"), Path("/usr/share/fonts"), Path("/usr/local/share/fonts"),
             Path.home() / ".local/share/fonts", Path.home() / ".fonts"]
_font_index = None
_font_cache = {}


def _index():
    global _font_index
    if _font_index is None:
        _font_index = []
        for d in FONT_DIRS:
            if d.exists():
                _font_index += [p for p in d.rglob("*") if p.suffix.lower() in (".ttf", ".otf")]
    return _font_index


def font_for(family, bold, size):
    key = (family, bold, round(size * 4))
    if key in _font_cache:
        return _font_cache[key]
    want = (family or "").lower().replace(" ", "")
    first = (family or "").lower().split(" ")[0]
    best, best_score = None, -1
    for p in _index():
        n = p.stem.lower().replace(" ", "").replace("-", "").replace("_", "")
        if not (n.startswith(want) or (len(first) > 3 and n.startswith(first) and "9pt" in want + n)):
            continue
        rest = n[len(want):] if n.startswith(want) else n[len(first):]
        if "italic" in n:
            continue  # never measure with an italic cut
        score = 1
        if bold and ("bold" in rest or "black" in rest or "wght" in rest or rest == ""):
            score += 2 if "bold" in rest else 1
        if not bold and (rest in ("", "regular") or "wght" in rest):
            score += 2
        if "italic" in rest:
            score -= 3
        if score > best_score:
            best, best_score = p, score
    f = None
    if best:
        try:
            f = ImageFont.truetype(str(best), max(4, round(size * 4)))
            try:
                f.set_variation_by_name("Bold" if bold else "Regular")
            except Exception:
                pass
            stem = best.stem.lower()
            # PowerPoint fakes bold when only Medium/SemiBold exists: measure a little wider to stay safe
            f.bold_exact = ("bold" in stem and "semibold" not in stem) or "black" in stem or "wght" in stem or "variable" in stem
            f.path_ = best
        except Exception:
            f = None
    _font_cache[key] = f
    return f


def text_width(text, family, bold, size):
    f = font_for(family, bold, size)
    if f is None:  # font not installed: average-glyph estimate
        return len(text) * size * (0.58 if bold else 0.52)
    w = f.getlength(text) / 4
    return w * 1.05 if bold and not getattr(f, "bold_exact", True) else w


def wrap(text, family, bold, size, width_pt, spacing=0):
    lines = []
    for para in text.split("\n"):
        words, cur = para.split(" "), ""
        for w in words:
            t = (cur + " " + w).strip()
            if text_width(t, family, bold, size) + spacing * len(t) <= width_pt or not cur:
                cur = t
            else:
                lines.append(cur)
                cur = w
        lines.append(cur)
    return lines


def _rel_lum(h):
    lin = [(c / 255) / 12.92 if c / 255 <= 0.03928 else ((c / 255 + 0.055) / 1.055) ** 2.4
           for c in (int(h[i:i + 2], 16) for i in (0, 2, 4))]
    return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]


def contrast_ratio(a, b):
    la, lb = sorted((_rel_lum(a), _rel_lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def contain_box(x, y, w, h, aspect, align="c"):
    """Largest box of the given aspect inside (x, y, w, h). align: c (centre) | l | r | t | b (LTR terms)."""
    if aspect > w / h:
        nw, nh = w, w / aspect
    else:
        nw, nh = h * aspect, h
    nx = x + {"l": 0, "r": w - nw}.get(align, (w - nw) / 2)
    ny = y + {"t": 0, "b": h - nh}.get(align, (h - nh) / 2)
    return nx, ny, nw, nh


# ------------------------------------------------------------------ deck context
class Deck:
    def __init__(self, spec, base_dir):
        self.spec = spec
        self.base = Path(base_dir)
        b = spec["brand"]
        if isinstance(b, str):
            bp = (self.base / b) if not Path(b).is_absolute() else Path(b)
            self.brand = json.loads(bp.read_text(encoding="utf8"))
            self.brand_dir = bp.parent
        else:
            self.brand, self.brand_dir = b, self.base
        c = self.brand["colors"]
        self.c = {k: v.lstrip("#").upper() for k, v in c.items()}
        self.c.setdefault("surface", self.mix(self.c["bg"], self.c["text"], 0.07))
        self.c.setdefault("line", self.mix(self.c["bg"], self.c["text"], 0.18))
        self.c.setdefault("muted", self.mix(self.c["bg"], self.c["text"], 0.55))
        self.c.setdefault("accent_text", "FFFFFF")
        self.lang = spec.get("lang", "en")
        self.rtl = self.lang in ("he", "ar")
        f = self.brand["fonts"]
        if self.rtl:
            self.fonts = {"heading": f.get("he_heading", f["heading"]), "body": f.get("he_body", f["body"])}
        else:
            self.fonts = {"heading": f["heading"], "body": f["body"]}
        self.caps_labels = not self.rtl and self.brand.get("caps_labels", True)
        self.radius = self.brand.get("radius", 0.14)
        # art direction: bold (poster type, colour fields) | editorial (meta strips, ghost type, rules) | soft (calm cards)
        self.art = spec.get("art", self.brand.get("art", "bold"))
        self.art_tighten = {"bold": 1.0, "editorial": 0.7, "soft": 0.3}.get(self.art, 0.7)
        self.title = spec.get("title", "")
        # brand shape language, so different brands don't share one signature: photo mask + offset block.
        # Derived from the brand's corner radius unless set: square brands get square photos, round brands round ones.
        r = self.radius
        sh = self.brand.get("shapes", {})
        self.photo_mask = sh.get("photo", "arch" if 0.1 <= r < 0.25 else ("round" if r >= 0.25 else None))
        self.photo_offset = sh.get("offset", r < 0.25)
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = E(W), E(H)
        self.tmp = Path(tempfile.mkdtemp(prefix="podium-"))
        self.warnings = []
        for role, fam in self.fonts.items():
            if font_for(fam, True, 20) is None:
                self.warnings.append(f"FONT '{fam}' ({role}) is not installed - PowerPoint will substitute it and text fitting "
                                     f"is estimated. Run: python engine/deck.py fonts <brand.json>")
        # small accent text must stay readable: derive a darker/lighter 'ink' version when the accent is too light
        self.accent_ink = self.c["accent"]
        for _ in range(12):
            if contrast_ratio(self.accent_ink, self.c["bg"]) >= 4.5:
                break
            self.accent_ink = self.mix(self.accent_ink, self.c["text"], 0.15)
        self.n = 0
        self.total = len(spec["slides"])

    @staticmethod
    def mix(a, b, t):
        ra, ga, ba = (int(a[i:i + 2], 16) for i in (0, 2, 4))
        rb, gb, bb = (int(b[i:i + 2], 16) for i in (0, 2, 4))
        return "%02X%02X%02X" % (round(ra + (rb - ra) * t), round(ga + (gb - ga) * t), round(ba + (bb - ba) * t))

    def video(self, sd):
        """Prepare a slide's video: re-encode to H.264 (<=1080p, ~8 Mbps, no audio unless `audio: true`), take a
        poster frame when none is given, read the duration. Cached across builds. Returns a dict or None."""
        v = sd["video"] if isinstance(sd["video"], dict) else {"path": sd["video"]}
        get = lambda k, dflt: v.get(k, sd.get(k, dflt))
        src = self.path(v["path"])
        if not src or not Path(src).exists():
            self.warnings.append(f"slide {self.n + 1}: missing video {v['path']}")
            return None
        audio = bool(get("audio", False))
        ff, fp = shutil.which("ffmpeg"), shutil.which("ffprobe")
        st = Path(src).stat()
        key = hashlib.md5(f"{Path(src).resolve()}|{st.st_mtime}|{st.st_size}|{audio}".encode()).hexdigest()[:12]
        cache = Path(tempfile.gettempdir()) / "podium-media"
        cache.mkdir(exist_ok=True)
        out, poster = cache / f"{Path(src).stem}-{key}.mp4", get("poster", None)
        poster = self.path(poster) if poster else None
        if ff:
            if not out.exists():
                cmd = [ff, "-v", "error", "-y", "-i", src, "-vf", "scale='min(1920,iw)':-2", "-c:v", "libx264",
                       "-preset", "medium", "-b:v", "8M", "-maxrate", "10M", "-bufsize", "16M", "-pix_fmt", "yuv420p",
                       "-profile:v", "high", "-movflags", "+faststart"]
                cmd += ["-c:a", "aac", "-b:a", "160k"] if audio else ["-an"]
                r = subprocess.run(cmd + [str(out)], capture_output=True)
                if r.returncode or not out.exists():
                    self.warnings.append(f"video {v['path']}: re-encode failed, embedding the original")
                    out = Path(src)
            if not poster:
                poster = cache / f"{Path(src).stem}-{key}.jpg"
                if not poster.exists():
                    subprocess.run([ff, "-v", "error", "-y", "-ss", str(get("poster_at", 0.5)), "-i", src,
                                    "-frames:v", "1", "-q:v", "2", str(poster)], capture_output=True)
                poster = str(poster)
        else:
            out = Path(src)
            self.warnings.append("ffmpeg not found: video embedded as-is (no re-encode, no automatic poster)")
        if not poster or not Path(poster).exists():
            self.warnings.append(f"video {v['path']}: no poster frame - add \"poster\": \"frame.jpg\"")
            return None
        dur = 0.0
        if fp:
            r = subprocess.run([fp, "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(out)],
                               capture_output=True, text=True)
            try:
                dur = float(r.stdout.strip())
            except ValueError:
                pass
        return {"src": str(out), "poster": str(poster), "dur_ms": int(dur * 1000) or 10000,
                "loop": bool(get("loop", True)), "muted": bool(get("muted", not audio)),
                "autoplay": bool(get("autoplay", True))}

    def path(self, p):
        if not p:
            return None
        q = Path(p)
        if not q.is_absolute():
            q = (self.base / p) if (self.base / p).exists() else (self.brand_dir / p)
        return str(q)


# ------------------------------------------------------------------ primitives
class S:
    """One slide being built. All coordinates are in inches in LTR space; RTL mirrors on placement."""

    def __init__(self, d: Deck, sd):
        self.d, self.sd = d, sd
        self.slide = d.prs.slides.add_slide(d.prs.slide_layouts[6])
        d.n += 1
        self.no = d.n
        # base | inverse | accent. On light brands structural slides flip to the dark text colour for rhythm;
        # dark brands stay dark (a cream section slide breaks a moody deck).
        light_brand = contrast_ratio(d.c["bg"], "000000") > contrast_ratio(d.c["bg"], "FFFFFF")
        flips = {"section": "inverse", "closing": "inverse"} if light_brand else {}
        tone = sd.get("tone", flips.get(sd["type"], "base"))
        if sd.get("image") and sd["type"] in ("closing", "statement", "cover"):
            tone = sd.get("tone", "base")  # photo slides keep the brand ground under the dimming
        c = d.c
        self.bg, self.fg, self.muted, self.line, self.surface = c["bg"], c["text"], c["muted"], c["line"], c["surface"]
        if tone == "inverse":
            self.bg, self.fg = c["text"], c["bg"]
            self.muted, self.line = d.mix(self.bg, self.fg, 0.55), d.mix(self.bg, self.fg, 0.18)
            self.surface = d.mix(self.bg, self.fg, 0.07)
        elif tone == "accent":
            self.bg, self.fg = c["accent"], c["accent_text"]
            self.muted, self.line = d.mix(self.bg, self.fg, 0.7), d.mix(self.bg, self.fg, 0.35)
            self.surface = d.mix(self.bg, "000000", 0.12)
        self.accent = c["accent"] if tone != "accent" else self.fg
        for _ in range(12):  # secondary (muted) text must still read: 4.5:1 against this slide's ground
            if contrast_ratio(self.muted, self.bg) >= 4.5:
                break
            self.muted = d.mix(self.muted, self.fg, 0.15)
        self.last_h = self.last_w = 0  # size of the last fitted text block (inches), for tight stacking
        self.ink = self.accent  # accent for small text, nudged until it reads on this slide's ground
        for _ in range(12):
            if contrast_ratio(self.ink, self.bg) >= 4.5:
                break
            self.ink = d.mix(self.ink, self.fg, 0.15)
        # reading ink: secondary text in long-text layouts. Brighter than `muted` (which is for labels and meta),
        # still a step below the main text so the hierarchy holds; never under 7:1 so paragraphs read on any ground
        self.body = self.readable(d.mix(self.bg, self.fg, 0.8), self.bg, 7.0, self.fg)
        self.photos = []        # big photo areas (LTR x, w): running labels get a pill there
        self.fields = []        # flat colour fields (LTR x, w, ink): running labels just take the field's ink
        self.full_bleed = False  # slide is a full-bleed photo or colour field: no meta strip
        self.media = []         # embedded videos on this slide: (shape id, duration ms, loop, muted, autoplay)
        self._vused = False     # the slide's video has been placed (it replaces the slide's main image once)
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(self.bg)

    # -- geometry
    def X(self, x, w):
        return W - x - w if self.d.rtl else x

    def _clean(self, shp):
        st = shp._element.find(qn("p:style"))
        if st is not None:
            shp._element.remove(st)
        return shp

    def rect(self, x, y, w, h, color, alpha=None, radius=0, shape=None):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        s = self._clean(self.slide.shapes.add_shape(kind, E(self.X(x, w)), E(y), E(w), E(h)))
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(color)
        if alpha is not None:
            clr = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
        s.line.fill.background()
        if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = min(0.5, radius / max(0.01, min(w, h)))
        return s

    def shadow(self, shp, blur=0.4, dist=0.1, alpha=0.10):
        """Soft drop shadow (cards on light grounds get depth; on dark grounds it is invisible, so skipped)."""
        if shp is None or not self.bg_is_light():
            return shp
        spPr = shp._element.find(qn("p:spPr"))
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(blur * 914400)), dist=str(int(dist * 914400)),
                              dir="5400000", algn="t", rotWithShape="0")
        clr = etree.SubElement(sh, qn("a:srgbClr"), val=self.d.c["text"])
        etree.SubElement(clr, qn("a:alpha"), val=str(int(alpha * 100000)))
        return shp

    def readable(self, color, ground, ratio=4.5, toward=None):
        """Nudge `color` toward `toward` (default: the ink that contrasts more with ground) until it reads."""
        if toward is None:
            toward = "FFFFFF" if contrast_ratio("FFFFFF", ground) > contrast_ratio("000000", ground) else "000000"
        for _ in range(14):
            if contrast_ratio(color, ground) >= ratio:
                break
            color = self.d.mix(color, toward, 0.15)
        return color

    def hline(self, x, y, w, color=None, weight=0.75):
        return self.rect(x, y, w, weight / 72, color or self.line)

    def vline(self, x, y, h, color=None, weight=0.75):
        return self.rect(x, y, weight / 72, h, color or self.line)

    def text(self, x, y, w, h, txt, role="body", color=None, align="l", anchor="t", bold=None,
             max_pt=None, min_pt=None, caps=False, tracking=0, font=None, ls=None, accent_color=None,
             outline=None, warn=True, para_space=0):
        """Text box that fits itself: shrinks from max to min size until the text fits w x h."""
        if txt is None or str(txt).strip() == "":
            return None
        if w <= 0.05 or h <= 0.05:  # a zero/negative box corrupts the file; report instead
            self.d.warnings.append(f"slide {self.no}: no room for text {str(txt)[:30]!r}")
            return None
        txt = str(txt)
        raw = txt                                 # keeps *accent* markers for the runs
        txt = txt.replace("*", "")                # measuring ignores them
        if role in ("display", "h1", "number") and tracking == 0:
            tracking = -0.018 * (max_pt or SCALE[role][0]) * self.d.art_tighten  # display type sits tight
        mx, mn, lsp, bld, frole = SCALE[role]
        mx, mn = max_pt or mx, min_pt or min(mn, max_pt or mn)
        bold = bld if bold is None else bold
        lsp = ls or lsp
        family = font or self.d.fonts[frole]
        if caps and self.d.caps_labels:
            txt = txt.upper()
        width_pt, height_pt = w * 72, h * 72
        size = mx
        n_paras = txt.count("\n") + 1
        while True:
            lines = wrap(txt, family, bold, size, width_pt, tracking)
            longest = max(text_width(l, family, bold, size) for l in lines)
            need_h = len(lines) * size * lsp * 1.18 + (n_paras - 1) * size * para_space
            if (need_h <= height_pt and longest <= width_pt * 1.01) or size <= mn:
                break
            size = max(mn, size * 0.94)
        # no widows: if the last line holds a single word, rebalance the breaks (same line count, narrower measure)
        if (len(lines) > 1 and len(lines[-1].split()) == 1 and "*" not in raw and "\n" not in raw
                and role not in ("label", "small", "number")):
            for k in range(1, 9):
                alt = wrap(txt, family, bold, size, width_pt * (1 - 0.04 * k), tracking)
                if len(alt) > len(lines):
                    break
                if len(alt[-1].split()) > 1:
                    raw = txt = "\n".join(alt)
                    lines = alt
                    break
        if warn and (need_h > height_pt * 1.02 or longest > width_pt * 1.03):
            self.d.warnings.append(f"slide {self.no}: text does not fit ({role}, {size:.0f}pt): {txt[:40]!r}")
        self.last_h, self.last_w, self.last_pt = need_h / 72, longest / 72, size
        tb = self.slide.shapes.add_textbox(E(self.X(x, w)), E(y), E(w), E(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        for a in list(bodyPr):
            bodyPr.remove(a)
        etree.SubElement(bodyPr, qn("a:noAutofit"))
        al = {"l": "r", "r": "l"}.get(align, align) if self.d.rtl else align
        if caps and self.d.caps_labels:
            raw = raw.upper()
        for i, para in enumerate(raw.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[al]
            p.line_spacing = lsp
            if para_space:
                p.space_after = Pt(size * para_space)
            self.runs(p, para, size, bold, family, color or self.fg, accent_color, tracking, outline, color)
        return tb

    def runs(self, p, para, size, bold, family, color, accent_color=None, tracking=0, outline=None, outline_color=None):
        """Write one paragraph as runs: Hebrew paragraphs RTL, Latin words / figures inside Hebrew as en-US runs,
        *word* in the accent colour. Shared by text boxes and table cells."""
        pPr = p._p.get_or_add_pPr()
        # Hebrew text, or currency/units in a Hebrew deck, reads right to left (₪ then sits left of the number,
        # the same way it does inside Hebrew sentences)
        if HEB.search(para) or (self.d.rtl and re.search(r"[₪€$£]", para)):
            pPr.set("rtl", "1")
        mixed = self.d.rtl and HEB.search(para)
        for k, seg0 in enumerate(re.split(r"\*", para)):
            if not seg0:
                continue
            # in a Hebrew line, Latin words / figures with symbols ("A · 2026", "$1.2M") become their own runs
            # tagged en-US, which is what PowerPoint uses to keep them in left-to-right order
            pieces = ltr_split(seg0) if mixed else [(seg0, False)]
            for seg, is_ltr in pieces:
                r = p.add_run()
                r.text = seg
                f = r.font
                f.size, f.bold, f.name = Pt(size), bold, family
                f.color.rgb = RGBColor.from_string((accent_color or self.accent) if k % 2 else color)
                rPr = r._r.get_or_add_rPr()
                if tracking:
                    rPr.set("spc", str(int(tracking * 100)))
                cs = etree.SubElement(rPr, qn("a:cs"))
                cs.set("typeface", family)
                if is_ltr:
                    rPr.set("lang", "en-US")
                elif HEB.search(seg) or mixed:
                    rPr.set("lang", "he-IL")
                if outline:
                    sf = rPr.find(qn("a:solidFill"))
                    if sf is not None:
                        rPr.remove(sf)
                    ln = etree.Element(qn("a:ln"), w=str(int(outline * 12700)))
                    lf = etree.SubElement(ln, qn("a:solidFill"))
                    etree.SubElement(lf, qn("a:srgbClr"), val=outline_color or self.line)
                    rPr.insert(0, ln)
                    rPr.insert(1, etree.Element(qn("a:noFill")))

    def fit_size(self, txt, w, h, role="body", max_pt=None, min_pt=None, bold=None, para_space=0):
        """Largest size (pt) at which txt fits w x h, without drawing anything."""
        if not txt:
            return max_pt or SCALE[role][0]
        mx, mn, lsp, bld, frole = SCALE[role]
        mx, mn = max_pt or mx, min_pt or min(mn, max_pt or mn)
        bold = bld if bold is None else bold
        fam = self.d.fonts[frole]
        txt = str(txt).replace("*", "")
        size, n_par = mx, txt.count("\n") + 1
        while True:
            lines = wrap(txt, fam, bold, size, w * 72)
            need = len(lines) * size * lsp * 1.18 + (n_par - 1) * size * para_space
            if (need <= h * 72 and max(text_width(l, fam, bold, size) for l in lines) <= w * 72 * 1.01) or size <= mn:
                return size
            size = max(mn, size * 0.94)

    def common(self, texts, w, h, role="body", max_pt=None, min_pt=None):
        """One size for a repeated set (rows, cards, steps): the size the longest item needs."""
        return min(self.fit_size(t, w, h, role, max_pt, min_pt) for t in texts if t) if any(texts) else max_pt

    def ghost(self, x, y, w, h, txt, size=260, align="l"):
        """Huge outlined word behind the content (may run off the slide)."""
        return self.text(x, y, w, h, txt, "display", color=self.d.mix(self.bg, self.fg, 0.22), max_pt=size,
                         min_pt=size, align=align, outline=1.25, anchor="m", warn=False)

    def _free(self, x0, x1):
        """True if the LTR span [x0, x1] is not covered by a big photo."""
        return all(x1 < px or x0 > px + pw for px, pw in self.photos)

    def tag(self, x, y, w, txt, color, align="l", bold=True):
        """A running label at a fixed spot. Over a photo it gets a small pill in the slide ground so it stays
        in place and readable (running elements never move or disappear)."""
        tr = 1 if self.d.caps_labels else 0
        tw = min(w, text_width(txt.upper() if self.d.caps_labels else txt, self.d.fonts["body"], bold, 9.5) / 72 + 0.12)
        tx = x if align == "l" else x + w - tw
        for f_ in self.fields:  # on a flat colour field: no pill, the field's own text colour
            fx, fw, ink = f_[:3]
            if tx < fx + fw and tx + tw > fx and y + 0.2 >= (f_[3] if len(f_) > 3 else 0):
                color = ink
                break
        else:
            if not self._free(tx, tx + tw):
                self.rect(tx - 0.12, y - 0.05, tw + 0.24, 0.34, self.bg, radius=0.17)
        self.text(x, y, w, 0.24, txt, "label", color=color, caps=True, tracking=tr, align=align, bold=bold,
                  max_pt=9.5, min_pt=9.5, warn=False)

    def strip(self, left=None, right=None):
        """Editorial meta strip: deck title on the left, section on the right, at fixed positions on every
        content slide; hairline across the free part."""
        if not (left or right):
            return
        half = (W - 2 * M) / 2
        if left:
            self.tag(M, 0.38, half, left, self.fg)
        if right:
            self.tag(W - M - half, 0.38, half, right, self.muted, align="r", bold=False)
        if not self.full_bleed:  # one full-width rule, the same on every content slide; it stops at photos
            x0 = M
            for px, pw in sorted(self.photos) + [(W - M, 0)]:
                if px - 0.2 > x0:
                    self.hline(x0, 0.7, min(px - 0.2, W - M) - x0, self.line, 0.75)
                x0 = max(x0, px + pw + 0.2)

    def label(self, x, y, w, txt, color=None, align="l"):
        return self.text(x, y, w, 0.3, txt, "label", color=color or self.ink, align=align,
                         caps=True, tracking=1.2 if self.d.caps_labels else 0)

    def framed(self, x, y, w, h, path, mask=None, offset=0.32, block=None):
        """Photo with a solid colour block offset behind it, optionally masked (arch | circle | round)."""
        fit = path.get("fit") if isinstance(path, dict) else None
        if (fit or self.sd.get("fit")) == "contain":  # screenshot: the frame takes the image's own shape
            p = self.d.path(path["path"] if isinstance(path, dict) else path)
            if p and Path(p).exists():
                im = Image.open(p)
                x, y, w, h = contain_box(x, y, w - offset, h - offset, im.width / im.height)
        if block is not False:
            self.rect(x + offset, y + offset, w, h, block or self.accent,
                      radius=self.d.radius if mask == "round" else 0,
                      shape={"arch": MSO_SHAPE.ROUND_2_SAME_RECTANGLE, "circle": MSO_SHAPE.OVAL}.get(mask))
            if mask == "arch":  # same arch geometry as the photo so the two curves match
                av = self.slide.shapes[-1]._element.find(qn("p:spPr")).find(qn("a:prstGeom")).find(qn("a:avLst"))
                for g in list(av): av.remove(g)
                etree.SubElement(av, qn("a:gd"), name="adj1", fmla="val 50000")
                etree.SubElement(av, qn("a:gd"), name="adj2", fmla="val 0")
        pic = self.image(x, y, w, h, path, radius=self.d.radius if mask == "round" else 0)
        if pic is not None and mask in ("arch", "circle"):
            geom = pic._element.find(qn("p:spPr")).find(qn("a:prstGeom"))
            geom.set("prst", "round2SameRect" if mask == "arch" else "ellipse")
            av = geom.find(qn("a:avLst"))
            for g in list(av): av.remove(g)
            if mask == "arch":  # round2SameRect needs both adjust values or PowerPoint rejects the file
                etree.SubElement(av, qn("a:gd"), name="adj1", fmla="val 50000")
                etree.SubElement(av, qn("a:gd"), name="adj2", fmla="val 0")
        return pic

    def image(self, x, y, w, h, path, flip=None, radius=0, focus=None, fit=None):
        """Picture in a box. fit=cover (default) crops around `focus` ([fx, fy], 0-1, default centre);
        fit=contain shows the whole image, uncropped, framed by a hairline (screenshots, UI, charts-as-images).
        In RTL decks photos flip so subjects face the text (screenshots never flip)."""
        if isinstance(path, dict):
            focus, fit, flip, path = path.get("focus", focus), path.get("fit", fit), path.get("flip", flip), path["path"]
        focus = focus or self.sd.get("focus") or (0.5, 0.5)
        fit = fit or self.sd.get("fit", "cover")
        v = self.sd.get("_video")
        if v and not self._vused and path == self.sd.get("image"):  # the slide's main image slot plays the video
            return self.movie(x, y, w, h, v, focus=focus, radius=radius, fit=fit)
        p = self.d.path(path)
        if not p or not Path(p).exists():
            self.rect(x, y, w, h, self.surface, radius=radius)
            self.d.warnings.append(f"slide {self.no}: missing image {path}")
            return None
        if fit == "contain":
            im = Image.open(p)
            ia, ba = im.width / im.height, w / h
            if abs(ia / ba - 1) > 0.04:
                # the whole image, as large as the box allows, sitting on the slide ground with a hairline edge
                # (no grey panel around it: a screenshot floating in a box reads as a placeholder)
                x, y, w, h = contain_box(x, y, w, h, ia, self.sd.get("align", "c"))
                pic = self.slide.shapes.add_picture(p, E(self.X(x, w)), E(y), E(w), E(h))
                self._edge(pic, radius, w, h)
                return pic
            flip = False   # nearly the box's own shape: fill it (a sliver is cropped at most), never mirrored
        else:
            flip = self.sd.get("flip", self.d.rtl) if flip is None else flip
        if flip:
            im = Image.open(p)
            out = self.d.tmp / (hashlib.sha256(str(Path(p).resolve()).encode()).hexdigest()[:20] + "_rtl" + Path(p).suffix)
            if not out.exists():
                im.transpose(Image.FLIP_LEFT_RIGHT).save(out)
            p = str(out)
        im = Image.open(p)
        if h >= H * 0.6 and w < W * 0.9:
            self.photos.append((x, w))
        pic = self.slide.shapes.add_picture(p, E(self.X(x, w)), E(y), E(w), E(h))
        ia, ba = im.width / im.height, w / h
        fx, fy = focus
        if flip:
            fx = 1 - fx
        if ia > ba:  # too wide: crop left/right around the focal point
            keep = ba / ia
            left = min(max(fx - keep / 2, 0), 1 - keep)
            pic.crop_left, pic.crop_right = left, 1 - keep - left
        else:        # too tall: crop top/bottom around the focal point
            keep = ia / ba
            top = min(max(fy - keep / 2, 0), 1 - keep)
            pic.crop_top, pic.crop_bottom = top, 1 - keep - top
        if fit == "contain":
            self._edge(pic, radius, w, h)
        elif radius:
            self._round(pic, radius, w, h)
        return pic

    def movie(self, x, y, w, h, v, focus=(0.5, 0.5), radius=0, fit="cover", edge=None):
        """Embedded video in a box, with its poster frame (what previews, PDFs and the editor show).
        cover: cropped to fill the box around `focus` (PowerPoint crops video like a picture); contain: the whole
        frame, as large as the box allows. Plays automatically and loops unless the slide says otherwise
        (timing is written by motion.animate)."""
        self._vused = True
        with Image.open(v["poster"]) as im:
            ia = im.width / im.height
        if fit == "contain":
            x, y, w, h = contain_box(x, y, w, h, ia, self.sd.get("align", "c"))
        mv = self.slide.shapes.add_movie(v["src"], E(self.X(x, w)), E(y), E(w), E(h), poster_frame_image=v["poster"],
                                         mime_type="video/mp4")
        ba = w / h
        fx, fy = focus
        if fit != "contain" and abs(ia / ba - 1) > 0.01:
            if ia > ba:
                keep = ba / ia
                left = min(max(fx - keep / 2, 0), 1 - keep)
                mv.crop_left, mv.crop_right = left, 1 - keep - left
            else:
                keep = ia / ba
                top = min(max(fy - keep / 2, 0), 1 - keep)
                mv.crop_top, mv.crop_bottom = top, 1 - keep - top
        if radius:
            self._round(mv, radius, w, h)
        if edge or (edge is None and fit == "contain"):
            self._edge(mv, radius, w, h)
        if h >= H * 0.6 and w < W * 0.9:
            self.photos.append((x, w))
        self.media.append((mv.shape_id, v["dur_ms"], v["loop"], v["muted"], v["autoplay"]))
        return mv

    @staticmethod
    def _round(pic, radius, w, h):
        geom = pic._element.find(qn("p:spPr")).find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        for g in list(av): av.remove(g)
        etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(min(50000, radius / min(w, h) * 100000))}")

    def _edge(self, pic, radius, w, h):
        """Screenshot treatment: small corner radius, a hairline in the slide's rule colour so a light UI does not
        bleed into a light ground (or a dark one into a dark ground), and a soft shadow on light grounds."""
        if radius:
            self._round(pic, min(radius, 0.12), w, h)
        pic.line.color.rgb = RGBColor.from_string(self.d.mix(self.bg, self.fg, 0.22))
        pic.line.width = Pt(0.75)
        self.shadow(pic, blur=0.35, dist=0.08, alpha=0.12)

    def cutout(self, x, y, w, path, flip=None):
        """Transparent PNG placed at its natural aspect (height follows width)."""
        p = self.d.path(path)
        im = Image.open(p)
        h = w * im.height / im.width
        flip = self.sd.get("flip", self.d.rtl) if flip is None else flip
        if flip:
            out = self.d.tmp / (hashlib.sha256(str(Path(p).resolve()).encode()).hexdigest()[:20] + "_rtl.png")
            if not out.exists():
                im.transpose(Image.FLIP_LEFT_RIGHT).save(out)
            p = str(out)
        return self.slide.shapes.add_picture(p, E(self.X(x, w)), E(y), E(w), E(h))

    def _side_luma(self, path, flip):
        """Brightness (0-1) of the side of a photo the text sits on, as displayed (after the RTL flip).
        Mean blended with the bright end, so a lit window behind the words counts."""
        p = self.d.path(path if not isinstance(path, dict) else path["path"])
        try:
            g = Image.open(p).convert("L")
        except Exception:
            return 0.5
        g.thumbnail((240, 240))
        if flip:
            g = g.transpose(Image.FLIP_LEFT_RIGHT)
        half = g.crop((g.width // 2, 0, g.width, g.height)) if self.d.rtl else g.crop((0, 0, g.width // 2, g.height))
        px = sorted(half.getdata())
        mean = sum(px) / len(px) / 255
        hi = px[int(len(px) * 0.85)] / 255
        return 0.65 * mean + 0.35 * hi

    def bg_image(self, path, dim=None, flip=None, scrim=True):
        """Full-bleed photo + overall dim + a soft gradient scrim behind the text side (left in LTR, right in RTL).
        The dim adapts to the photo: a dark photo is barely dimmed (it would turn to mud), a bright one is dimmed
        until the text side is dark (or light) enough for the type. An explicit `dim` is the most that is applied
        unless the type needs more to read."""
        self.full_bleed = True
        if isinstance(path, dict):
            flip = path.get("flip", flip)
        flip_ = self.sd.get("flip", self.d.rtl) if flip is None else flip
        self.image(0, 0, W, H, path, flip=flip)
        L = self._side_luma(path, flip_)
        if self.bg_is_light():   # dark type on a washed photo: the text side must end up light
            tgt, lb = 0.80, _rel_lum(self.bg) ** (1 / 2.2)
            need = (tgt - L) / max(0.05, lb - L) if L < tgt else 0.0
        else:                    # light type: the text side must end up dark
            tgt, lb = 0.26, _rel_lum(self.bg) ** (1 / 2.2)
            need = (L - tgt) / max(0.05, L - lb) if L > tgt else 0.0
        need = min(0.82, max(0.0, need))
        eff = max(0.12, need) if dim is None else max(need, min(dim, need + 0.15))
        eff = min(0.85, max(0.08, eff))
        self.rect(0, 0, W, H, self.bg, alpha=eff)
        if scrim:
            r = self._clean(self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(0), E(0), E(W), E(H)))
            r.line.fill.background()
            spPr = r._element.find(qn("p:spPr"))
            for old in spPr.findall(qn("a:solidFill")): spPr.remove(old)
            g = etree.Element(qn("a:gradFill"), rotWithShape="1")
            gl = etree.SubElement(g, qn("a:gsLst"))
            # opaque on the text side, clear on the far side; lighter when the photo is already dark there
            k = min(1.0, max(0.45, 0.35 + need * 1.3))
            stops = [(0, 0.55 * k), (55000, 0.25 * k), (100000, 0.0)]
            if self.d.rtl:
                stops = [(0, 0.0), (45000, 0.25 * k), (100000, 0.55 * k)]
            for pos, a in stops:
                gs = etree.SubElement(gl, qn("a:gs"), pos=str(pos))
                c = etree.SubElement(gs, qn("a:srgbClr"), val=self.bg)
                etree.SubElement(c, qn("a:alpha"), val=str(int(a * 100000)))
            etree.SubElement(g, qn("a:lin"), ang="0", scaled="0")
            spPr.find(qn("a:prstGeom")).addnext(g)

    # -- chrome
    def chrome(self, section=None):
        """Footer (brand mark left, page number right) and top strip (section left, deck title right),
        always at the same fixed spots. Anything whose spot is covered by a photo is left out, never moved."""
        if self.sd.get("chrome") is False or self.sd["type"] == "cover":
            return
        b = self.d.brand
        y = H - 0.52
        label = b.get("label", b.get("name", ""))
        logo = b.get("logo_on_light" if self.bg_is_light() else "logo") or b.get("logo")
        lp = self.d.path(logo) if logo else None
        # footer: brand mark bottom-left, page number bottom-right, same spot on every content slide
        if lp and Path(lp).exists():  # the logo replaces the brand label: one asset, deck-wide
            im = Image.open(lp)
            ratio = im.width / im.height
            lh = 0.26
            lw = lh * ratio
            if lw > 2.4:  # very wide wordmarks: cap the width, keep the aspect
                lw, lh = 2.4, 2.4 / ratio
            if not self._free(M, M + lw):
                self.rect(M - 0.12, y - 0.07, lw + 0.24, lh + 0.1, self.bg, radius=0.17)
            self.slide.shapes.add_picture(lp, E(self.X(M, lw)), E(y - 0.02 + (0.26 - lh) / 2), E(lw), E(lh))
        else:
            self.tag(M, y, 3.0, label, self.muted)
        if self.sd.get("page_number", True):
            self.tag(W - M - 2.0, y, 2.0, f"{self.no:02d} / {self.d.total:02d}", self.muted, align="r", bold=False)
        if section and section.strip().lower() == label.strip().lower():
            section = None  # never print the brand twice
        if self.d.art in ("bold", "editorial") and self.sd.get("strip", True):
            left = self.d.title or section          # the left label never changes sides
            self.strip(left, section if self.d.title else None)
        elif section:
            self.tag(M, 0.42, 6.0, section, self.muted)

    def bg_is_light(self):
        r, g, b = (int(self.bg[i:i + 2], 16) for i in (0, 2, 4))
        return (0.299 * r + 0.587 * g + 0.114 * b) > 140


# ------------------------------------------------------------------ archetypes
def a_cover(s: S, c):
    v = c.get("variant", "image" if c.get("image") else "type")
    # titles are bottom-anchored; the kicker sits just above the real text height, not above the box
    if v == "image":
        s.bg_image(c["image"], c.get("dim"))
        s.text(col(0), 2.2, span(c.get("width", 10)), 3.1, c["title"], "display", anchor="b", max_pt=c.get("max_pt", 120))
        s.label(col(0), 5.3 - s.last_h - 0.45, span(8), c.get("kicker"))
        s.text(col(0), 5.45, span(7), 0.9, c.get("subtitle"), "lead")
    elif v == "split":
        s.image(col(7), 0, W - col(7), H, c["image"])
        s.text(col(0), 1.5, span(6), 3.4, c["title"], "display", anchor="b", max_pt=c.get("max_pt", 96))
        s.label(col(0), 4.9 - s.last_h - 0.45, span(6), c.get("kicker"))
        s.hline(col(0), 5.15, span(2), s.accent, 3)
        s.text(col(0), 5.4, span(6), 1.0, c.get("subtitle"), "lead")
    else:  # type-only, oversized
        s.text(col(0), 1.2, span(12), 3.9, c["title"], "display", anchor="b", max_pt=c.get("max_pt", 150))
        s.label(col(0), 5.1 - s.last_h - 0.45, span(8), c.get("kicker"))
        s.hline(col(0), 5.35, span(12))
        s.text(col(0), 5.6, span(7), 0.95, c.get("subtitle"), "lead")
    if v == "type":  # meta lines (date, presenter, url) sit opposite the subtitle
        for i, m in enumerate((c.get("meta") or [])[:3]):
            s.text(col(8), 5.6 + i * 0.3, span(4), 0.28, m, "small", color=s.muted, align="r")
    s.chrome()


def a_section(s: S, c):
    s.text(col(0), 1.0, span(6), 4.4, c.get("number", f"{s.no:02d}"), "number", color=s.accent, max_pt=260, min_pt=120,
           anchor="m")
    s.label(col(7), 2.0, span(5), c.get("kicker"))
    s.text(col(7), 2.4, span(5), 1.9, c["title"], "h1", anchor="t", max_pt=60)
    s.text(col(7), 4.35, span(5), 1.6, c.get("text"), "lead", color=s.muted)
    if c.get("cutout"):  # transparent figure standing in front of the numeral
        s.cutout(c.get("cutout_x", col(3)), c.get("cutout_y", 0.9), c.get("cutout_w", 3.2), c["cutout"])
    s.chrome(c.get("section"))


def a_statement(s: S, c):
    if c.get("image"):
        s.bg_image(c["image"], c.get("dim"))
    s.label(col(0), 1.3, span(8), c.get("kicker"))
    s.text(col(0), 1.7, span(c.get("width", 10)), 3.7, c["title"], "display", max_pt=c.get("max_pt", 104))
    if c.get("text"):
        s.hline(col(0), 5.55, span(1), s.accent, 3)
        s.text(col(0), 5.75, span(6), 0.95, c["text"], "lead", color=s.fg if c.get("image") else s.muted)
    s.chrome(c.get("section"))


def a_agenda(s: S, c):
    items = [it if isinstance(it, dict) else {"title": it} for it in c["items"][:8]]
    s.label(col(0), TOP, span(4), c.get("kicker"))
    s.text(col(0), TOP + 0.35, span(4), 2.2, c["title"], "h1")
    s.text(col(0), TOP + 2.7, span(4), 2.0, c.get("text"), "body", color=s.muted)
    top, bot = TOP + 0.1, BOTTOM - 0.1
    rh = (bot - top) / len(items)
    cx, cw = col(5), span(7)
    has_text = any(it.get("text") for it in items)
    for i, it in enumerate(items):
        y = top + i * rh
        s.hline(cx, y, cw)
        th = min(0.48, rh * (0.5 if has_text else 0.8))
        ty = y + (rh - th) / 2 if not it.get("text") else y + 0.1
        s.text(cx, ty, 0.75, th, f"{i + 1:02d}", "h3", color=s.accent, anchor="m")
        s.text(cx + 0.85, ty, cw - 0.85 - (1.3 if it.get("meta") else 0), th, it["title"], "h3", anchor="m", max_pt=22)
        if it.get("text"):
            s.text(cx + 0.85, ty + th, cw - 0.85, rh - th - 0.12, it["text"], "small", color=s.muted)
        if it.get("meta"):
            s.text(cx + cw - 1.25, ty, 1.25, th, it["meta"], "small", color=s.muted, align="r", anchor="m")
    s.hline(cx, bot, cw)
    s.chrome(c.get("section"))


def _head(s: S, c, width=8, tight=False):
    s.label(col(0), TOP, span(width), c.get("kicker"))
    s.text(col(0), TOP + 0.32, span(width), 1.15, c["title"], "h1", max_pt=c.get("title_pt", 46))
    if tight:  # dense layouts: the content starts under the real title, not under a fixed two-line box
        y = TOP + 0.32 + s.last_h + 0.2
        if c.get("text"):
            s.text(col(0), y, span(width), 0.75, c["text"], "lead", color=s.body, max_pt=19)
            y += s.last_h + 0.15
        return y + 0.3
    if c.get("text"):
        s.text(col(0), TOP + 1.5, span(width), 0.75, c["text"], "lead", color=s.muted)
        return TOP + 2.45
    return TOP + 1.65


def metric(s: S, x, y, w, h, value, unit=None, color=None, max_pt=96, min_pt=40):
    """Big number with an optional smaller unit beside it (₪, €, M, h). Returns the y just below the number.
    In RTL decks the unit lands on the left of the number, as it does in Hebrew text."""
    m_ = re.match(r"^\s*([₪€$£])\s*([\d.,]+\s*[KMB]?)\s*$|^\s*([\d.,]+\s*[KMB]?)\s*([₪€$£])\s*$", str(value))
    if m_ and not unit:  # "380 ₪" / "₪380" -> number + unit, so every figure is formatted the same way
        value, unit = (m_.group(2), m_.group(1)) if m_.group(1) else (m_.group(3), m_.group(4))
    uw = 0
    if unit:  # measured, so a long unit ("מיליון ₪", "M ILS") never wraps under the number
        uw = min(w * 0.42, 0.18 + text_width(unit, s.d.fonts["heading"], True, max_pt * 0.38) / 72)
    s.text(x, y, w - uw, h, value, "number", color=color, max_pt=max_pt, min_pt=min_pt)
    nh, nw, npt = s.last_h, s.last_w, s.last_pt
    if unit:
        s.text(x + nw + 0.08, y + nh * 0.32, max(uw, w - nw - 0.08), nh * 0.6, unit, "h2", color=color, max_pt=npt * 0.38,
               min_pt=npt * 0.25, anchor="t", bold=True, warn=False)
    return y + nh


def a_kpi(s: S, c):
    y0 = _head(s, c)
    m = c["metrics"][:4]
    n = len(m)
    rows = 1                                   # one row reads as a scoreboard; 4 figures still get ~80pt
    per = n
    cw = (W - 2 * M) / per
    rh = min((BOTTOM - y0) / rows, 3.0)
    y0 = y0 + ((BOTTOM - y0) - rh * rows) * 0.4   # few metrics: centre them instead of leaving the bottom empty
    hl = c.get("highlight", [0])
    for i, it in enumerate(m):
        r, k = divmod(i, per)
        x, y = M + k * cw, y0 + r * rh
        s.hline(x, y, cw - (G if k < per - 1 else 0))
        if k:
            s.vline(x - G / 2, y + 0.15, rh - 0.3)
        pad = 0.25 if k else 0
        yb = metric(s, x + pad, y + 0.2, cw - pad - G, rh * 0.5, it["value"], it.get("unit"),
                    color=s.accent if i in hl else s.fg, max_pt=120 if per <= 2 else (104 if per == 3 else 84))
        s.hline(x + pad, yb + 0.12, 0.45, s.accent if i in hl else s.fg, 2.25)   # dash + label tight under the number
        s.text(x + pad, yb + 0.28, cw - pad - G, y + rh - yb - 0.35, it["label"], "body",
               caps=not it.get("note"), tracking=0.6 if not it.get("note") else 0)
    s.chrome(c.get("section"))


def a_process(s: S, c):
    y0 = _head(s, c)
    steps = c["steps"][:5]
    n = len(steps)
    cw = (W - 2 * M - (n - 1) * G) / n
    # big step numerals on a ruled line: scale contrast instead of small dots; the highlighted step carries the accent
    note_h = 0.7 if c.get("note") else 0
    numh = min(1.25, max(0.8, (BOTTOM - y0 - note_h - 2.6) * 0.6))
    yb = y0 + 0.15
    ry = yb + numh + 0.12                      # the rule under the numerals
    bh = BOTTOM - ry - 0.95 - note_h
    sz = s.common([st.get("text") for st in steps], cw, bh, "body", max_pt=17)  # one size for all steps
    tz = s.common([st["title"] for st in steps], cw, 0.6, "h3", max_pt=24)
    npt = min(s.fit_size(st.get("n", f"{i + 1:02d}"), cw, numh, "number", max_pt=96, min_pt=40) for i, st in enumerate(steps))
    hl = c.get("highlight", [0])
    quiet = s.readable(s.d.mix(s.bg, s.fg, 0.32), s.bg, 3.0, s.fg)   # large type: 3:1 is enough
    s.hline(M, ry, W - 2 * M, s.line, 1)
    for i, st in enumerate(steps):
        x = M + i * (cw + G)
        on = i in hl
        s.text(x, yb, cw, numh, st.get("n", f"{i + 1:02d}"), "number", color=s.accent if on else quiet,
               max_pt=npt, min_pt=npt, anchor="b")
        s.hline(x, ry - 0.02, cw if i < n - 1 else cw, s.accent if on else s.fg, 3 if on else 1.25)
        s.text(x, ry + 0.28, cw - 0.4, 0.6, st["title"], "h3", max_pt=tz, min_pt=tz)
        s.text(x, ry + 0.9, cw, bh, st.get("text"), "body", color=s.muted, max_pt=sz, min_pt=sz)
    if c.get("note"):
        s.rect(M, BOTTOM - 0.62, W - 2 * M, 0.5, s.surface, radius=s.d.radius / 2)
        s.text(M + 0.25, BOTTOM - 0.53, W - 2 * M - 0.5, 0.32, c["note"], "body", anchor="m")
    s.chrome(c.get("section"))


def a_timeline(s: S, c):
    y0 = _head(s, c)
    ev = c["events"][:8]
    n = len(ev)
    zig = n > 4                       # many events alternate above/below the line
    step = (W - 2 * M) / n
    if not zig:  # few events: the dates become display numerals above the line, the story sits below
        hl_ = c.get("highlight", [n - 1])
        wh = min(1.3, (BOTTOM - y0) * 0.36)
        ymid = y0 + 0.1 + wh + 0.22
        # one size for every date, small enough that each stays on one line
        hf = s.d.fonts["heading"]
        wpt = max(24, min([72, wh * 72 / 1.1] + [(step - G - 0.15) * 72 * 100 / max(1, text_width(e["when"], hf, True, 100))
                                                 for e in ev if e.get("when")]))
        s.hline(M, ymid, W - 2 * M, s.line, 1.25)
        quiet = s.readable(s.d.mix(s.bg, s.fg, 0.4), s.bg, 3.0, s.fg)
        tz = s.common([e["title"] for e in ev], step - G, 0.95, "h2", max_pt=30)
        for i, e in enumerate(ev):
            x = M + i * step
            hl = i in hl_
            s.text(x, y0 + 0.1, step - G, wh, e.get("when"), "number", color=s.accent if hl else quiet,
                   max_pt=wpt, min_pt=wpt, anchor="b")
            if hl:  # the milestone that matters: an accent segment on the line + a filled node
                s.hline(x, ymid - 0.02, step - G, s.accent, 3.5)
            s.rect(x, ymid - 0.11, 0.22, 0.22, s.accent if hl else s.fg, shape=MSO_SHAPE.OVAL)
            s.text(x, ymid + 0.4, step - G, 0.95, e["title"], "h2", max_pt=tz, min_pt=tz)
            s.text(x, ymid + 0.4 + s.last_h + 0.2, step - G, BOTTOM - ymid - s.last_h - 0.65, e.get("text"), "body",
                   color=s.muted, max_pt=17)
        s.chrome(c.get("section"))
        return
    ymid = y0 + (BOTTOM - y0) * 0.5
    s.hline(M, ymid, W - 2 * M, s.fg, 1.5)
    dot = 0.26
    for i, e in enumerate(ev):
        x = M + i * step
        hl = i in c.get("highlight", [n - 1])
        s.rect(x, ymid - dot / 2, dot, dot, s.accent if hl else s.bg, shape=MSO_SHAPE.OVAL)
        o = s.slide.shapes[-1]; o.line.fill.solid(); o.line.fill.fore_color.rgb = RGBColor.from_string(s.accent if hl else s.fg); o.line.width = Pt(1.5)
        tw = (step * 2 if zig else step) - G
        up = zig and i % 2 == 0
        y = ymid - 1.75 if up else ymid + 0.4
        s.label(x, y, tw, e.get("when"), color=s.accent if hl else s.muted)
        if zig:
            s.text(x, y + 0.32, tw, 0.65, e["title"], "h3", max_pt=24)
            s.text(x, y + 0.95, tw, 0.75, e.get("text"), "body", color=s.muted)
        else:  # few events: give each one real presence
            s.text(x, y + 0.35, tw, 1.0, e["title"], "h2", max_pt=30)
            s.text(x, y + 1.4, tw, BOTTOM - y - 1.5, e.get("text"), "lead", color=s.muted)
    s.chrome(c.get("section"))


def a_comparison(s: S, c):
    y0 = _head(s, c)
    L, R = c["left"], c["right"]
    pw = (W - 2 * M - G) / 2
    ph = BOTTOM - y0 - 0.1
    n_rows = max(4, len(L["items"][:6]), len(R["items"][:6]))
    sz = s.common(L["items"][:6] + R["items"][:6], pw - 1.15, (ph - 1.0) / n_rows, "body", max_pt=16)
    for k, side in enumerate((L, R)):
        x = M + k * (pw + G)
        good = k == 1
        fill = s.d.mix(s.bg, s.fg, 0.9) if good and c.get("contrast", True) else s.surface
        fg = s.bg if good and c.get("contrast", True) else s.fg
        mut = s.d.mix(fill, fg, 0.6)
        s.rect(x, y0, pw, ph, fill, radius=s.d.radius)
        s.label(x + 0.35, y0 + 0.3, pw - 0.7, side["title"], color=s.accent if good else mut)
        items = side["items"][:6]
        ih = (ph - 1.0) / n_rows
        for i, it in enumerate(items):
            yy = y0 + 0.85 + i * ih
            s.text(x + 0.35, yy, 0.35, 0.35, "✓" if good else "–", "h3", color=s.accent if good else mut, max_pt=16)
            s.text(x + 0.8, yy, pw - 1.15, ih, it, "body", color=fg, max_pt=sz, min_pt=sz)
    s.chrome(c.get("section"))


def a_cards(s: S, c):
    y0 = _head(s, c)
    cards = c["cards"][:6]
    n = len(cards)
    per = 3 if n in (3, 5, 6) else (2 if n == 4 else n)
    rows = (n + per - 1) // per
    cw = (W - 2 * M - (per - 1) * G) / per
    avail = BOTTOM - y0 - (rows - 1) * G
    # size cards to their content (title + text lines), never taller than the space
    hl = c.get("highlight", [])
    if rows == 1 and avail >= 2.6:
        # one row of peers: tall tiles that fill the space, numeral on top, title + text anchored to the bottom
        ch = avail - 0.12
        sz = s.common([cd.get("text") for cd in cards], cw - 0.7, min(2.0, ch - 2.2), "body", max_pt=17)
        tz = s.common([cd["title"] for cd in cards], cw - 0.7, 0.9, "h2", max_pt=28)
        # all tiles share one title line count and one text height, so titles sit on one baseline
        th = max(len(wrap(cd["title"], s.d.fonts["heading"], True, tz, (cw - 0.7) * 72)) for cd in cards) * tz * 1.05 * 1.18 / 72
        bh_ = max(len(wrap(cd.get("text") or "", s.d.fonts["body"], False, sz, (cw - 0.7) * 72)) for cd in cards) * sz * 1.35 * 1.18 / 72
        for i, cd in enumerate(cards):
            x = M + i * (cw + G)
            on = i in hl
            fill = s.accent if on else s.surface
            fg = s.d.c["accent_text"] if on else s.fg
            s.shadow(s.rect(x, y0, cw, ch, fill, radius=s.d.radius))
            s.text(x + 0.35, y0 + 0.3, cw - 0.7, 1.1, cd.get("n", f"{i + 1:02d}"), "number", color=fg if on else s.accent,
                   max_pt=64, min_pt=40)
            yb = y0 + ch - 0.35
            if cd.get("text"):
                s.text(x + 0.35, yb - bh_, cw - 0.7, bh_ + 0.05, cd["text"], "body",
                       color=s.readable(s.d.mix(fill, fg, 0.72), fill) if on else s.muted, max_pt=sz, min_pt=sz)
                yb -= bh_ + 0.18
            s.text(x + 0.35, yb - th, cw - 0.7, th + 0.05, cd["title"], "h2", color=fg, max_pt=tz, min_pt=tz, anchor="b")
            s.hline(x + 0.35, yb - th - 0.25, 0.5, fg if on else s.accent, 2.5)
        s.chrome(c.get("section"))
        return
    longest = max(len(wrap(cd.get("text") or "", s.d.fonts["body"], False, 14, (cw - 0.6) * 72)) for cd in cards)
    need = 1.45 + longest * 14 * 1.35 * 1.18 / 72 + 0.35
    ch = min(avail / rows, max(need, 2.1))
    y0 = y0 + (avail - ch * rows) * 0.35   # leftover space goes mostly below, a little above
    trole = "small" if rows > 1 else "body"
    compact = ch < 2.0                       # short cards: number sits beside the title, text gets the room
    ty, by = (0.22, 0.72) if compact else (0.8, 1.3)
    sz = s.common([cd.get("text") for cd in cards], cw - 0.6, ch - by - 0.2, trole)
    tz = s.common([cd["title"] for cd in cards], cw - 0.6, 0.5, "h3")
    for i, cd in enumerate(cards):
        r, k = divmod(i, per)
        x, y = M + k * (cw + G), y0 + r * (ch + G)
        on = i in hl
        fill = s.accent if on else s.surface
        fg = s.d.c["accent_text"] if on else s.fg
        s.shadow(s.rect(x, y, cw, ch, fill, radius=s.d.radius))
        nw = 0.62 if compact else 0
        s.text(x + 0.3, y + (0.22 if compact else 0.25), 0.6 if compact else cw - 0.6, 0.5, cd.get("n", f"{i + 1:02d}"),
               "h3" if compact else "h2", color=fg if on else s.accent, max_pt=tz if compact else 26)
        s.text(x + 0.3 + nw, y + ty, cw - 0.6 - nw, 0.5, cd["title"], "h3", color=fg, max_pt=tz, min_pt=tz)
        s.text(x + 0.3, y + by, cw - 0.6, ch - by - 0.2, cd.get("text"), trole, color=fg if on else s.muted, max_pt=sz, min_pt=sz)
    s.chrome(c.get("section"))


def a_image_text(s: S, c):
    left = c.get("image_side", "left") == "left"
    iw = W * c.get("image_share", 0.46)
    s.image(0 if left else W - iw, 0, iw, H, c["image"])
    tx = iw + 0.6 if left else M
    tw = W - iw - 0.6 - M
    s.label(tx, 1.4, tw, c.get("kicker"))
    s.text(tx, 1.75, tw, 1.9, c["title"], "h1", max_pt=48)
    s.text(tx, 3.75, tw, 1.3, c.get("text"), "lead", color=s.muted)
    for i, b in enumerate((c.get("bullets") or [])[:4]):
        yy = 5.2 + i * 0.42
        s.rect(tx, yy + 0.1, 0.09, 0.09, s.accent)
        s.text(tx + 0.25, yy, tw - 0.25, 0.4, b, "body")
    if c.get("chrome", True):
        s.chrome()


def a_quote(s: S, c):
    if c.get("image"):
        s.image(col(8), TOP, span(4), BOTTOM - TOP, c["image"], radius=s.d.radius)
    if c.get("cutout"):
        s.cutout(c.get("cutout_x", col(9)), c.get("cutout_y", TOP + 0.4), c.get("cutout_w", span(3)), c["cutout"])
    tw = span(7) if (c.get("image") or c.get("cutout")) else span(10)
    s.text(col(0), TOP - 0.1, 1.0, 1.0, "“", "display", color=s.accent, max_pt=110, font=s.d.fonts["heading"])
    s.text(col(0), TOP + 0.85, tw, 3.0, c["quote"], "h1", max_pt=c.get("max_pt", 48))
    s.hline(col(0), TOP + 4.05, 0.6, s.accent, 2.25)
    s.text(col(0), TOP + 4.25, tw, 0.4, c.get("name"), "h3")
    s.text(col(0), TOP + 4.65, tw, 0.4, c.get("role"), "small", color=s.muted)
    s.chrome(c.get("section"))


def a_team(s: S, c):
    y0 = _head(s, c)
    ppl = c["people"][:4]
    n = len(ppl)
    cw = (W - 2 * M - (n - 1) * G) / n
    ih = min(cw * 1.15, BOTTOM - y0 - 1.0)
    for i, p in enumerate(ppl):
        x = M + i * (cw + G)
        s.image(x, y0, cw, ih, p["photo"], radius=s.d.radius)
        s.text(x, y0 + ih + 0.18, cw, 0.4, p["name"], "h3")
        s.text(x, y0 + ih + 0.55, cw, 0.4, p.get("role"), "small", color=s.muted)
    s.chrome(c.get("section"))


def a_chart(s: S, c):
    s.label(col(0), TOP, span(4), c.get("kicker"))
    s.text(col(0), TOP + 0.32, span(4), 1.9, c["title"], "h1", max_pt=42)
    s.text(col(0), TOP + 2.35, span(4), 1.6, c.get("text"), "body", color=s.muted)
    if c.get("takeaway"):
        tk = c["takeaway"]
        if len(tk) <= 10:  # a figure: big
            s.text(col(0), BOTTOM - 1.6, span(4), 1.0, tk, "number", color=s.accent, max_pt=64, min_pt=36)
        else:              # a phrase: headline size, never display size
            s.text(col(0), BOTTOM - 1.75, span(4), 1.15, tk, "h2", color=s.accent, max_pt=26)
            if len(tk.split()) > 6:
                s.d.warnings.append(f"slide {s.no}: chart takeaway should be a figure or 2-6 words: {tk[:40]!r}")
        s.text(col(0), BOTTOM - 0.6, span(4), 0.4, c.get("takeaway_label"), "small", color=s.muted)
    kind = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE_MARKERS, "doughnut": XL_CHART_TYPE.DOUGHNUT}[c.get("chart", "column")]
    cats = c["categories"]
    series = c["series"]
    if s.d.rtl:  # Hebrew reads right to left: categories (and the doughnut legend) start on the right
        cats = list(reversed(cats))
        series = [dict(x, values=list(reversed(x["values"]))) for x in series]
    data = CategoryChartData()
    data.categories = cats
    for se in series:
        data.add_series(se["name"], se["values"])
    x, w = col(5), span(7)
    gf = s.slide.shapes.add_chart(kind, E(s.X(x, w)), E(TOP + 0.3), E(w), E(BOTTOM - TOP - 0.4), data)
    ch = gf.chart
    ch.has_title = False
    ch.font.size = Pt(12)
    ch.font.name = s.d.fonts["body"]
    ch.font.color.rgb = RGBColor.from_string(s.muted)
    ch.has_legend = len(series) > 1 or kind == XL_CHART_TYPE.DOUGHNUT
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    palette = [s.accent, s.fg, s.muted, s.line]
    plot = ch.plots[0]
    if kind == XL_CHART_TYPE.DOUGHNUT:
        hl_i = c.get("highlight", 0)
        if s.d.rtl:
            hl_i = len(cats) - 1 - hl_i
        quiet = [s.muted, s.d.mix(s.muted, s.bg, 0.45), s.d.mix(s.muted, s.bg, 0.7), s.line]
        for i, pt in enumerate(plot.series[0].points):  # accent only on the segment the takeaway is about
            col_ = s.accent if i == hl_i else quiet[(i - (1 if i > hl_i else 0)) % len(quiet)]
            pt.format.fill.solid(); pt.format.fill.fore_color.rgb = RGBColor.from_string(col_)
            pt.format.line.color.rgb = RGBColor.from_string(s.bg); pt.format.line.width = Pt(2)
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(14); dl.font.bold = True; dl.font.color.rgb = RGBColor.from_string(s.bg)
        dl.number_format = c.get("format", "General"); dl.number_format_is_linked = False
        dl.show_value = True; dl.show_category_name = False; dl.show_percentage = False
        ch.legend.font.size = Pt(13)
    else:
        hl = c.get("highlight")
        for i, se in enumerate(plot.series):
            col_ = palette[i % len(palette)] if len(series) > 1 else (s.muted if hl is not None else s.accent)
            if kind == XL_CHART_TYPE.LINE_MARKERS:
                se.format.line.color.rgb = RGBColor.from_string(col_); se.format.line.width = Pt(3)
                se.smooth = False
            else:
                se.format.fill.solid(); se.format.fill.fore_color.rgb = RGBColor.from_string(col_)
            if hl is not None and len(series) == 1 and kind != XL_CHART_TYPE.LINE_MARKERS:
                idx = (len(cats) - 1 - hl) if s.d.rtl else hl
                p = se.points[idx]; p.format.fill.solid(); p.format.fill.fore_color.rgb = RGBColor.from_string(s.accent)
        try:
            plot.gap_width = 60
        except Exception:
            pass
        va = ch.value_axis
        if kind != XL_CHART_TYPE.LINE_MARKERS:
            va.minimum_scale = 0  # bars always start at zero: no exaggerated differences
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string(s.line)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(10)
        ca = ch.category_axis
        ca.format.line.color.rgb = RGBColor.from_string(s.line)
        ca.tick_labels.font.size = Pt(13)
        clean = c.get("clean", len(series) == 1 and kind != XL_CHART_TYPE.LINE_MARKERS)
        if clean:  # designer default: values sit on the bars, no axis or gridlines competing with them
            va.visible = False
            va.has_major_gridlines = False
            ca.tick_labels.font.size = Pt(14)
        if s.d.rtl:
            va_el = ch._chartSpace.find(".//" + qn("c:valAx"))
            cr = va_el.find(qn("c:crosses"))
            if cr is not None: cr.set("val", "max")
        if c.get("labels", True):
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(18 if clean else 11); dl.font.bold = True
            dl.font.color.rgb = RGBColor.from_string(s.fg)
            dl.font.name = s.d.fonts["heading"]
            dl.number_format = c.get("format", "General"); dl.number_format_is_linked = False
            try:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            except Exception:
                pass
    s.chrome(c.get("section"))


def _plan_card(s: S, x, y, w, h, p, fill, fg, price_col, label_col, big_pt, button, fz=None):
    """One plan: name, display price (+unit, period), rule, features, and a pill button pinned to the bottom."""
    mut = s.readable(s.d.mix(fill, fg, 0.68), fill, toward=fg)
    pad = 0.42
    s.text(x + pad, y + 0.42, w - 2 * pad, 0.3, p["name"], "label", color=label_col, caps=True,
           tracking=1.2 if s.d.caps_labels else 0)
    yb = metric(s, x + pad, y + 0.78, w - 2 * pad, big_pt / 72 * 1.12, p["price"], p.get("unit"), color=price_col,
                max_pt=big_pt, min_pt=34)
    if p.get("period"):
        s.text(x + pad, yb + 0.06, w - 2 * pad, 0.34, p["period"], "body", color=mut, max_pt=15)
        yb += 0.42
    s.hline(x + pad, yb + 0.22, w - 2 * pad, s.d.mix(fill, fg, 0.3), 0.75)
    feats = p.get("features", [])[:6]
    fy = yb + 0.42
    bottom = y + h - (1.05 if p.get("cta") else 0.35)
    fh = min(0.5, (bottom - fy) / max(1, len(feats)))
    fit = s.common(feats, w - 2 * pad - 0.4, fh, "body", max_pt=16, min_pt=11) if feats else 16
    fz = min(fz, fit) if fz else fit
    for k, f in enumerate(feats):
        s.text(x + pad, fy + k * fh, 0.32, fh, "✓", "body", color=price_col, max_pt=fz, anchor="m", warn=False)
        s.text(x + pad + 0.4, fy + k * fh, w - 2 * pad - 0.4, fh, f, "body", color=fg, max_pt=fz, min_pt=fz, anchor="m")
    if p.get("cta"):
        bfill, bink = button
        bh = 0.56
        s.rect(x + pad, y + h - 0.35 - bh, w - 2 * pad, bh, bfill, radius=bh / 2)
        s.text(x + pad + 0.2, y + h - 0.35 - bh, w - 2 * pad - 0.4, bh, p["cta"], "h3", color=bink, align="c",
               anchor="m", max_pt=16)


def a_pricing(s: S, c):
    plans = c["plans"][:3]
    n = len(plans)
    hl = c.get("highlight", [n // 2 if n > 1 else 0])
    acc, acc_ink = s.d.c["accent"], s.d.c["accent_text"]
    if n == 1:
        # single offer: the argument on the left, the offer on a full-height panel in the inverse ground on the right
        p = plans[0]
        px0 = col(7) - 0.3
        panel, ink = s.fg, s.bg
        s.rect(px0, 0, W - px0, H, panel)
        s.fields.append((px0, W - px0, s.readable(s.d.mix(panel, ink, 0.7), panel, toward=ink)))
        s.photos.append((px0, W - px0))
        s.label(col(0), TOP + 0.35, span(6), c.get("kicker"))
        s.text(col(0), TOP + 0.7, span(6) + 0.2, 2.3, c["title"], "h1", max_pt=c.get("title_pt", 56))
        y = TOP + 0.7 + s.last_h + 0.35
        s.text(col(0), y, span(6), 1.4, c.get("text"), "lead", color=s.muted)
        if c.get("text"):
            y += s.last_h + 0.4
        pts = (c.get("points") or [])[:4]
        if pts:
            y = max(y, BOTTOM - 0.55 * len(pts) - 0.1)
            for i, pt in enumerate(pts):
                yy = y + i * 0.55
                s.hline(col(0), yy, span(6) - 0.3)
                s.text(col(0), yy + 0.1, 0.4, 0.4, "✓", "body", color=s.ink, max_pt=16, anchor="m", warn=False)
                s.text(col(0) + 0.45, yy + 0.1, span(6) - 0.75, 0.4, pt, "body", max_pt=16, anchor="m")
        price_col = acc if contrast_ratio(acc, panel) >= 3 else ink
        lab = s.readable(acc, panel, 4.5, ink) if contrast_ratio(acc, panel) >= 3 else s.readable(s.d.mix(panel, ink, 0.7), panel, toward=ink)
        bt = (acc, acc_ink) if contrast_ratio(acc, panel) >= 1.6 else (ink, panel)
        _plan_card(s, px0 + 0.3, TOP - 0.05, W - M - px0 - 0.3 + 0.42, BOTTOM - TOP + 0.05, p, panel, ink, price_col, lab,
                   120, bt)
        s.chrome(c.get("section"))
        return
    if n == 2:  # argument column + two tall plan cards
        s.label(col(0), TOP, span(4), c.get("kicker"))
        s.text(col(0), TOP + 0.32, span(4), 2.4, c["title"], "h1", max_pt=c.get("title_pt", 48))
        s.text(col(0), TOP + 0.32 + s.last_h + 0.3, span(4), 2.0, c.get("text"), "lead", color=s.muted)
        x0, top = col(4) + 0.2, TOP
        cw = (W - M - x0 - G) / 2
    else:
        s.label(col(0), TOP, span(10), c.get("kicker"))      # header sized to the real title, so cards get the room
        s.text(col(0), TOP + 0.32, span(10), 1.15, c["title"], "h1", max_pt=c.get("title_pt", 46))
        top = TOP + 0.32 + s.last_h + 0.3
        if c.get("text"):
            s.text(col(0), top - 0.1, span(10), 0.75, c["text"], "lead", color=s.muted)
            top += s.last_h + 0.2
        top += 0.3                                             # headroom for the raised card
        x0 = M
        cw = (W - 2 * M - 2 * G) / 3
    fz_all = s.common([f for q in plans for f in q.get("features", [])[:6]], cw - 0.84 - 0.4, 0.46, "body",
                      max_pt=16, min_pt=12)
    for i, p in enumerate(plans):
        on = i in hl
        x = x0 + i * (cw + G)
        big = 88 if n == 2 else 76
        # every card must hold its features at a readable rhythm: shrink the price before squeezing the list
        for q in plans:
            room = BOTTOM - top + (0.3 if n == 3 else 0) - (0.78 + (0.42 if q.get("period") else 0) + 0.42
                                                             + 0.48 * len(q.get("features", [])[:6]) + (1.05 if q.get("cta") else 0.35))
            big = max(48, min(big, room * 72 / 1.12))
        y = top
        if not any(q.get("cta") for q in plans):  # no button to pin to the bottom: cards hug their content,
            need = max(0.78 + big / 72 * 1.12 + (0.42 if q.get("period") else 0) + 0.42
                       + 0.5 * len(q.get("features", [])[:6]) + 0.4 for q in plans)
            y = max(top, BOTTOM - need)         # standing on the footer line, so nothing floats
        y -= 0.3 if on and n == 3 else 0
        h = BOTTOM - y
        fill = acc if on else s.surface
        fg = acc_ink if on else s.fg
        card = s.rect(x, y, cw, h, fill, radius=s.d.radius)
        s.shadow(card, alpha=0.16 if on else 0.08)
        price_col = fg if on else (s.accent if contrast_ratio(s.accent, fill) >= 3 else fg)
        lab = fg if on else s.readable(s.accent, fill)
        bt = (fg, fill) if on else (acc, acc_ink)
        if p.get("badge"):
            bw = min(cw - 0.84, text_width(p["badge"], s.d.fonts["body"], True, 10) / 72 + 0.5)
            s.rect(x + cw - 0.42 - bw, y + 0.36, bw, 0.36, fg if on else acc, radius=0.18)
            s.text(x + cw - 0.42 - bw, y + 0.36, bw, 0.36, p["badge"], "label", color=fill if on else acc_ink, align="c",
                   anchor="m", max_pt=10, caps=True, warn=False)
        _plan_card(s, x, y, cw, h, p, fill, fg, price_col, lab, big, bt, fz=fz_all)
    s.chrome(c.get("section"))


def a_gallery(s: S, c):
    y0 = _head(s, c)
    imgs = c["images"][:3]
    n = len(imgs)
    weights = c.get("weights") or ([2, 1] if n == 2 else [1] * n)
    tot = sum(weights)
    avail = W - 2 * M - (n - 1) * G
    x = M
    ih = BOTTOM - y0 - (0.5 if any(isinstance(i, dict) and i.get("caption") for i in imgs) else 0)
    for i, im in enumerate(imgs):
        im = im if isinstance(im, dict) else {"path": im}
        w = avail * weights[i] / tot
        s.image(x, y0, w, ih, im, radius=s.d.radius)
        if im.get("caption"):
            s.text(x, y0 + ih + 0.12, w, 0.35, im["caption"], "small", color=s.muted)
        x += w + G
    s.chrome(c.get("section"))


def a_closing(s: S, c):
    split = c.get("variant") == "split" and c.get("image")   # mirrors the split cover: photo right, ask left
    if split:
        s.image(col(7), 0, W - col(7), H, c["image"])
    elif c.get("image"):
        s.bg_image(c["image"], c.get("dim"))
    tw = span(6) if split else span(10)
    s.label(col(0), 1.0, tw, c.get("kicker"))
    two = bool(c.get("title2"))
    s.text(col(0), 1.35, tw, 1.75 if two else 3.3, c["title"], "display", max_pt=104 if not split else 84, anchor="b")
    if two:
        s.text(col(0), 3.12, tw, 1.75, c["title2"], "display", color=s.accent, max_pt=104 if not split else 84, anchor="t")
    rows = c.get("contacts", [])[:4]
    if c.get("image") and not split and rows:  # contact rows on a photo get a quiet panel so they read
        s.rect(col(0) - 0.25, 4.85, span(7) + 0.5, 0.45 + len(rows) * 0.4, s.bg, alpha=0.72, radius=s.d.radius)
    s.hline(col(0), 5.05, span(6) if split else span(7))
    for i, r in enumerate(rows):
        s.text(col(0), 5.27 + i * 0.4, span(2), 0.34, r[0], "label", color=s.ink, caps=True, anchor="m")
        s.text(col(2), 5.24 + i * 0.4, span(4 if split else 5), 0.38, r[1], "body", bold=True, anchor="m")
    if c.get("note") and not split:
        s.rect(col(8), 4.9, span(4), 1.5, s.surface if not c.get("image") else s.bg, radius=s.d.radius)
        s.text(col(8) + 0.3, 5.05, span(4) - 0.6, 1.2, c["note"], "lead", anchor="m")
    s.chrome()


# ------------------------------------------------------------------ bold / editorial compositions
def v_cover_block(s: S, c):
    """Poster cover: stacked title on the left, colour field on the right with an arched photo crossing its edge."""
    fx = W * 0.64
    s.rect(fx, 0, W - fx, H, s.accent)
    s.full_bleed = True
    if c.get("image"):
        s.framed(fx - 1.55, 1.05, 3.9, 5.35, c["image"], mask=c.get("mask", s.d.photo_mask or "round"), block=False)
    s.label(col(0), 1.0, span(7), c.get("kicker"))
    s.text(col(0), 1.35, fx - 1.9 - M, 3.7, c["title"], "display", max_pt=118, ls=0.88, anchor="b")
    s.hline(col(0), 5.3, 0.9, s.accent, 3)
    s.text(col(0), 5.5, fx - 2.2 - M, 1.0, c.get("subtitle"), "lead", color=s.muted)
    s.chrome()


def v_statement_block(s: S, c):
    """Statement on a full accent field, optional outlined ghost word behind it."""
    field = s.d.c["accent"]                      # the brand accent, whatever the slide tone
    s.rect(0, 0, W, H, field)
    s.full_bleed = True
    ink = s.d.c["accent_text"]
    hot = s.d.c["bg"] if contrast_ratio(s.d.c["bg"], field) >= 3 else s.d.c["text"]
    if c.get("ghost"):
        s.text(-0.3, 3.4, W + 0.6, 4.4, c["ghost"], "display", color=s.d.mix(field, ink, 0.28), max_pt=300,
               min_pt=200, outline=1.5, anchor="b", warn=False)
    s.text(col(0), 1.0, span(10), 0.3, c.get("kicker"), "label", color=ink, caps=True, tracking=1.2 if s.d.caps_labels else 0)
    s.text(col(0), 1.4, span(c.get("width", 11)), 3.9, c["title"], "display", color=ink, accent_color=hot if contrast_ratio(hot, field) >= 3 else ink,
           max_pt=c.get("max_pt", 110), ls=0.9)
    s.text(col(0), 5.5, span(6), 1.0, c.get("text"), "lead", color=ink)
    if c.get("chrome") is False:
        return
    s.text(col(0), H - 0.52, span(6), 0.24, s.d.brand.get("label", ""), "label", color=ink, caps=True)
    s.text(col(9), H - 0.52, span(3), 0.24, f"{s.no:02d} / {s.d.total:02d}", "label", color=ink, align="r", bold=False)


def v_section_ghost(s: S, c):
    """Section with a giant outlined numeral running off the slide."""
    num = c.get("number", f"{s.no:02d}")
    # the outlined numeral lives between the header rule and the footer, clear of the title column
    s.text(col(6) + 0.4, 0.85, W - col(6) + 0.8, BOTTOM - 0.85, num, "number", color=s.d.mix(s.bg, s.fg, 0.25),
           max_pt=420, min_pt=260, outline=1.5, anchor="m", warn=False)
    if c.get("cutout"):
        s.cutout(c.get("cutout_x", col(7)), c.get("cutout_y", 0.95), c.get("cutout_w", 3.3), c["cutout"])
    s.label(col(0), 2.55, span(6), c.get("kicker"))
    s.text(col(0), 2.9, span(6), 2.1, c["title"], "display", max_pt=88, ls=0.9)
    s.text(col(0), 5.15, span(5), 1.2, c.get("text"), "lead", color=s.muted)
    s.chrome(c.get("section"))


def v_kpi_hero(s: S, c):
    """One number carries the slide; the others support it in a column."""
    m = c["metrics"][:4]
    hi = (c.get("highlight") or [0])[0]           # the highlighted metric is the hero
    hero, rest = m[hi], [x for k, x in enumerate(m) if k != hi]
    s.label(col(0), TOP, span(7), c.get("kicker"))
    s.text(col(0), TOP + 0.32, span(7), 1.3, c["title"], "h1", max_pt=40)
    y = TOP + 0.32 + s.last_h + 0.3                 # two-line titles push the hero down instead of colliding
    if c.get("text"):                              # source / context note: never dropped
        s.text(col(0), y, span(7), 0.7, c["text"], "body", color=s.muted, max_pt=15)
        y += s.last_h + 0.2
    yb = metric(s, col(0), y, span(7), BOTTOM - y - 1.1, hero["value"], hero.get("unit"), color=s.accent,
                max_pt=220, min_pt=110)
    s.hline(col(0), yb + 0.12, 0.9, s.accent, 3)
    s.text(col(0), yb + 0.32, span(6), BOTTOM - yb - 0.35, hero["label"], "lead")
    if rest:
        rh = min(1.9, (BOTTOM - TOP) / len(rest))
        y_start = max(TOP + 0.2, yb + 0.25 - rh * len(rest))  # the column ends on the hero number's baseline
        for i, it in enumerate(rest):
            y = y_start + i * rh
            s.hline(col(8), y, span(4))
            yv = metric(s, col(8), y + 0.2, span(4), rh * 0.5, it["value"], it.get("unit"), max_pt=60, min_pt=34)
            s.text(col(8), yv + 0.1, span(4), y + rh - yv - 0.15, it["label"], "body", color=s.muted)
    s.chrome(c.get("section"))


def v_kpi_bars(s: S, c):
    """Stepped bands: each metric is a band whose length shows its weight."""
    y0 = _head(s, c)
    m = c["metrics"][:4]
    n = len(m)
    bh = min(1.0, (BOTTOM - y0 - 0.1) / n)
    weights = c.get("weights") or [1 - i * (0.55 / max(1, n - 1)) for i in range(n)]
    top_w = max(weights) or 1
    weights = [max(0.35, w / top_w) for w in weights]   # any scale (0-1, %, raw numbers) -> longest band = full width
    for i, it in enumerate(m):
        y = y0 + i * bh
        w = (W - M) * weights[i]
        fill = s.accent if i == 0 else s.d.mix(s.bg, s.fg, 0.12 + 0.1 * i)
        ink = s.d.c["accent_text"] if i == 0 else s.fg
        s.rect(0, y, w, bh - 0.06, fill)
        s.text(M, y, w - M - 3.2, bh - 0.06, it["label"], "h3", color=ink, anchor="m", caps=True, tracking=0.5)
        s.text(w - 3.0, y, 2.75, bh - 0.06, f'{it["value"]}{" " + it["unit"] if it.get("unit") else ""}', "h1", color=ink, anchor="m", align="r", max_pt=44)
    s.chrome(c.get("section"))


def v_cards_rows(s: S, c):
    """Editorial numbered rows: big numerals, title and text on hairlines."""
    cards = c["cards"][:5]
    s.label(col(0), TOP, span(4), c.get("kicker"))
    s.text(col(0), TOP + 0.35, span(4), 2.6, c["title"], "h1", max_pt=48)
    s.text(col(0), TOP + 3.1, span(4), 1.8, c.get("text"), "body", color=s.muted)
    rh = (BOTTOM - TOP) / len(cards)
    hl = c.get("highlight", [])
    sz = s.common([cd.get("text") for cd in cards], span(6) - 0.5, rh * 0.55 - 0.15, "body")
    tz = s.common([cd["title"] for cd in cards], span(6) - 0.5, rh * 0.42, "h3")
    for i, cd in enumerate(cards):
        y = TOP + i * rh
        s.hline(col(5), y, span(7))
        s.text(col(5), y + 0.12, span(1) + 0.4, rh - 0.2, cd.get("n", f"{i + 1:02d}"), "number",
               color=s.accent if i in hl else s.fg, max_pt=min(64, rh * 50), min_pt=28)
        s.text(col(6) + 0.5, y + 0.15, span(6) - 0.5, rh * 0.42, cd["title"], "h3", max_pt=tz, min_pt=tz)
        s.text(col(6) + 0.5, y + 0.15 + rh * 0.4, span(6) - 0.5, rh * 0.55 - 0.15, cd.get("text"), "body", color=s.muted,
               max_pt=sz, min_pt=sz)
    s.hline(col(5), BOTTOM, span(7))
    s.chrome(c.get("section"))


def v_image_framed(s: S, c):
    """Photo inside the margins with an offset colour block and an arch/round/circle mask; text beside it."""
    left = c.get("image_side", "left") == "left"
    iw, ih = span(5), BOTTOM - TOP - 0.3
    ix = col(0) if left else col(7)
    # screenshots (fit: contain) are never masked or cropped; photos default to the arch
    mask = None if c.get("fit") == "contain" else c.get("mask", s.d.photo_mask)
    s.framed(ix, TOP, iw, ih, c["image"], mask=mask, offset=0.3, block=None if c.get("offset", s.d.photo_offset) else False)
    tx = col(6) + 0.3 if left else col(0)
    tw = span(6) - 0.3
    s.label(tx, TOP + 0.5, tw, c.get("kicker"))
    s.text(tx, TOP + 0.85, tw, 2.0, c["title"], "h1", max_pt=50)
    s.text(tx, TOP + 2.95, tw, 1.3, c.get("text"), "lead", color=s.muted)
    for i, b in enumerate((c.get("bullets") or [])[:4]):
        yy = TOP + 4.35 + i * 0.44
        s.text(tx, yy, 0.5, 0.4, f"{i + 1:02d}", "label", color=s.ink, anchor="m")
        s.text(tx + 0.55, yy, tw - 0.55, 0.4, b, "body", anchor="m")
    s.chrome(c.get("section"))


def v_image_full(s: S, c):
    """Few words + a strong photo: the photo takes the whole slide, the line sits on it (instead of a 50/50 split
    that leaves half the slide empty)."""
    s.bg_image(c["image"], c.get("dim"))
    s.text(col(0), 1.6, span(c.get("width", 8)), 3.4, c["title"], "display", max_pt=c.get("max_pt", 84), anchor="b")
    s.label(col(0), 5.0 - s.last_h - 0.45, span(6), c.get("kicker"))
    s.text(col(0), 5.25, span(6), 1.0, c.get("text"), "lead")
    s.chrome(c.get("section"))


def v_comparison_split(s: S, c):
    """Two full-height halves: the old way in a quiet tone, the new way on the accent."""
    L, R = c["left"], c["right"]
    s.rect(0, 0, W / 2, H, s.surface)
    s.rect(W / 2, 0, W / 2, H, s.accent)
    s.full_bleed = True
    ink = s.d.c["accent_text"]
    for k, (side, fg, mut, mark) in enumerate(((L, s.fg, s.muted, "–"), (R, ink, s.d.mix(s.accent, ink, 0.75), "✓"))):
        x = M if k == 0 else W / 2 + M
        w = W / 2 - 2 * M
        s.text(x, 0.9, w, 0.3, side["title"], "label", color=mut, caps=True, tracking=1.2 if s.d.caps_labels else 0)
        head = side.get("headline") or (c["title"] if k == 1 else c.get("left_headline", ""))
        s.text(x, 1.25, w, 1.4, head, "h1", color=fg, max_pt=40)
        items = side["items"][:5]
        ih = min(0.75, (BOTTOM - 3.0) / len(items))
        sz = s.common(L["items"][:5] + R["items"][:5], w - 0.5, ih - 0.1, "body", max_pt=17)
        for i, it in enumerate(items):
            y = 3.0 + i * ih
            s.hline(x, y, w, s.d.mix(s.surface if k == 0 else s.accent, fg, 0.25), 0.75)
            s.text(x, y + 0.08, 0.4, ih - 0.1, mark, "h3", color=mut if k == 0 else ink, anchor="m", max_pt=18)
            s.text(x + 0.5, y + 0.08, w - 0.5, ih - 0.1, it, "body", color=fg, anchor="m", max_pt=sz, min_pt=sz)
    s.text(M, H - 0.52, W / 2 - 2 * M, 0.24, s.d.brand.get("label", ""), "label", color=s.muted, caps=True)
    s.text(W / 2 + M, H - 0.52, W / 2 - 2 * M, 0.24, f"{s.no:02d} / {s.d.total:02d}", "label", color=ink, align="r", bold=False)


def v_closing_block(s: S, c):
    """Closing on the accent field: the ask in giant type, contacts below."""
    s.rect(0, 0, W, H, s.accent)
    s.full_bleed = True
    ink = s.d.c["accent_text"]
    s.text(col(0), 1.0, span(8), 0.3, c.get("kicker"), "label", color=ink, caps=True, tracking=1.2 if s.d.caps_labels else 0)
    s.text(col(0), 1.35, span(11), 1.9, c["title"], "display", color=ink, max_pt=120, anchor="b", ls=0.88)
    if c.get("title2"):  # second line in whichever brand colour is NOT the ink, so the two lines contrast
        alt = s.d.c["bg"] if contrast_ratio(ink, s.d.c["text"]) < contrast_ratio(ink, s.d.c["bg"]) else s.d.c["text"]
        s.text(col(0), 3.2, span(11), 1.9, c["title2"], "display", color=alt, max_pt=120, ls=0.88)
    s.hline(col(0), 5.2, span(12), s.d.mix(s.accent, ink, 0.4))
    for i, r in enumerate((c.get("contacts") or [])[:3]):
        x = col(i * 4)
        s.text(x, 5.4, span(4), 0.3, r[0], "label", color=s.d.mix(s.accent, ink, 0.7), caps=True)
        s.text(x, 5.72, span(4), 0.45, r[1], "h3", color=ink, max_pt=20)


# ------------------------------------------------------------------ long-text archetypes
def _paras(t):
    """Body text: list of paragraphs or one string with blank lines -> single string with \n between paragraphs."""
    if isinstance(t, list):
        return "\n".join(t)
    return re.sub(r"\n\s*\n", "\n", t or "")


def _balance(paras, k=2):
    """Split paragraphs into two columns of similar length (by words), keeping reading order."""
    if len(paras) <= 1:
        return [paras, []]
    words = [len(p.split()) for p in paras]
    tot, best = sum(words), None
    for cut in range(1, len(paras)):
        diff = abs(sum(words[:cut]) - tot / 2)
        if best is None or diff < best[0]:
            best = (diff, cut)
    return [paras[:best[1]], paras[best[1]:]]


def _lines_h(s, txt, w, size, role="body", bold=None, ls=None):
    """Height (in) a block of text takes at a given size in a given width."""
    if not txt:
        return 0
    mx, mn, lsp, bld, frole = SCALE[role]
    b = bld if bold is None else bold
    # measured on a slightly narrower measure: PowerPoint breaks a little earlier than the font metrics suggest,
    # and stacked paragraphs must never collide
    return len(wrap(str(txt).replace("*", ""), s.d.fonts[frole], b, size, w * 72 * 0.965)) * size * (ls or lsp) * 1.18 / 72


def _flow_size(s, paras, w, h, max_pt=21, min_pt=12, gap=0.75, ls=1.42):
    """Largest body size at which the paragraphs, stacked with `gap` (in ems) between them, fit w x h."""
    sz = max_pt
    while sz > min_pt:
        need = sum(_lines_h(s, p_, w, sz, ls=ls) for p_ in paras) + (len(paras) - 1) * sz * gap / 72
        if need <= h - 0.05:
            break
        sz = max(min_pt, sz * 0.95)
    return sz


def _flow(s, x, y, w, h, paras, sz, colors, gap=0.75, ls=1.42):
    """Stack paragraphs as separate boxes, so each gets the no-widow rebalance. colors: per paragraph, the last
    entry repeats. Warns (once) when the stack runs past h."""
    y_end = y + h
    for i, para in enumerate(paras):
        hh = _lines_h(s, para, w, sz, ls=ls)
        s.text(x, y, w, hh + 0.08, para, "body", color=colors[min(i, len(colors) - 1)], max_pt=sz, min_pt=sz, ls=ls)
        y += hh + sz * gap / 72
    if y - sz * gap / 72 > y_end + 0.05:
        s.d.warnings.append(f"slide {s.no}: body text runs past its column at {sz:.0f}pt - shorten it")
    return y


def _title_block(s: S, c, x, y, w, h, max_pt=44):
    """Kicker + title + short accent rule, sized to the real title. Returns the y under the rule."""
    s.label(x, y, w, c.get("kicker"))
    s.text(x, y + 0.35, w, h, c["title"], "h1", max_pt=max_pt)
    yb = y + 0.35 + s.last_h + 0.32
    s.hline(x, yb, 0.8, s.accent, 2.25)
    return yb + 0.3


def a_text(s: S, c):
    """Reading slide. essay: title column + one long column. columns: two balanced columns, or one column and a
    pull quote. Type is set as large as the copy allows (up to 21pt) with generous leading, like a printed report."""
    v = c.get("variant", "essay" if not c.get("quote") else "columns")
    body = _paras(c["body"])
    if v == "essay":
        y = _title_block(s, c, col(0), TOP, span(4), 2.6, 44)
        bx, bw = col(5) + 0.2, span(7) - 0.2
        top = TOP + 0.42
        paras = body.split("\n")
        sz = _flow_size(s, paras, bw, BOTTOM - top)
        # the standfirst sits in the reading ink, never larger than the body it introduces
        s.text(col(0), y, span(4) - 0.2, BOTTOM - y, c.get("lead"), "lead", color=s.body, max_pt=max(15, min(20, sz)))
        s.vline(col(5) - G / 2 + 0.05, top + 0.05, BOTTOM - top - 0.2, s.line)
        # the opening paragraph carries full-strength ink, the rest the reading ink: a quiet way in
        _flow(s, bx, top, bw, BOTTOM - top, paras, sz, [s.fg, s.body])
    else:  # columns
        y0 = _head(s, c, width=10, tight=True)
        has_q = bool(c.get("quote"))
        paras = body.split("\n")
        if has_q:  # quote takes the side: text runs as one comfortable column
            tw = span(7)
            _flow(s, M, y0, tw, BOTTOM - y0, paras, _flow_size(s, paras, tw, BOTTOM - y0, 19), [s.fg])
            qx, qw = col(8), span(4)
            qz = s.fit_size(c["quote"], qw - 0.4, BOTTOM - y0 - 0.9, "h2", max_pt=30)
            qh = _lines_h(s, c["quote"], qw - 0.4, qz, "h2")
            s.text(qx + 0.4, y0, qw - 0.4, qh + 0.1, c["quote"], "h2", max_pt=qz, min_pt=qz)
            by = y0 + qh + 0.3
            s.hline(qx + 0.4, by, 0.45, s.accent, 1.5)
            s.text(qx + 0.4, by + 0.14, qw - 0.4, 0.5, c.get("quote_by"), "small", color=s.body, max_pt=13)
            bh = (0.14 + s.last_h) if c.get("quote_by") else 0
            s.vline(qx, y0 + 0.08, by + bh - y0 - 0.08, s.accent, 2.25)
        else:
            cols = [ch for ch in _balance(paras) if ch]
            cw = (W - 2 * M - G * 3) / 2
            sz = min(_flow_size(s, ch, cw, BOTTOM - y0, 19) for ch in cols)
            for k, chunk in enumerate(cols):
                _flow(s, M + k * (cw + G * 3), y0, cw, BOTTOM - y0, chunk, sz, [s.fg])
    s.chrome(c.get("section"))


def a_case(s: S, c):
    """Case study: challenge -> approach -> result as paragraphs, with the headline result as a big number.
    The result column is the payoff: accent rule, full-strength ink; the first two read in the reading ink."""
    s.label(col(0), TOP, span(8), c.get("kicker"))
    s.text(col(0), TOP + 0.32, span(8), 1.3, c["title"], "h1", max_pt=44)
    yt = TOP + 0.32 + s.last_h
    if c.get("metric"):
        s.text(col(8), TOP - 0.05, span(4), 1.25, c["metric"]["value"], "number", color=s.accent, max_pt=88, min_pt=28,
               align="r")
        s.text(col(8), TOP - 0.05 + s.last_h + 0.12, span(4), 0.5, c["metric"]["label"], "small", color=s.body, align="r")
    y0 = max(TOP + 1.95, yt + 0.5)
    parts = [("challenge", c.get("challenge_label", "Challenge")), ("approach", c.get("approach_label", "Approach")),
             ("result", c.get("result_label", "Result"))]
    parts = [(k, lab) for k, lab in parts if c.get(k)]
    cw = (W - 2 * M - (len(parts) - 1) * G * 2) / len(parts)
    sz = min(_flow_size(s, _paras(c[k]).split("\n"), cw, BOTTOM - y0 - 0.72, 20, gap=0.55, ls=1.4) for k, _ in parts)
    for i, (k, lab) in enumerate(parts):
        x = M + i * (cw + G * 2)
        res = k == "result"
        s.hline(x, y0, cw, s.accent if res else s.line, 3 if res else 1)
        s.text(x, y0 + 0.22, cw, 0.3, f"{i + 1:02d}  {lab}", "label", color=s.ink if res else s.fg, caps=True,
               tracking=1 if s.d.caps_labels else 0)
        _flow(s, x, y0 + 0.7, cw, BOTTOM - y0 - 0.72, _paras(c[k]).split("\n"), sz, [s.fg if res else s.body],
              gap=0.55, ls=1.4)
    s.chrome(c.get("section"))


def _point(s: S, x, y, w, h, i, p, tz, sz, stack=True, num_pt=40, lead="lead", body="text"):
    """One detailed point in a grid cell: hairline on top with an accent tick at its start, numeral, bold lead, text."""
    s.hline(x, y, w, s.line, 1)
    s.hline(x, y - 0.012, 0.45, s.accent, 2.5)
    n = p.get("n", f"{i + 1:02d}")
    if stack:  # numeral above the lead
        s.text(x, y + 0.22, w, num_pt / 72 * 1.1, n, "number", color=s.accent, max_pt=num_pt, min_pt=num_pt, warn=False)
        ty = y + 0.22 + num_pt / 72 * 1.1 + 0.1
        tx, tw = x, w
    else:      # numeral beside the lead, on its first line
        s.text(x, y + 0.22, 0.7, tz / 72 * 1.4, n, "h3", color=s.accent, max_pt=tz, min_pt=tz, warn=False)
        ty, tx, tw = y + 0.22, x + 0.75, w - 0.75
    lh = _lines_h(s, p[lead], tw, tz, "h3")
    s.text(tx, ty, tw, lh + 0.06, p[lead], "h3", max_pt=tz, min_pt=tz)
    s.text(tx, ty + lh + 0.14, tw, max(0.3, y + h - (ty + lh + 0.14)), p.get(body), "body", color=s.body,
           max_pt=sz, min_pt=sz, ls=1.4)


def _grid_sizes(s, items, lead_key, text_key, tw, cell_h, head_h, tmax=24, bmax=19):
    """Shared lead / text sizes for a grid of cells: the largest that every cell can hold. When the cells are
    tight the lead steps down first (to 16pt), so the answers / explanations never drop under a readable 13pt."""
    texts = [it.get(text_key) for it in items if it.get(text_key)]
    tz = s.common([it[lead_key] for it in items], tw, 0.8, "h3", max_pt=tmax, min_pt=16)
    while True:
        lh = max(_lines_h(s, it[lead_key], tw, tz, "h3") for it in items)
        room = cell_h - 0.22 - head_h - lh - 0.14 - 0.05
        sz = bmax
        while texts and sz > 12 and max(_lines_h(s, t, tw, sz, ls=1.4) for t in texts) > room:
            sz = max(12, sz * 0.95)
        if sz >= 13.5 or tz <= 16:
            return tz, sz
        tz = max(16, tz * 0.92)


def _title_cell(s: S, c, x, y, w, h):
    """The slide's title set inside the first grid cell (used when the item count leaves one cell free)."""
    s.label(x, y + 0.02, w, c.get("kicker"))
    s.text(x, y + 0.37, w, h - 0.45, c["title"], "h1", max_pt=c.get("title_pt", 40))
    if c.get("text"):
        ty = y + 0.37 + s.last_h + 0.25
        s.text(x, ty, w - 0.2, y + h - ty, c["text"], "body", color=s.body, max_pt=16)


def _five_cells(n_cells=5):
    """3x2 grid geometry with the first cell kept for the title. Returns (title cell, item cells)."""
    cw = (W - 2 * M - 2 * G * 2) / 3
    ch = (BOTTOM - TOP - 0.4) / 2
    cells = []
    for i in range(1, n_cells + 1):
        r, k = divmod(i, 3)
        cells.append((M + k * (cw + G * 2), TOP + 0.1 + r * (ch + 0.4), cw, ch - 0.1))
    return (M, TOP, cw, ch), cells


def a_bullets(s: S, c):
    """Detailed points: each a bold lead line and one or two full sentences, set in a grid that fills the slide.
    3 or fewer: columns. 4: 2x2. 5: title cell + five cells (3x2). 6: header + two columns of three."""
    pts = c["points"][:6]
    n = len(pts)
    if n == 5:
        tc, cells = _five_cells()
        tz, sz = _grid_sizes(s, pts, "lead", "text", tc[2], cells[0][3], 34 / 72 * 1.1 + 0.1, 24, 18)
        _title_cell(s, c, *tc)
        for i, (p, cell) in enumerate(zip(pts, cells)):
            _point(s, *cell, i, p, tz, sz, num_pt=34)
        s.chrome(c.get("section"))
        return
    y0 = _head(s, c, width=10, tight=True)
    if n <= 3 and c.get("variant") != "list":
        # three or fewer: columns across the full width, a display numeral heads each point
        cw = (W - 2 * M - (n - 1) * G * 2) / n
        y0 += 0.2
        tz, sz = _grid_sizes(s, pts, "lead", "text", cw, BOTTOM - y0, 56 / 72 * 1.1 + 0.1, 26, 20)
        for i, p in enumerate(pts):
            _point(s, M + i * (cw + G * 2), y0, cw, BOTTOM - y0, i, p, tz, sz, num_pt=56)
        s.chrome(c.get("section"))
        return
    two = n > 3
    per = (n + 1) // 2 if two else n
    cw = (W - 2 * M - G * 3) / 2 if two else span(9)
    y0 += 0.1
    rh = (BOTTOM - y0) / per
    tmax, bmax = (24, 19) if per <= 2 else (21, 17)
    stack = False
    if per == 2:  # numeral above the lead when the cells have the height for it at a comfortable size
        tz, sz = _grid_sizes(s, pts, "lead", "text", cw, rh - 0.15, 30 / 72 * 1.1 + 0.1, tmax, bmax)
        stack = sz >= 16 and tz >= 20
    if not stack:
        tz, sz = _grid_sizes(s, pts, "lead", "text", cw - 0.75, rh - 0.15, 0, tmax, bmax)
    for i, p in enumerate(pts):
        k, r = divmod(i, per)
        _point(s, M + k * (cw + G * 3), y0 + r * rh, cw, rh - 0.15, i, p, tz, sz, stack=stack, num_pt=30)
    s.chrome(c.get("section"))


def a_faq(s: S, c):
    """Questions and answers on hairlines, set as large as the space allows. 3 or fewer: columns; 5: title cell +
    five cells; 4 and 6: two columns."""
    qa = c["items"][:6]
    n = len(qa)
    if n == 5:
        tc, cells = _five_cells()
        tz, sz = _grid_sizes(s, qa, "q", "a", tc[2], cells[0][3], 0, 22, 18)
        _title_cell(s, c, *tc)
    else:
        y0 = _head(s, c, width=10, tight=True)
        if n <= 3:
            cw = (W - 2 * M - (n - 1) * G * 2) / n
            cells = [(M + i * (cw + G * 2), y0, cw, BOTTOM - y0) for i in range(n)]
        else:
            per = (n + 1) // 2
            cw = (W - 2 * M - G * 3) / 2
            rh = (BOTTOM - y0) / per
            cells = [(M + (i // per) * (cw + G * 3), y0 + (i % per) * rh, cw, rh - 0.12) for i in range(n)]
        tz, sz = _grid_sizes(s, qa, "q", "a", cw, cells[0][3], 0, 22 if n > 3 else 26, 18 if n > 3 else 20)
    for (x, y, w, h), it in zip(cells, qa):
        s.hline(x, y, w, s.line, 1)
        s.hline(x, y - 0.012, 0.45, s.accent, 2.5)
        qh = _lines_h(s, it["q"], w, tz, "h3")
        s.text(x, y + 0.22, w, qh + 0.06, it["q"], "h3", max_pt=tz, min_pt=tz)
        s.text(x, y + 0.22 + qh + 0.14, w, max(0.3, y + h - (y + 0.22 + qh + 0.14)), it["a"], "body", color=s.body,
               max_pt=sz, min_pt=sz, ls=1.4)
    s.chrome(c.get("section"))


# ------------------------------------------------------------------ bold layouts for the remaining archetypes
def v_process_stairs(s: S, c):
    """Steps as rising blocks: each step a column that is taller than the last; the final step in accent."""
    steps = c["steps"][:5]
    n = len(steps)
    s.label(col(0), TOP, span(8), c.get("kicker"))
    s.text(col(0), TOP + 0.32, span(8), 1.1, c["title"], "h1", max_pt=44)
    top_min, base = TOP + 1.7, BOTTOM - 0.08  # blocks stand on the footer line, never under it
    cw = (W - 2 * M - (n - 1) * 0.08) / n
    for i, st in enumerate(steps):
        x = M + i * (cw + 0.08)
        h = (base - top_min) * (0.68 + 0.32 * i / max(1, n - 1))  # first block still tall enough for its text
        y = base - h
        last = i == n - 1
        fill = s.accent if last else s.d.mix(s.bg, s.fg, 0.06 + 0.06 * i)
        ink = s.d.c["accent_text"] if last else s.fg
        # body text on each tile: muted only when it still reads (4.5:1) on that tile, else full-strength text
        body = ink if last else s.muted
        for _ in range(12):
            if contrast_ratio(body, fill) >= 4.5:
                break
            body = s.d.mix(body, ink, 0.2)
        num_col = ink  # one accent per set: only the final block carries it
        if contrast_ratio(num_col, fill) < 3:
            num_col = ink
        s.rect(x, y, cw, h + 0.05, fill, radius=s.d.radius * 0.6)
        s.text(x + 0.25, y + 0.2, cw - 0.5, 0.95, st.get("n", f"{i + 1:02d}"), "number", color=num_col, max_pt=58, min_pt=36)
        s.text(x + 0.25, y + 1.15, cw - 0.5, 0.6, st["title"], "h3", color=ink)
        s.text(x + 0.25, y + 1.75, cw - 0.5, min(1.8, h - 2.0), st.get("text"), "small" if n > 4 else "body",
               color=body, max_pt=15)
    s.chrome(c.get("section"))  # running elements stay put; they sit on pills where the blocks reach them


def v_timeline_band(s: S, c):
    """A full-width accent band crosses the slide; milestones sit on it, dates above, details below."""
    ev = c["events"][:6]
    n = len(ev)
    y0 = _head(s, c) + 0.35
    by, bh = y0 + 1.05, 1.05
    s.rect(0, by, W, bh, s.accent)
    ink = s.d.c["accent_text"]
    step = (W - 2 * M) / n
    tz = s.common([e["title"] for e in ev], step - G - 0.3, bh, "h3", max_pt=19)
    for i, e in enumerate(ev):
        x = M + i * step
        hl = i in c.get("highlight", [])
        s.text(x, y0 + 0.05, step - G, 0.9, e.get("when"), "h2", color=s.accent if hl else s.fg, max_pt=30, anchor="b")
        s.rect(x, by + bh / 2 - 0.09, 0.18, 0.18, ink, shape=MSO_SHAPE.OVAL)
        if i:
            s.vline(x - G / 2, by + 0.18, bh - 0.36, s.d.mix(s.accent, ink, 0.35))
        s.text(x + 0.3, by, step - G - 0.3, bh, e["title"], "h3", color=ink, anchor="m", max_pt=tz, min_pt=tz)
        s.text(x, by + bh + 0.3, step - G, BOTTOM - by - bh - 0.35, e.get("text"), "lead", color=s.muted, max_pt=19)
    s.chrome(c.get("section"))


def v_agenda_split(s: S, c):
    """Half the slide is an accent field with the title and item count; the list sits on the other half."""
    items = [it if isinstance(it, dict) else {"title": it} for it in c["items"][:7]]
    fw = W * 0.4
    s.rect(0, 0, fw, H, s.accent)
    s.fields.append((0, fw, s.d.c["accent_text"]))
    ink = s.d.c["accent_text"]
    s.text(M, 1.0, fw - 2 * M, 0.3, c.get("kicker"), "label", color=ink, caps=True, tracking=1.2 if s.d.caps_labels else 0)
    s.text(M, 1.35, fw - 2 * M, 2.6, c["title"], "h1", color=ink, max_pt=52)
    s.text(M, 4.1, fw - 2 * M, 2.0, f"{len(items):02d}", "number", color=s.d.mix(s.accent, ink, 0.45), max_pt=150, min_pt=90)
    s.text(M, BOTTOM - 0.6, fw - 2 * M, 0.6, c.get("text"), "small", color=ink)
    x, w = fw + 0.7, W - fw - 0.7 - M
    rh = (BOTTOM - TOP + 0.2) / len(items)
    for i, it in enumerate(items):
        y = TOP - 0.2 + i * rh
        s.hline(x, y, w)
        s.text(x, y, 0.9, rh, f"{i + 1:02d}", "h2", color=s.accent, anchor="m", max_pt=28)
        s.text(x + 1.0, y + (0.08 if it.get("text") else 0), w - 1.0 - (1.4 if it.get("meta") else 0),
               rh * (0.5 if it.get("text") else 1), it["title"], "h3", anchor="m" if not it.get("text") else "b", max_pt=22)
        if it.get("text"):
            s.text(x + 1.0, y + rh * 0.52, w - 1.0, rh * 0.45, it["text"], "small", color=s.muted)
        if it.get("meta"):
            s.text(x + w - 1.35, y, 1.35, rh, it["meta"], "small", color=s.muted, align="r", anchor="m")
    s.hline(x, TOP - 0.2 + len(items) * rh, w)
    s.chrome(c.get("section"))


def v_team_circles(s: S, c):
    """Round portraits with an offset accent ring, names centred under them."""
    y0 = _head(s, c)
    ppl = c["people"][:4]
    n = len(ppl)
    cw = (W - 2 * M - (n - 1) * G) / n
    d = min(cw - 0.5, BOTTOM - y0 - 1.25, 3.0)
    for i, p in enumerate(ppl):
        cx = M + i * (cw + G) + (cw - d) / 2
        ring = s.rect(cx + 0.14, y0 + 0.1, d, d, s.bg, shape=MSO_SHAPE.OVAL)
        ring.line.fill.solid(); ring.line.fill.fore_color.rgb = RGBColor.from_string(s.accent); ring.line.width = Pt(2.5)
        pic = s.image(cx, y0, d, d, p["photo"])
        if pic is not None:
            pic._element.find(qn("p:spPr")).find(qn("a:prstGeom")).set("prst", "ellipse")
        s.text(M + i * (cw + G), y0 + d + 0.3, cw, 0.42, p["name"], "h3", align="c")
        s.text(M + i * (cw + G), y0 + d + 0.72, cw, 0.4, p.get("role"), "small", color=s.muted, align="c")
    s.chrome(c.get("section"))


def v_gallery_mosaic(s: S, c):
    """One big photo bleeding off the left, two stacked on the right; captions as small labels on the photos."""
    imgs = [im if isinstance(im, dict) else {"path": im} for im in c["images"][:3]]
    s.label(col(0), TOP, span(5), c.get("kicker"))
    s.text(col(0), TOP + 0.32, span(5), 1.1, c["title"], "h1", max_pt=40)
    y0 = TOP + 1.55
    if any((im.get("fit") or c.get("fit")) == "contain" for im in imgs):
        _mosaic_contain(s, c, imgs, y0)
        return
    bw = W * 0.6
    s.image(0, y0, bw, H - y0, imgs[0])
    s.photos.append((0, bw))
    rest = imgs[1:]
    if rest:
        rh = (H - y0 - (len(rest) - 1) * 0.1) / len(rest)
        for k, im in enumerate(rest):
            s.image(bw + 0.1, y0 + k * (rh + 0.1), W - bw - 0.1, rh, im)
            s.photos.append((bw + 0.1, W - bw - 0.1))
    else:
        s.text(bw + 0.5, y0 + 0.2, W - bw - 0.5 - M, 2.5, c.get("text"), "lead", color=s.muted)
    for k, im in enumerate(imgs):
        if not im.get("caption"):
            continue
        if k == 0:
            x, y, w = 0.35, y0 + 0.25, bw - 0.7  # top of the photo: clear of the footer, same baseline as the others
        else:
            rh = (H - y0 - (len(rest) - 1) * 0.1) / len(rest)
            x, y, w = bw + 0.45, y0 + (k - 1) * (rh + 0.1) + 0.25, W - bw - 0.8
        tw = min(w, text_width(im["caption"], s.d.fonts["body"], True, 11) / 72 * 1.12 + 0.6)
        s.rect(x, y, tw, 0.38, s.d.c["bg"], radius=0.19)
        s.text(x + 0.25, y, tw - 0.4, 0.38, im["caption"], "small", anchor="m", max_pt=11, min_pt=11, bold=True,
               color=s.d.c["text"], warn=False)
    pass  # running header/footer stay (pills over photos)
    s.chrome(c.get("section"))


def _mosaic_contain(s: S, c, imgs, y0):
    """Mosaic for screenshots / slides / UI: nothing is cropped, so the cells take each image's own shape.
    One large image and up to two stacked beside it, sized together to fill the content area inside the margins,
    each with a hairline edge; captions sit under the images as quiet labels."""
    ars = []
    for im in imgs:
        p = s.d.path(im["path"])
        try:
            with Image.open(p) as f:
                ars.append(f.width / f.height)
        except Exception:
            ars.append(16 / 9)
    cap = 0.42 if any(im.get("caption") for im in imgs) else 0
    # the header text (if any) goes beside the title, opposite it
    s.text(col(6), TOP + 0.36, span(6), 1.1, c.get("text"), "body", color=s.body, max_pt=17)
    top, bot = y0 + 0.05, BOTTOM - 0.05
    avail_w, avail_h = W - 2 * M, bot - top
    rest = ars[1:]
    gap = G * 1.5
    if not rest:
        x, y, w, h = contain_box(M, top, avail_w, avail_h - cap, ars[0], "t")
        boxes = [(x, y, w, h)]
    else:
        # big image height Hb; the stack beside it shares the same total height: sum(hi) + gaps (+captions) = Hb
        # stack width ws: hi = ws / ar_i  ->  ws = (Hb - extra) / sum(1/ar_i)
        k = len(rest)
        extra = (k - 1) * gap + (k - 1) * cap
        inv = sum(1 / a for a in rest)
        def widths(hb):
            return hb * ars[0], max(0.5, (hb - extra) / inv)
        hb = avail_h - cap
        wb, ws = widths(hb)
        if wb + gap + ws > avail_w:   # too wide: scale the whole group down
            lo, hi_ = 0.5, hb
            for _ in range(30):
                mid = (lo + hi_) / 2
                a, b = widths(mid)
                lo, hi_ = (mid, hi_) if a + gap + b <= avail_w else (lo, mid)
            hb = lo
            wb, ws = widths(hb)
        x0 = M   # start-aligned under the title (mirrored in RTL); any slack is left at the far end
        boxes = [(x0, top, wb, hb)]
        yy = top
        for a in rest:
            hs = ws / a
            boxes.append((x0 + wb + gap, yy, ws, hs))
            yy += hs + gap + cap
    for (x, y, w, h), im in zip(boxes, imgs):
        s.image(x, y, w, h, dict(im, fit="contain"), radius=s.d.radius)
        if im.get("caption"):
            s.hline(x, y + h + 0.14, 0.3, s.accent, 1.5)
            s.text(x + 0.42, y + h + 0.04, w - 0.42, 0.3, im["caption"], "small", color=s.body, max_pt=12, min_pt=10,
                   anchor="m", bold=True, warn=False)
    s.chrome(c.get("section"))


def v_quote_bleed(s: S, c):
    """Photo bleeds over one half; the quote is set large on the other with a giant accent quote mark."""
    img = c.get("image")
    iw = W * 0.42
    if img:
        s.image(W - iw, 0, iw, H, img)
    tw = (W - iw if img else W) - M - 0.7
    s.text(col(0), 0.55, 1.6, 1.6, "“", "display", color=s.accent, max_pt=200, min_pt=160, font=s.d.fonts["heading"], warn=False)
    s.text(col(0), 2.0, tw, 3.2, c["quote"], "h1", max_pt=c.get("max_pt", 44), ls=1.08)
    s.hline(col(0), 5.45, 0.9, s.accent, 3)
    s.text(col(0), 5.65, tw, 0.4, c.get("name"), "h3")
    s.text(col(0), 6.05, tw, 0.4, c.get("role"), "small", color=s.muted)
    s.chrome(c.get("section"))


# ------------------------------------------------------------------ business archetypes (structured content)
NUM_CELL = re.compile(r"^\s*[-+−–]?\s*[₪€$£]?\s*[\d.,]+\s*(%|[KMBkmb]|x|×|pp|pt|h|d|ms)?\s*[₪€$£]?\s*[↑↓▲▼]?\s*$")


def _num(v):
    """Leading figure of a value as a float ("12,400" -> 12400, "3.2K" -> 3200), or None."""
    m_ = re.search(r"([\d][\d,]*\.?\d*)\s*([KMBkmb])?", str(v))
    if not m_:
        return None
    try:
        x = float(m_.group(1).replace(",", ""))
    except ValueError:
        return None
    return x * {"k": 1e3, "m": 1e6, "b": 1e9}.get((m_.group(2) or "").lower(), 1)


def _ink_on(s: S, fill):
    """Text colour that reads on a fill: the slide's ground or its text colour, whichever contrasts more."""
    return s.bg if contrast_ratio(s.bg, fill) >= contrast_ratio(s.fg, fill) else s.fg


def _he(s: S, en, he):
    return he if s.d.rtl else en


def _initials(name):
    parts = [p_ for p_ in re.split(r"\s+", str(name).strip()) if p_]
    return "".join(p_[0] for p_ in parts[:2]).upper()


def _avatar(s: S, x, y, d, person, fill=None, ink=None, ring=None):
    """Round portrait (photo cropped to a circle) or initials on a disc."""
    photo = person.get("photo") if isinstance(person, dict) else None
    if photo:
        pic = s.image(x, y, d, d, photo, flip=False)
        if pic is not None:
            pic._element.find(qn("p:spPr")).find(qn("a:prstGeom")).set("prst", "ellipse")
            if ring:
                pic.line.color.rgb = RGBColor.from_string(ring)
                pic.line.width = Pt(1.5)
        return pic
    fill = fill or s.d.mix(s.bg, s.fg, 0.12)
    disc = s.rect(x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    if ring:
        disc.line.fill.solid(); disc.line.fill.fore_color.rgb = RGBColor.from_string(ring); disc.line.width = Pt(1.5)
    s.text(x, y, d, d, _initials(person.get("name", "") if isinstance(person, dict) else person), "h3",
           color=ink or _ink_on(s, fill), align="c", anchor="m", max_pt=d * 72 * 0.36, min_pt=8, warn=False)
    return disc


def _arrow(s: S, x1, y1, x2, y2, color, weight=1.25, head=True):
    """Straight connector in LTR coordinates (mirrored in RTL) with an optional arrowhead at the end."""
    from pptx.enum.shapes import MSO_CONNECTOR
    if s.d.rtl:
        x1, x2 = W - x1, W - x2
    cn = s.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    s._clean(cn)
    cn.line.color.rgb = RGBColor.from_string(color)
    cn.line.width = Pt(weight)
    if head:
        ln = cn.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return cn


# ---- table
def _cell_borders(cell, sides):
    """sides: {"L"|"R"|"T"|"B": (hex, pt) | None}. Unlisted sides get no line."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for i, side in enumerate("LRTB"):
        spec = sides.get(side)
        ln = etree.Element(qn(f"a:ln{side}"), w=str(int((spec[1] if spec else 0) * 12700)), cmpd="sng")
        if spec:
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=spec[0])
        else:
            etree.SubElement(ln, qn("a:noFill"))
        tcPr.insert(i, ln)


def a_table(s: S, c):
    """Native, editable PowerPoint table. Numbers align on their last digit; Hebrew tables run right to left.
    bold: header on the inverse band, a highlighted column becomes an accent band; editorial: hairlines only;
    soft: a rounded card with quiet zebra rows."""
    cols_ = [str(h) for h in c["columns"][:7]]
    nc = len(cols_)
    rows = [[str(v) if v is not None else "" for v in r][:nc] for r in c["rows"][:8]]
    rows = [r + [""] * (nc - len(r)) for r in rows]
    nr = len(rows)
    art = s.d.art
    hc, hr = c.get("highlight_col"), c.get("highlight_row")
    total = c.get("total", False)
    numeric = [k > 0 and sum(1 for r in rows if NUM_CELL.match(r[k]) or r[k] in ("", "–", "-", "—")) == nr
               for k in range(nc)]
    side = nc <= 4 and not c.get("wide")
    if side:   # narrow table: the argument in a column beside it (asymmetric, like the chart slide)
        s.label(col(0), TOP, span(4), c.get("kicker"))
        s.text(col(0), TOP + 0.32, span(4), 2.3, c["title"], "h1", max_pt=c.get("title_pt", 42))
        y = TOP + 0.32 + s.last_h + 0.3
        s.text(col(0), y, span(4) - 0.2, BOTTOM - y - 1.8, c.get("text"), "body", color=s.body, max_pt=17)
        if c.get("takeaway"):
            s.text(col(0), BOTTOM - 1.55, span(4), 1.0, c["takeaway"], "number", color=s.accent, max_pt=64, min_pt=32)
            s.text(col(0), BOTTOM - 0.55, span(4), 0.45, c.get("takeaway_label"), "small", color=s.muted)
        tx, tw, ty, bot = col(5), span(7), TOP + 0.05, BOTTOM
    else:
        ty = _head(s, c, width=10, tight=True)
        tx, tw, bot = M, W - 2 * M, BOTTOM - (0.4 if c.get("source") else 0)
    pad = 0.14 if art != "soft" else 0.2
    inset = 0.22 if art == "soft" else 0
    tx_, tw_ = tx + inset, tw - 2 * inset
    fb = s.d.fonts["body"]
    # column widths from measured content: numbers take what they need, the label column takes the rest
    nat = []
    for k in range(nc):
        cells = [r[k] for r in rows]
        wpt = max([text_width(t, fb, k == 0 or (total and i == nr - 1), 15) for i, t in enumerate(cells)] +
                  [text_width(cols_[k].upper() if s.d.caps_labels else cols_[k], fb, True, 10) * 1.1])
        nat.append(min(wpt / 72, 3.6 if k == 0 else 2.6) + 2 * pad)
    extra = tw_ - sum(nat)
    if extra > 0:
        widths = [nat[0] + extra * 0.45] + [w_ + extra * 0.55 / max(1, nc - 1) for w_ in nat[1:]] if nc > 1 else [tw_]
    else:
        widths = [w_ * tw_ / sum(nat) for w_ in nat]
    head_h = 0.5
    room = bot - ty - head_h - 2 * inset - (0.1 if art == "soft" else 0)
    rh = max(0.36, min(0.66, room / nr))
    sz = min(16, min(s.fit_size(r[k], widths[k] - 2 * pad, rh - 0.06, "body", max_pt=16, min_pt=10.5,
                                 bold=k == 0) for r in rows for k in range(nc) if r[k]))
    hz = min(10.5, min(s.fit_size(h, widths[k] - 2 * pad, head_h - 0.08, "label", max_pt=10.5, min_pt=8)
                       for k, h in enumerate(cols_) if h))
    th = head_h + nr * rh
    if art == "soft":
        card = s.rect(tx, ty, tw, th + 2 * inset + 0.05, s.surface, radius=s.d.radius)
        s.shadow(card, alpha=0.07)
    gf = s.slide.shapes.add_table(nr + 1, nc, E(s.X(tx_, tw_)), E(ty + inset), E(tw_), E(th))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for a_ in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(a_, "0")
    if s.d.rtl:
        tblPr.set("rtl", "1")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # "No Style, No Grid": every line below is ours
    for k, w_ in enumerate(widths):
        tbl.columns[k].width = E(w_)
    tbl.rows[0].height = E(head_h)
    for r in range(nr):
        tbl.rows[r + 1].height = E(rh)
    ground = s.surface if art == "soft" else s.bg
    acc, acc_ink = s.d.c["accent"], s.d.c["accent_text"]
    hair = (s.d.mix(ground, s.fg, 0.16), 0.75)
    strong = (s.fg, 1.5)
    tint = s.d.mix(ground, acc, 0.13 if s.bg_is_light() else 0.2)

    def put(cell, txt, size, bold, color, fill, align, family, borders, caps=False):
        cell.margin_left = cell.margin_right = E(pad)
        cell.margin_top = cell.margin_bottom = E(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if fill:
            cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            cell.fill.background()
        _cell_borders(cell, borders)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        # Hebrew tables are right-aligned throughout (the table itself runs right to left); in English, figures
        # align right so their last digits line up
        al = "r" if s.d.rtl else align
        p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[al]
        if caps and s.d.caps_labels:
            txt = txt.upper()
        s.runs(p, txt, size, bold, family, color, tracking=0.8 if caps and s.d.caps_labels else 0)
        if not txt:
            p.add_run().text = ""

    for k in range(nc):   # header
        on = hc == k
        if art == "bold":
            fill = acc if on else s.fg
            ink = acc_ink if on else s.bg
            b = {}
        elif art == "editorial":
            fill, ink = None, (s.ink if on else s.muted)
            b = {"B": strong, **({"T": (acc, 2.5)} if on else {})}
        else:
            fill, ink = (tint if on else None), (s.ink if on else s.muted)
            b = {"B": hair}
        put(tbl.cell(0, k), cols_[k], hz, True, ink, fill, "r" if numeric[k] else "l", fb, b, caps=True)
    for r, row in enumerate(rows):
        last = r == nr - 1
        is_total = total and last
        on_r = hr == r
        for k, v in enumerate(row):
            on_c = hc == k
            fill, ink, bold = None, s.fg, k == 0 or is_total
            b = {} if last and not is_total else {"B": hair}
            if art == "soft" and r % 2 == 1:
                fill = s.d.mix(ground, s.bg, 0.6)
                b = {}
            if art == "soft" and not (r % 2 == 1):
                b = {}
            if on_c:
                if art == "bold":
                    fill, bold = tint, True
                elif art == "editorial":
                    bold = True
                else:
                    fill = tint
            if on_r:
                if art == "bold":
                    fill, ink, bold = acc, acc_ink, True
                else:
                    fill, bold = (None if art == "editorial" else tint), True
                    b = {"T": (acc, 1.5), "B": (acc, 1.5)} if art == "editorial" else b
                    if numeric[k]:
                        ink = s.ink
            if is_total:
                b = {"T": strong}
            if not on_r and not k == 0 and not numeric[k] and not is_total:
                ink = s.body
            put(tbl.cell(r + 1, k), v, sz, bold, ink, fill, "r" if numeric[k] else "l",
                fb, b)
    if c.get("source"):
        s.text(tx, ty + inset * 2 + th + 0.18, tw, 0.3, c["source"], "small", color=s.muted, max_pt=11)
    s.chrome(c.get("section"))


# ---- funnel
def a_funnel(s: S, c):
    """Stages narrowing to the result. bold: a continuous funnel silhouette (trapezoids, darkest at the top, the
    result in accent); editorial: outlined tiers; soft: rounded bars. Values sit inside, labels on leaders beside,
    step conversion between the rows."""
    st = c["stages"][:6]
    n = len(st)
    art = s.d.art
    y0 = _head(s, c, width=8, tight=True) - 0.1
    fx, fw = M, span(6) + 0.1
    lx, lw = col(7) + 0.2, span(5) - 0.2
    gap = 0.1
    th = min(1.0, (BOTTOM - y0 - (n - 1) * gap) / n)
    y0 += max(0.0, (BOTTOM - y0 - n * th - (n - 1) * gap) * 0.25)
    vals = [_num(x.get("value")) for x in st]
    if all(v is not None and v > 0 for v in vals) and vals[0]:
        ratios = [max(0.24, (v / vals[0]) ** 0.38) for v in vals]
        for i in range(1, n):          # a funnel never widens
            ratios[i] = min(ratios[i], ratios[i - 1] - 0.04)
    else:
        ratios = [1 - 0.72 * i / max(1, n - 1) for i in range(n)]
    ratios = [max(0.2, r) for r in ratios]
    hl = c.get("highlight", [n - 1])
    cx = fx + fw / 2
    vz = min(s.fit_size(x.get("value", ""), max(0.6, fw * ratios[min(i + 1, n - 1)] * 0.85 - 0.3), th * 0.8, "number",
                        max_pt=min(52, th * 72 * 0.62), min_pt=16) for i, x in enumerate(st))
    tz = s.common([x["label"] for x in st], lw, th * 0.48, "h3", max_pt=22, min_pt=13)
    bz = s.common([x.get("text") for x in st], lw, th * 0.52, "small", max_pt=13.5, min_pt=10)
    for i, x in enumerate(st):
        y = y0 + i * (th + gap)
        wt = fw * ratios[i]
        wb = fw * (ratios[i + 1] if i < n - 1 else ratios[i] * 0.84)
        on = i in hl
        if art == "editorial":
            fill = s.d.c["accent"] if on else None
        elif art == "soft":
            fill = s.d.c["accent"] if on else s.d.mix(s.bg, s.fg, 0.1 + 0.07 * i)
        else:
            fill = s.d.c["accent"] if on else s.d.mix(s.bg, s.fg, max(0.28, 0.92 - 0.16 * i))
        ink = s.d.c["accent_text"] if on else (s.fg if fill is None else _ink_on(s, fill))
        if art == "soft":
            sh = s.rect(cx - wt / 2, y, wt, th, fill, radius=min(th / 2, s.d.radius * 1.5))
        else:
            sh = s.rect(cx - wt / 2, y, wt, th, fill or s.bg, shape=MSO_SHAPE.TRAPEZOID)
            sh.adjustments[0] = max(0.0, (wt - wb) / 2 / min(wt, th))
            sh._element.find(qn("p:spPr")).find(qn("a:xfrm")).set("flipV", "1")  # wide edge on top
            if fill is None:
                sh.fill.background()
                sh.line.fill.solid(); sh.line.fill.fore_color.rgb = RGBColor.from_string(s.fg); sh.line.width = Pt(1)
        bw = min(wt, wb) if art != "soft" else wt
        s.text(cx - bw / 2 + 0.15, y, bw - 0.3, th, x.get("value"), "number", color=ink, max_pt=vz, min_pt=vz,
               align="c", anchor="m")
        # leader from the tier's edge (at mid height) to its label
        ex = cx + ((wt + wb) / 4 if art != "soft" else wt / 2) + 0.14
        s.hline(ex, y + th / 2, lx - 0.22 - ex, s.accent if on else s.line, 1.25 if on else 0.75)
        s.rect(lx - 0.27, y + th / 2 - 0.045, 0.09, 0.09, s.accent if on else s.fg, shape=MSO_SHAPE.OVAL)
        has_t = bool(x.get("text"))
        s.text(lx, y + (0.02 if has_t else 0), lw, th * (0.5 if has_t else 1), x["label"], "h3",
               color=s.ink if on else s.fg, max_pt=tz, min_pt=tz, anchor="b" if has_t else "m")
        if has_t:
            s.text(lx, y + th * 0.52 + 0.02, lw, th * 0.48, x["text"], "small", color=s.body, max_pt=bz, min_pt=bz)
        # step conversion, set between this stage and the next, on the funnel's axis edge
        if i < n - 1 and c.get("conversion", True):
            conv = st[i + 1].get("conversion")
            if conv is None and vals[i] and vals[i + 1] is not None and all(v is not None for v in vals):
                conv = f"{vals[i + 1] / vals[i] * 100:.0f}%" if vals[i + 1] / vals[i] >= 0.1 else f"{vals[i + 1] / vals[i] * 100:.1f}%"
            if conv:
                yy = y + th + gap / 2 - 0.13
                ex2 = cx + (wb / 2 if art != "soft" else min(wt, fw * ratios[i + 1]) / 2) + 0.2
                s.text(ex2, yy, 1.4, 0.26, f"↓ {conv}", "label", color=s.muted, max_pt=10.5, min_pt=9, anchor="m",
                       warn=False)
    s.chrome(c.get("section"))


# ---- matrix / SWOT
def a_matrix(s: S, c):
    """2x2 matrix: the argument on the side, four quadrants with axes. The one quadrant that matters takes the
    accent. variant swot: a full-width 2x2 with giant S / W / O / T letters."""
    if c.get("variant") == "swot":
        return _swot(s, c)
    q = c["quadrants"][:4]
    art = s.d.art
    s.label(col(0), TOP, span(4), c.get("kicker"))
    s.text(col(0), TOP + 0.32, span(4) - 0.1, 2.4, c["title"], "h1", max_pt=c.get("title_pt", 42))
    y = TOP + 0.32 + s.last_h + 0.3
    s.text(col(0), y, span(4) - 0.3, BOTTOM - y, c.get("text"), "body", color=s.body, max_pt=17)
    mx0, my0 = col(4) + 0.75, TOP + 0.05
    mw, mh = W - M - mx0, BOTTOM - 0.55 - my0
    gap = 0.1 if art != "editorial" else 0
    qw, qh = (mw - gap) / 2, (mh - gap) / 2
    hl = c.get("highlight")
    pad = 0.3
    tz = s.common([x["title"] for x in q], qw - 2 * pad, 0.8, "h2", max_pt=24, min_pt=15)
    bodies = [x.get("text") or "\n".join(x.get("items", [])[:4]) for x in q]
    bz = s.common(bodies, qw - 2 * pad, qh - 1.15, "body", max_pt=16, min_pt=11)
    for i, x in enumerate(q):
        r, k = divmod(i, 2)
        qx, qy = mx0 + k * (qw + gap), my0 + r * (qh + gap)
        on = hl == i
        if art == "editorial":
            fill = s.d.mix(s.bg, s.d.c["accent"], 0.1) if on else None
        else:
            fill = s.d.c["accent"] if on else s.surface
        if fill:
            s.rect(qx, qy, qw, qh, fill, radius=s.d.radius if art == "soft" else s.d.radius * 0.4)
        ink = s.d.c["accent_text"] if (on and art != "editorial") else s.fg
        sub = s.readable(s.d.mix(fill or s.bg, ink, 0.75), fill or s.bg, 4.5, ink)
        tag = x.get("tag") or ["01", "02", "03", "04"][i]
        s.text(qx + pad, qy + 0.26, qw - 2 * pad, 0.28, tag, "label", color=ink if on and art != "editorial" else s.ink,
               caps=True, tracking=1 if s.d.caps_labels else 0)
        s.text(qx + pad, qy + 0.56, qw - 2 * pad, 0.8, x["title"], "h2", color=ink, max_pt=tz, min_pt=tz)
        s.text(qx + pad, qy + 0.62 + s.last_h + 0.12, qw - 2 * pad, qh - 1.15, bodies[i], "body", color=sub,
               max_pt=bz, min_pt=bz)
    if art == "editorial":   # the cross: two rules through the centre
        s.hline(mx0, my0 + mh / 2, mw, s.fg, 1.25)
        s.vline(mx0 + mw / 2, my0, mh, s.fg, 1.25)
    # axes: arrows along the outer edges, the label at the far (high) end
    ax = mx0 - 0.28
    _arrow(s, ax, my0 + mh, ax, my0 + 0.05, s.fg, 1.25)
    _arrow(s, mx0, my0 + mh + 0.28, mx0 + mw - 0.05, my0 + mh + 0.28, s.fg, 1.25)
    xl, yl = c.get("x_axis"), c.get("y_axis")
    if xl:
        s.text(mx0, my0 + mh + 0.38, mw - 0.05, 0.3, xl, "label", color=s.fg, align="r", caps=True,
               tracking=1 if s.d.caps_labels else 0, warn=False)
    if yl:   # rotated along the vertical axis, reading upward
        tb = s.text(ax - 0.3 - mh / 2, my0 + mh / 2 - 0.15, mh, 0.3, yl, "label", color=s.fg,
                    align="l" if s.d.rtl else "r", caps=True, tracking=1 if s.d.caps_labels else 0, warn=False)
        if tb is not None:
            tb.rotation = 270
    s.chrome(c.get("section"))


def _swot(s: S, c):
    q = c["quadrants"][:4]
    art = s.d.art
    letters = ["S", "W", "O", "T"]
    names = _he(s, ["Strengths", "Weaknesses", "Opportunities", "Threats"],
                ["חוזקות", "חולשות", "הזדמנויות", "איומים"])
    y0 = _head(s, c, width=10, tight=True) - 0.05
    gap = 0.1 if art != "editorial" else 0.3
    qw, qh = (W - 2 * M - gap) / 2, (BOTTOM - y0 - gap) / 2
    hl = c.get("highlight")
    lw = min(1.5, qh * 0.95)
    pad = 0.28
    bodies = [x.get("text") or "\n".join("· " + t for t in x.get("items", [])[:4]) for x in q]
    bz = s.common(bodies, qw - lw - 2 * pad, qh - 0.85, "body", max_pt=16, min_pt=11)
    for i, x in enumerate(q):
        r, k = divmod(i, 2)
        qx, qy = M + k * (qw + gap), y0 + r * (qh + gap)
        on = hl == i
        if art == "editorial":
            fill = None
            s.hline(qx, qy, qw, s.accent if on else s.fg, 2.5 if on else 1)
        else:
            fill = s.d.c["accent"] if on else s.surface
            s.rect(qx, qy, qw, qh, fill, radius=s.d.radius if art == "soft" else s.d.radius * 0.4)
        ink = s.d.c["accent_text"] if (on and fill) else s.fg
        sub = s.readable(s.d.mix(fill or s.bg, ink, 0.78), fill or s.bg, 4.5, ink)
        # the giant initial: outlined on the quiet quadrants, solid on the one that matters
        lc = ink if on and fill else (s.accent if on else s.d.mix(fill or s.bg, s.fg, 0.4))
        s.text(qx + 0.12, qy + 0.02, lw, qh - 0.04, letters[i], "number", color=lc, max_pt=min(150, qh * 72 * 0.9),
               min_pt=60, anchor="t", outline=None if on else 1.25, align="c", warn=False)
        tx, tw = qx + lw + pad, qw - lw - 2 * pad
        s.text(tx, qy + pad, tw, 0.4, x.get("title") or names[i], "h3", color=ink, max_pt=20)
        s.text(tx, qy + pad + 0.5, tw, qh - pad - 0.6, bodies[i], "body", color=sub, max_pt=bz, min_pt=bz)
    s.chrome(c.get("section"))


# ---- org chart
def a_org(s: S, c):
    """Org chart: the lead in an inverse card at the top, direct reports on a bus line, their teams on a spine
    under each. One node may carry the accent (`highlight`: [report] or [report, child])."""
    root, reps = c["root"], c["reports"][:4]
    n = len(reps)
    art = s.d.art
    hl = c.get("highlight") or []
    rw, rh = span(3) + 0.4, 1.0
    rx, ry = max((W - rw) / 2, col(4) + 0.3), TOP + 0.05
    kids_n = max([len(r.get("children", [])[:3]) for r in reps] + [0])
    cy = min(3.25, BOTTOM - 0.92 - 0.3 - kids_n * 0.64)       # the reports row: as low as the teams allow
    bus = cy - 0.36
    s.label(col(0), TOP, rx - M - 0.4, c.get("kicker"))
    s.text(col(0), TOP + 0.32, rx - M - 0.4, bus - TOP - 0.5, c["title"], "h1", max_pt=c.get("title_pt", 40))
    if c.get("text"):   # opposite the title, so the lead's card stands between the two
        tx_ = rx + rw + 0.45
        s.text(tx_, TOP + 0.08, W - M - tx_, bus - TOP - 0.3, c["text"], "body", color=s.body, max_pt=16)
    s.shadow(s.rect(rx, ry, rw, rh, s.fg, radius=s.d.radius))
    ad = 0.62
    _avatar(s, rx + 0.22, ry + (rh - ad) / 2, ad, root, fill=s.d.mix(s.fg, s.bg, 0.22), ink=s.bg)
    s.text(rx + 0.22 + ad + 0.18, ry + 0.16, rw - ad - 0.6, 0.42, root["name"], "h3", color=s.bg, max_pt=20, anchor="b")
    s.text(rx + 0.22 + ad + 0.18, ry + 0.58, rw - ad - 0.6, 0.3, root.get("role"), "small",
           color=s.readable(s.d.mix(s.fg, s.bg, 0.7), s.fg), max_pt=12)
    cw = (W - 2 * M - (n - 1) * G) / n
    ch = 0.92
    centers = [M + i * (cw + G) + cw / 2 for i in range(n)]
    s.vline(rx + rw / 2, ry + rh, bus - ry - rh, s.fg, 1)
    if n > 1:
        s.hline(centers[0], bus, centers[-1] - centers[0], s.fg, 1)
    kids_all = [k for r in reps for k in r.get("children", [])[:3]]
    nz = s.common([r["name"] for r in reps], cw - 1.05, 0.38, "h3", max_pt=18, min_pt=12)
    kz = s.common([k["name"] for k in kids_all], cw - 0.75, 0.3, "body", max_pt=15, min_pt=10.5)
    rz = s.common([k.get("role") for k in kids_all], cw - 0.75, 0.26, "small", max_pt=12, min_pt=9)
    krow = min(0.66, (BOTTOM - (cy + ch) - 0.25) / 3)
    for i, r in enumerate(reps):
        x = M + i * (cw + G)
        on = hl[:1] == [i] and len(hl) == 1
        s.vline(centers[i], bus, cy - bus, s.fg, 1)
        if art == "editorial":
            fill = None
            s.hline(x, cy, cw, s.accent if on else s.fg, 2.5 if on else 1.25)
        else:
            fill = s.d.c["accent"] if on else s.surface
            card = s.rect(x, cy, cw, ch, fill, radius=s.d.radius)
            if art == "soft":
                s.shadow(card, alpha=0.07)
        ink = s.d.c["accent_text"] if on and fill else s.fg
        ad2 = 0.5
        _avatar(s, x + 0.2, cy + (ch - ad2) / 2, ad2, r, fill=s.d.mix(fill or s.bg, ink, 0.14), ink=ink)
        s.text(x + 0.85, cy + 0.12, cw - 1.05, 0.38, r["name"], "h3", color=ink, max_pt=nz, min_pt=nz, anchor="b")
        s.text(x + 0.85, cy + 0.5, cw - 1.05, 0.32, r.get("role"), "small",
               color=s.readable(s.d.mix(fill or s.bg, ink, 0.72), fill or s.bg, 4.5, ink), max_pt=12, min_pt=9)
        kids = r.get("children", [])[:3]
        if kids:
            sx = x + 0.2 + ad2 / 2
            top_k = cy + ch
            last_mid = cy + ch + 0.25 + (len(kids) - 1) * krow + krow * 0.42
            s.vline(sx, top_k, last_mid - top_k, s.line, 1)
            for j, k in enumerate(kids):
                ky = cy + ch + 0.25 + j * krow
                on_k = hl == [i, j]
                s.hline(sx, ky + krow * 0.42, 0.28, s.line, 1)
                s.rect(sx + 0.24, ky + krow * 0.42 - 0.04, 0.08, 0.08, s.accent if on_k else s.fg, shape=MSO_SHAPE.OVAL)
                s.text(sx + 0.45, ky + 0.02, cw - 0.75, krow * 0.5, k["name"], "body", bold=True,
                       color=s.ink if on_k else s.fg, max_pt=kz, min_pt=kz, anchor="b")
                s.text(sx + 0.45, ky + krow * 0.5 + 0.02, cw - 0.75, krow * 0.45, k.get("role"), "small",
                       color=s.muted, max_pt=rz, min_pt=rz)
    s.chrome(c.get("section"))


# ---- roadmap
def _period_index(periods, v, end=False):
    if isinstance(v, (int, float)):
        return float(v)
    if v in periods:
        return periods.index(v) + (1 if end else 0)
    return None


def a_roadmap(s: S, c):
    """Workstreams x periods. Bars span the periods they run in; work already done goes quiet, the item that
    matters takes the accent, a 'now' line cuts through. bold: solid bars; editorial: bars as heavy rules with the
    label above; soft: rounded pills."""
    periods = c["periods"][:6]
    lanes = c["lanes"][:4]
    npd, nl = len(periods), len(lanes)
    art = s.d.art
    y0 = _head(s, c, width=8, tight=True) - 0.1
    lw_ = span(2) + 0.25
    gx = M + lw_ + G
    gw = W - M - gx
    pw = gw / npd
    hy = y0
    ly0 = hy + 0.48
    lh = (BOTTOM - ly0) / nl
    now = c.get("now")
    if isinstance(now, str):
        now = _period_index(periods, now)
        now = None if now is None else now + 0.5
    for k, p in enumerate(periods):
        x = gx + k * pw
        past = now is not None and k + 1 <= now
        cur = now is not None and k <= now < k + 1
        s.text(x + 0.12, hy, pw - 0.24, 0.32, p, "label", color=s.ink if cur else (s.muted if past else s.fg),
               caps=True, tracking=1 if s.d.caps_labels else 0, anchor="m")
        s.vline(x, hy, BOTTOM - hy, s.line, 0.75)
    s.vline(gx + gw, hy, BOTTOM - hy, s.line, 0.75)
    if now is not None and 0 <= now <= npd:   # the 'now' line runs behind the bars; its pill sits on the header rule
        nx = gx + now * pw
        s.rect(nx - 0.012, ly0, 0.024, BOTTOM - ly0, s.d.c["accent"])
    items_all = [it["title"] for ln in lanes for it in ln.get("items", [])]
    lz = s.common([ln["name"] for ln in lanes], lw_, 0.45, "h3", max_pt=19, min_pt=12)
    hl = c.get("highlight")
    layout = []
    max_rows = 1
    for ln in lanes:
        rows_end, placed = [], []
        for it in ln.get("items", [])[:6]:
            a_ = _period_index(periods, it.get("from", 0))
            b_ = _period_index(periods, it.get("to", a_ + 1 if a_ is not None else 1), end=True)
            if a_ is None or b_ is None:
                s.d.warnings.append(f"slide {s.no}: roadmap item {it['title']!r} has an unknown period")
                continue
            for r_, e_ in enumerate(rows_end):
                if a_ >= e_ - 1e-6:
                    rows_end[r_] = b_
                    break
            else:
                rows_end.append(b_)
                r_ = len(rows_end) - 1
            placed.append((it, a_, b_, r_))
        max_rows = max(max_rows, len(rows_end))
        layout.append((ln, placed, max(1, len(rows_end))))
    bh = min(0.5, (lh - 0.3) / max_rows - 0.1)
    if art == "editorial":
        bh = min(0.55, bh)
    min_w = min([(b_ - a_) * pw for _, pl, _ in layout for _, a_, b_, _ in pl] or [pw])
    tz = s.common(items_all, max(0.5, min_w - 0.34), bh - 0.06 if art != "editorial" else 0.3, "body",
                  max_pt=15, min_pt=10)
    for li, (ln, placed, nrows) in enumerate(layout):
        y = ly0 + li * lh
        s.hline(M, y, W - 2 * M, s.fg if li == 0 else s.line, 1.25 if li == 0 else 0.75)
        s.text(M, y + 0.16, 0.5, 0.3, f"{li + 1:02d}", "label", color=s.ink, warn=False)
        s.text(M, y + 0.44, lw_ - 0.1, 0.45, ln["name"], "h3", max_pt=lz, min_pt=lz)
        if ln.get("text"):
            s.text(M, y + 0.44 + s.last_h + 0.06, lw_ - 0.1, lh - 0.9, ln["text"], "small", color=s.muted, max_pt=11.5)
        blk = nrows * (bh + 0.1) - 0.1
        by0 = y + max(0.18, (lh - blk) / 2)
        for it, a_, b_, r_ in placed:
            bx, bw_ = gx + a_ * pw + 0.06, (b_ - a_) * pw - 0.12
            by = by0 + r_ * (bh + 0.1)
            on = it.get("highlight") or (hl == [li, ln.get("items", []).index(it)])
            done = now is not None and b_ <= now
            if art == "editorial":
                colr = s.d.c["accent"] if on else (s.line if done else s.fg)
                s.rect(bx, by + bh - 0.07, bw_, 0.07, colr)
                s.text(bx, by, bw_, bh - 0.1, it["title"], "body", bold=not done,
                       color=s.ink if on else (s.muted if done else s.fg), max_pt=tz, min_pt=tz, anchor="b")
                continue
            if art == "soft":
                fill = s.d.c["accent"] if on else (s.surface if done else s.d.mix(s.bg, s.fg, 0.16))
            else:
                fill = s.d.c["accent"] if on else (s.surface if done else s.d.mix(s.bg, s.fg, 0.86))
            ink = s.d.c["accent_text"] if on else (s.muted if done else _ink_on(s, fill))
            if done and contrast_ratio(ink, fill) < 4.5:
                ink = s.readable(ink, fill)
            s.rect(bx, by, bw_, bh, fill, radius=bh / 2 if art == "soft" else s.d.radius * 0.35)
            s.text(bx + 0.17, by, bw_ - 0.34, bh, it["title"], "body", bold=True, color=ink, max_pt=tz, min_pt=tz,
                   anchor="m")
    if now is not None and 0 <= now <= npd:
        nx = gx + now * pw
        lab = c.get("now_label") or _he(s, "Now", "היום")
        tw_ = text_width(lab.upper() if s.d.caps_labels else lab, s.d.fonts["body"], True, 9.5) / 72 + 0.3
        s.rect(nx - tw_ / 2, ly0 - 0.14, tw_, 0.28, s.d.c["accent"], radius=0.14)
        s.text(nx - tw_ / 2, ly0 - 0.14, tw_, 0.28, lab, "label", color=s.d.c["accent_text"], align="c", anchor="m",
               max_pt=9.5, min_pt=9.5, caps=True, warn=False)
    s.hline(M, BOTTOM, W - 2 * M, s.line, 0.75)
    s.chrome(c.get("section"))


# ---- device mock-up
DEVICE_BODY = "1C1C1E"


def _media_aspect(s: S, c):
    v = s.sd.get("_video")
    p = v["poster"] if v else s.d.path(c.get("image"))
    try:
        with Image.open(p) as im:
            return im.width / im.height
    except Exception:
        return 16 / 10


def _screen(s: S, c, x, y, w, h, radius=0):
    """The device's screen content: the video if the slide has one, else the screenshot, never cropped."""
    v = s.sd.get("_video")
    s.rect(x, y, w, h, "000000", radius=radius)
    if v:
        return s.movie(x, y, w, h, v, fit="contain", radius=radius, edge=False)
    p = s.d.path(c.get("image"))
    if not p or not Path(p).exists():
        s.d.warnings.append(f"slide {s.no}: missing screen image {c.get('image')}")
        return None
    with Image.open(p) as im:
        ia = im.width / im.height
    bx, by, bw, bh = contain_box(x, y, w, h, ia)
    pic = s.slide.shapes.add_picture(p, E(s.X(bx, bw)), E(by), E(bw), E(bh))
    if radius:
        s._round(pic, radius, bw, bh)
    return pic


def _laptop(s: S, c, x, y, w, h):
    a = min(1.9, max(1.3, _media_aspect(s, c)))
    bez = 0.13
    base_h = 0.2
    # lid width L: screen (L-2b) x (L-2b)/a, lid = screen + bezel; base 1.14 L wide
    L = min(w / 1.14, (h - base_h - 2 * bez - 0.06) * a + 2 * bez)
    sw = L - 2 * bez
    shh = sw / a
    lid_h = shh + 2 * bez + 0.06
    tot_h = lid_h + base_h
    lx = x + (w - L) / 2
    ly = y + (h - tot_h) / 2
    lid = s.rect(lx, ly, L, lid_h, DEVICE_BODY, radius=0.16)
    lid.line.fill.solid(); lid.line.fill.fore_color.rgb = RGBColor.from_string(s.d.mix(s.bg, s.fg, 0.3) if not s.bg_is_light() else "3A3A3C")
    lid.line.width = Pt(0.75)
    s.shadow(lid, blur=0.6, dist=0.2, alpha=0.18)
    _screen(s, c, lx + bez, ly + bez, sw, shh)
    s.rect(lx + L / 2 - 0.025, ly + bez / 2 - 0.025, 0.05, 0.05, "3A3A3C", shape=MSO_SHAPE.OVAL)   # camera
    bw_ = L * 1.14
    bx = x + (w - bw_) / 2
    base = s.rect(bx, ly + lid_h - 0.01, bw_, base_h, "C9CBD0", shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
    base.rotation = 180  # rounded corners at the bottom
    s.rect(bx, ly + lid_h - 0.01, bw_, 0.045, "E4E5E8")
    s.rect(x + w / 2 - L * 0.08, ly + lid_h - 0.01, L * 0.16, 0.06, "A9ACB2", radius=0.03)   # thumb notch
    return lx, ly, L, tot_h


def _phone(s: S, c, x, y, w, h):
    a = _media_aspect(s, c)
    a = a if 0.4 <= a <= 0.62 else 0.46
    bez = 0.11
    ph_h = h
    ph_w = (ph_h - 2 * bez) * a + 2 * bez
    if ph_w > w:
        ph_w = w
        ph_h = (ph_w - 2 * bez) / a + 2 * bez
    px, py = x + (w - ph_w) / 2, y + (h - ph_h) / 2
    r = ph_w * 0.16
    body = s.rect(px, py, ph_w, ph_h, DEVICE_BODY, radius=r)
    body.line.fill.solid(); body.line.fill.fore_color.rgb = RGBColor.from_string("4A4A4E"); body.line.width = Pt(1.25)
    s.shadow(body, blur=0.6, dist=0.2, alpha=0.2)
    for by_, bh_ in ((0.22, 0.09), (0.33, 0.13)):  # side buttons
        s.rect(px - 0.035, py + ph_h * by_, 0.04, ph_h * bh_, "3A3A3C")
    s.rect(px + ph_w - 0.005, py + ph_h * 0.27, 0.04, ph_h * 0.15, "3A3A3C")
    _screen(s, c, px + bez, py + bez, ph_w - 2 * bez, ph_h - 2 * bez, radius=r - bez)
    s.rect(px + ph_w / 2 - ph_w * 0.15, py + bez + 0.1, ph_w * 0.3, 0.17, "000000", radius=0.085)   # island
    return px, py, ph_w, ph_h


def a_device(s: S, c):
    """A screenshot (or a video) inside a code-drawn laptop or phone, text beside it. The screen takes the
    image's own shape, so nothing is cropped. bold: the device stands on an accent field bleeding off the edge;
    editorial: on a quiet panel; soft: just the device and its shadow."""
    kind = c.get("variant", "laptop")
    left = c.get("image_side", "left") == "left"
    art = s.d.art
    if kind == "phone":
        dz_w = span(5)
        dx = col(0) if left else col(7)
    else:
        dz_w = span(7) + 0.2
        dx = col(0) if left else W - M - dz_w
    tx = (dx + dz_w + 0.7) if left else col(0)
    tw = (W - M - tx) if left else (dx - 0.7 - M)
    dy, dh = TOP - 0.05, BOTTOM - TOP + 0.05
    if art == "bold":
        fx0 = 0 if left else dx + dz_w * (0.42 if kind == "laptop" else 0.5)
        fx1 = dx + dz_w * (0.58 if kind == "laptop" else 0.5) if left else W
        s.rect(fx0, 1.75, fx1 - fx0, H - 1.75, s.d.c["accent"])
        s.fields.append((fx0, fx1 - fx0, s.d.c["accent_text"], 1.75))
    elif art == "editorial":
        s.rect(dx - 0.25, dy - 0.15, dz_w + 0.5, dh + 0.3, s.surface, radius=s.d.radius * 0.5)
    if kind == "phone":
        _phone(s, c, dx, dy + 0.05, dz_w, dh - 0.1)
    else:
        _laptop(s, c, dx, dy, dz_w, dh)
    ty = TOP + 0.45
    s.label(tx, ty, tw, c.get("kicker"))
    s.text(tx, ty + 0.35, tw, 2.0, c["title"], "h1", max_pt=c.get("title_pt", 46))
    y = ty + 0.35 + s.last_h + 0.3
    s.text(tx, y, tw, 1.4, c.get("text"), "lead", color=s.body, max_pt=19)
    if c.get("text"):
        y += s.last_h + 0.35
    bl = (c.get("bullets") or [])[:4]
    if bl:
        bh = min(0.5, (BOTTOM - 0.1 - y) / len(bl))
        bz = s.common(bl, tw - 0.6, bh - 0.06, "body", max_pt=16, min_pt=11)
        for i, b in enumerate(bl):
            yy = y + i * bh
            s.hline(tx, yy, tw, s.line, 0.75)
            s.text(tx, yy, 0.5, bh, f"{i + 1:02d}", "label", color=s.ink, anchor="m", warn=False)
            s.text(tx + 0.55, yy, tw - 0.55, bh, b, "body", anchor="m", max_pt=bz, min_pt=bz)
    s.chrome(c.get("section"))


# ---- testimonials
def a_testimonials(s: S, c):
    """2-4 customer voices. With a featured quote: it takes a full-height panel set large, the others stack beside
    it on hairlines. Without: equal columns. bold: the featured panel in the inverse ground; editorial: no panels,
    a rule between; soft: rounded cards."""
    qs = c["quotes"][:4]
    n = len(qs)
    art = s.d.art
    feat = c.get("featured", 0 if n >= 2 else None)
    mark = "”" if s.d.rtl else "“"
    if feat is None or feat >= n:
        y0 = _head(s, c, width=10, tight=True)
        cw = (W - 2 * M - (n - 1) * G) / n
        ch = BOTTOM - y0
        qz = s.common([q["quote"] for q in qs], cw - 0.7, ch - 2.1, "lead", max_pt=21, min_pt=13)
        for i, q in enumerate(qs):
            x = M + i * (cw + G)
            if art == "editorial":
                s.hline(x, y0, cw, s.fg, 1)
                fill = s.bg
            else:
                fill = s.surface
                card = s.rect(x, y0, cw, ch, fill, radius=s.d.radius)
                if art == "soft":
                    s.shadow(card, alpha=0.07)
            px = x + (0.35 if art != "editorial" else 0)
            pw_ = cw - (0.7 if art != "editorial" else 0.2)
            s.text(px, y0 + 0.2, 0.8, 0.8, mark, "display", color=s.accent, max_pt=72, min_pt=72,
                   font=s.d.fonts["heading"], warn=False)
            s.text(px, y0 + 1.0, pw_, ch - 2.1, q["quote"], "lead", max_pt=qz, min_pt=qz)
            _avatar(s, px, y0 + ch - 0.95, 0.6, q, fill=s.d.mix(fill, s.fg, 0.12))
            s.text(px + 0.75, y0 + ch - 0.98, pw_ - 0.75, 0.34, q.get("name"), "h3", max_pt=16, anchor="b")
            s.text(px + 0.75, y0 + ch - 0.62, pw_ - 0.75, 0.3, q.get("role"), "small", color=s.muted, max_pt=12)
        s.chrome(c.get("section"))
        return
    f, rest = qs[feat], [q for i, q in enumerate(qs) if i != feat]
    pw = col(7) - 0.1
    if art == "bold":
        panel, ink = s.fg, s.bg
        s.rect(0, 0, pw, H, panel)
        s.fields.append((0, pw, s.readable(s.d.mix(panel, ink, 0.7), panel, toward=ink)))
        s.photos.append((0, pw))
        fx, fw_ = M, pw - M - 0.6
    elif art == "soft":
        panel, ink = s.surface, s.fg
        s.shadow(s.rect(M, TOP - 0.1, pw - M - 0.3, BOTTOM - TOP + 0.1, panel, radius=s.d.radius))
        fx, fw_ = M + 0.45, pw - M - 1.2
    else:
        panel, ink = s.bg, s.fg
        s.vline(pw - 0.15, TOP, BOTTOM - TOP, s.line, 0.75)
        fx, fw_ = M, pw - M - 0.6
    sub = s.readable(s.d.mix(panel, ink, 0.7), panel, toward=ink)
    s.text(fx - 0.04, TOP + 0.05, 1.6, 1.4, mark, "display", color=s.d.c["accent"] if contrast_ratio(s.d.c["accent"], panel) >= 2 else ink,
           max_pt=150, min_pt=150, font=s.d.fonts["heading"], warn=False)
    s.text(fx, TOP + 1.35, fw_, BOTTOM - TOP - 2.6, f["quote"], "h1", color=ink, max_pt=c.get("max_pt", 38), min_pt=20,
           ls=1.1)
    s.hline(fx, BOTTOM - 1.05, 0.7, s.d.c["accent"], 2.5)
    _avatar(s, fx, BOTTOM - 0.82, 0.7, f, fill=s.d.mix(panel, ink, 0.16), ink=ink)
    s.text(fx + 0.88, BOTTOM - 0.84, fw_ - 0.9, 0.38, f.get("name"), "h3", color=ink, max_pt=19, anchor="b")
    s.text(fx + 0.88, BOTTOM - 0.44, fw_ - 0.9, 0.32, f.get("role"), "small", color=sub, max_pt=12.5)
    rx, rw = col(7) + 0.35, W - M - col(7) - 0.35
    s.label(rx, TOP, rw, c.get("kicker"))
    s.text(rx, TOP + 0.32, rw, 1.0, c.get("title"), "h2", max_pt=28)
    y0 = TOP + 0.32 + s.last_h + 0.4
    rh = (BOTTOM - y0) / max(1, len(rest))
    qz = s.common([q["quote"] for q in rest], rw, rh - 1.0, "body", max_pt=17, min_pt=11.5)
    for i, q in enumerate(rest):
        y = y0 + i * rh
        s.hline(rx, y, rw, s.line, 0.75)
        s.text(rx, y + 0.2, rw, rh - 1.0, f"{mark}{q['quote']}", "body", max_pt=qz, min_pt=qz, ls=1.35)
        _avatar(s, rx, y + rh - 0.68, 0.46, q, fill=s.d.mix(s.bg, s.fg, 0.12))
        s.text(rx + 0.6, y + rh - 0.72, rw - 0.6, 0.3, q.get("name"), "body", bold=True, max_pt=13.5, anchor="b")
        s.text(rx + 0.6, y + rh - 0.42, rw - 0.6, 0.28, q.get("role"), "small", color=s.muted, max_pt=11.5)
    s.chrome(c.get("section"))


# ---- logo wall
def _logo_asset(s: S, path, mono):
    """Trim a logo to its visible pixels; mono: recolour it to one ink (keeps its alpha). Returns (path, aspect)."""
    p = s.d.path(path)
    im = Image.open(p).convert("RGBA")
    a = im.getchannel("A")
    if a.getextrema()[0] == 255:   # no transparency: key out the corner colour
        g = im.convert("RGB")
        bg = g.getpixel((0, 0))
        from PIL import ImageChops
        diff = ImageChops.difference(g, Image.new("RGB", g.size, bg)).convert("L")
        a = diff.point(lambda v: min(255, v * 4))
    box = a.getbbox() or (0, 0, im.width, im.height)
    im, a = im.crop(box), a.crop(box)
    if mono:
        ink = mono if isinstance(mono, str) else s.d.mix(s.bg, s.fg, 0.82)
        solid = Image.new("RGBA", im.size, tuple(int(ink[i:i + 2], 16) for i in (0, 2, 4)) + (255,))
        solid.putalpha(a)
        im = solid
    else:
        im.putalpha(a)
    out = s.d.tmp / f"logo-{abs(hash((str(p), str(mono)))) % 10**8}.png"
    im.save(out)
    from PIL import ImageStat
    dens = ImageStat.Stat(a).mean[0] / 255   # how much ink the mark carries in its box
    return str(out), im.width / im.height, dens


def a_logos(s: S, c):
    """Client / partner wall: logos trimmed to their ink and sized by optical area (a wide wordmark and a square
    mark carry the same weight), in an even grid. mono (default) sets every logo in one ink so the wall reads as
    one surface. The claim and an optional stat sit beside it."""
    logos = [l if isinstance(l, dict) else {"path": l} for l in c["logos"][:12]]
    n = len(logos)
    art = s.d.art
    mono = c.get("mono", True)
    side = bool(c.get("stat") or c.get("text")) or n <= 8
    if side:
        s.label(col(0), TOP, span(4), c.get("kicker"))
        s.text(col(0), TOP + 0.32, span(4) - 0.1, 2.4, c["title"], "h1", max_pt=c.get("title_pt", 42))
        y = TOP + 0.32 + s.last_h + 0.3
        s.text(col(0), y, span(4) - 0.3, 1.6, c.get("text"), "body", color=s.body, max_pt=17)
        st = c.get("stat")
        if st:
            yb = metric(s, col(0), BOTTOM - 1.75, span(4), 1.15, st["value"], st.get("unit"), color=s.accent,
                        max_pt=84, min_pt=40)
            s.hline(col(0), yb + 0.1, 0.45, s.accent, 2.25)
            s.text(col(0), yb + 0.22, span(4) - 0.3, BOTTOM - yb - 0.2, st.get("label"), "small", color=s.body, max_pt=13)
        gx, gy, gw, gh = col(4) + 0.35, TOP + 0.05, W - M - col(4) - 0.35, BOTTOM - TOP - 0.05
    else:
        y0 = _head(s, c, width=10, tight=True)
        gx, gy, gw, gh = M, y0, W - 2 * M, BOTTOM - y0
    per = 2 if n <= 4 else (3 if n <= 9 else 4)
    if side and n in (7, 8):
        per = 4
    if not side:
        per = 3 if n <= 6 else 4
    rows = (n + per - 1) // per
    gap = 0.1 if art != "editorial" else 0
    cw, ch = (gw - (per - 1) * gap) / per, (gh - (rows - 1) * gap) / rows
    ch = min(ch, cw * 0.75)
    gy += (gh - rows * ch - (rows - 1) * gap) / 2
    assets = [_logo_asset(s, l["path"], mono) for l in logos]
    area = (cw * 0.6) * (ch * 0.36)   # optical size: equal ink area, capped by the cell
    for i, (lg, (p, ar, dens)) in enumerate(zip(logos, assets)):
        r, k = divmod(i, per)
        x, y = gx + k * (cw + gap), gy + r * (ch + gap)
        if art == "editorial":
            if k:
                s.vline(x, y + 0.15, ch - 0.3, s.line, 0.75)
            if r:
                s.hline(x + 0.15, y, cw - 0.3, s.line, 0.75)
        else:
            s.rect(x, y, cw, ch, s.surface, radius=s.d.radius if art == "soft" else s.d.radius * 0.4)
        # equal optical weight: wide marks are not shrunk to equal area (height ~ aspect^-0.4), light marks
        # (thin type, small ink) get a little more size
        lh_ = math.sqrt(area) * ar ** -0.4 * min(1.25, max(0.85, (0.33 / max(0.05, dens)) ** 0.3))
        lw_ = lh_ * ar
        k_ = min(1.0, cw * 0.8 / lw_, ch * 0.5 / lh_)
        lw_, lh_ = lw_ * k_, lh_ * k_
        s.slide.shapes.add_picture(p, E(s.X(x + (cw - lw_) / 2, lw_)), E(y + (ch - lh_) / 2), E(lw_), E(lh_))
    s.chrome(c.get("section"))


# ---- before / after
def a_beforeafter(s: S, c):
    """Two states, side by side. Images: they meet on one seam with a handle (the before/after slider everyone
    recognises), labels pinned in the corners. Text: the old state quiet and narrower, the new one on the inverse
    ground, an arrow on the seam."""
    b, a = c["before"], c["after"]
    labs = c.get("labels") or _he(s, ["Before", "After"], ["לפני", "אחרי"])
    y0 = _head(s, c, width=10, tight=True)
    cap = c.get("caption")
    bot = BOTTOM - (0.5 if cap else 0)
    acc, acc_ink = s.d.c["accent"], s.d.c["accent_text"]
    if b.get("image") and a.get("image"):
        x0, x1 = M, W - M
        mid = (x0 + x1) / 2
        side_cap = b.get("caption") or a.get("caption")
        ih = bot - y0 - (0.45 if side_cap else 0)
        for side, (st, lab, xx) in enumerate(((b, labs[0], x0), (a, labs[1], mid))):
            img = st["image"] if isinstance(st["image"], dict) else {"path": st["image"]}
            img.setdefault("flip", False)
            s.image(xx, y0, mid - x0, ih, img)
            lw_ = text_width(lab.upper() if s.d.caps_labels else lab, s.d.fonts["body"], True, 10) / 72 + 0.44
            lx = xx + 0.22 if side == 0 else x1 - 0.22 - lw_
            fill, ink = (s.bg, s.fg) if side == 0 else (acc, acc_ink)
            s.rect(lx, y0 + 0.22, lw_, 0.36, fill, radius=0.18)
            s.text(lx, y0 + 0.22, lw_, 0.36, lab, "label", color=ink, align="c", anchor="m", max_pt=10, min_pt=10,
                   caps=True, tracking=1 if s.d.caps_labels else 0, warn=False)
            if st.get("caption"):
                s.text(xx + (0 if side == 0 else 0.25), y0 + ih + 0.12, mid - x0 - 0.25, 0.3, st["caption"], "small",
                       color=s.body, max_pt=12.5, anchor="m", bold=True)
        s.rect(mid - 0.02, y0, 0.04, ih, "FFFFFF")
        d = 0.66
        hd = s.rect(mid - d / 2, y0 + ih / 2 - d / 2, d, d, "FFFFFF", shape=MSO_SHAPE.OVAL)
        s.shadow(hd, blur=0.25, dist=0.04, alpha=0.25) if s.bg_is_light() else None
        s.text(mid - d / 2, y0 + ih / 2 - d / 2, d, d, "‹ ›", "h3", color="1C1C1E", align="c", anchor="m", max_pt=17,
               min_pt=17, warn=False)
    else:
        bw_ = span(5)
        aw_ = W - 2 * M - bw_ - G
        h_ = bot - y0
        ax = M + bw_ + G
        s.rect(M, y0, bw_, h_, s.surface, radius=s.d.radius * (1 if s.d.art == "soft" else 0.5))
        panel = s.fg
        ink = s.bg
        s.rect(ax, y0, aw_, h_, panel, radius=s.d.radius * (1 if s.d.art == "soft" else 0.5))
        sub_a = s.readable(s.d.mix(panel, ink, 0.75), panel, toward=ink)
        for st, lab, x, w_, fg_, sub_, on in ((b, labs[0], M, bw_, s.fg, s.muted, False),
                                              (a, labs[1], ax, aw_, ink, sub_a, True)):
            px, pw_ = x + 0.42, w_ - 0.84
            lab_col = (acc if contrast_ratio(acc, panel) >= 3 else ink) if on else s.muted
            s.text(px, y0 + 0.36, pw_, 0.3, lab, "label", color=lab_col, caps=True,
                   tracking=1.2 if s.d.caps_labels else 0)
            s.text(px, y0 + 0.75, pw_, 1.4, st.get("title"), "h2" if not on else "h1", color=fg_ if on else s.body,
                   max_pt=26 if not on else 38)
            yy = y0 + 0.75 + s.last_h + 0.3
            its = st.get("items", [])[:5]
            if st.get("text"):
                s.text(px, yy, pw_, 1.2, st["text"], "body", color=sub_, max_pt=17)
                yy += s.last_h + 0.3
            if its:
                ih_ = min(0.55, (y0 + h_ - 0.3 - yy) / len(its))
                iz = s.common(its, pw_ - 0.45, ih_ - 0.06, "body", max_pt=16, min_pt=11)
                for k, it in enumerate(its):
                    s.hline(px, yy + k * ih_, pw_, s.d.mix(panel if on else s.surface, fg_, 0.22), 0.75)
                    s.text(px, yy + k * ih_, 0.35, ih_, "✓" if on else "–", "body",
                           color=(acc if contrast_ratio(acc, panel) >= 3 else ink) if on else s.muted, anchor="m",
                           max_pt=iz, warn=False)
                    s.text(px + 0.4, yy + k * ih_, pw_ - 0.45, ih_, it, "body", color=fg_ if on else s.body,
                           anchor="m", max_pt=iz, min_pt=iz)
        d = 0.72
        hx = M + bw_ + G / 2 - d / 2
        s.rect(hx, y0 + h_ / 2 - d / 2, d, d, acc, shape=MSO_SHAPE.OVAL)
        s.text(hx, y0 + h_ / 2 - d / 2, d, d, "←" if s.d.rtl else "→", "h2", color=acc_ink, align="c", anchor="m",
               max_pt=24, min_pt=24, warn=False)
    if cap:
        s.hline(M, BOTTOM - 0.22, 0.35, s.accent, 1.5)
        s.text(M + 0.5, BOTTOM - 0.38, W - 2 * M - 0.5, 0.32, cap, "small", color=s.body, max_pt=12.5, anchor="m")
    s.chrome(c.get("section"))


# ---- pyramid
def a_pyramid(s: S, c):
    """3-5 tiers, the apex narrowest. One silhouette cut into tiers, each tier's story on a leader beside it.
    The apex (or `highlight`) takes the accent. bold: solid graded tiers; editorial: outlined; soft: light tints."""
    tiers = c["tiers"][:5]
    n = len(tiers)
    art = s.d.art
    y0 = _head(s, c, width=8, tight=True) - 0.05
    px, pw = M + 0.1, span(5) + 0.4
    lx, lw = col(6) + 0.2, span(6) - 0.2
    gap = 0.08
    T = BOTTOM - y0
    th = (T - (n - 1) * gap) / n
    hl = c.get("highlight", [0])
    cx = px + pw / 2
    tz = s.common([t["title"] for t in tiers], lw, th * 0.48, "h3", max_pt=22, min_pt=13)
    bz = s.common([t.get("text") for t in tiers], lw, th * 0.52, "small", max_pt=13.5, min_pt=10)
    for i, t in enumerate(tiers):
        y = y0 + i * (th + gap)
        z0, z1 = i * (th + gap), i * (th + gap) + th
        wt, wb = pw * z0 / T, pw * z1 / T
        on = i in hl
        if art == "editorial":
            fill = s.d.c["accent"] if on else None
        elif art == "soft":
            fill = s.d.c["accent"] if on else s.d.mix(s.bg, s.fg, 0.08 + 0.05 * i)
        else:
            fill = s.d.c["accent"] if on else s.d.mix(s.bg, s.fg, max(0.22, 0.9 - 0.15 * (i - (1 if 0 in hl else 0))))
        if i == 0:
            sh = s.rect(cx - wb / 2, y, wb, th, fill or s.bg, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        else:
            sh = s.rect(cx - wb / 2, y, wb, th, fill or s.bg, shape=MSO_SHAPE.TRAPEZOID)
            sh.adjustments[0] = max(0.0, (wb - wt) / 2 / min(wb, th))
        if fill is None:
            sh.fill.background()
            sh.line.fill.solid(); sh.line.fill.fore_color.rgb = RGBColor.from_string(s.fg); sh.line.width = Pt(1)
        ink = s.d.c["accent_text"] if on else (s.fg if fill is None else _ink_on(s, fill))
        ny = y + (th * 0.42 if i == 0 else 0)
        s.text(cx - 0.6, ny, 1.2, th - (th * 0.42 if i == 0 else 0), t.get("n", f"{i + 1:02d}"), "h3", color=ink,
               align="c", anchor="m", max_pt=min(22, th * 72 * 0.4), min_pt=11, warn=False)
        ex = cx + (wt + wb) / 4 + 0.12
        s.hline(ex, y + th / 2, lx - 0.25 - ex, s.accent if on else s.line, 1.25 if on else 0.75)
        has_t = bool(t.get("text"))
        s.text(lx, y + (0.02 if has_t else 0), lw, th * (0.5 if has_t else 1), t["title"], "h3",
               color=s.ink if on else s.fg, max_pt=tz, min_pt=tz, anchor="b" if has_t else "m")
        if has_t:
            s.text(lx, y + th * 0.52 + 0.02, lw, th * 0.48, t["text"], "small", color=s.body, max_pt=bz, min_pt=bz)
    s.chrome(c.get("section"))


# ---- video
def a_video(s: S, c):
    """Video slide. full (default): the clip plays full-bleed behind a dim + scrim, a short line set large on it.
    framed: the clip in a frame inside the margins (cropped only if `fit: cover`), the caption beside it, with
    the running time as a small meta chip."""
    v = s.sd.get("_video")
    if not v:
        s.d.warnings.append(f"slide {s.no}: video slide without a playable video")
    if c.get("variant", "full") == "full":
        if c.get("image"):
            s.bg_image(c["image"], c.get("dim"))
        s.text(col(0), 1.6, span(c.get("width", 9)), 3.4, c["title"], "display", max_pt=c.get("max_pt", 88), anchor="b")
        s.label(col(0), 5.0 - s.last_h - 0.45, span(6), c.get("kicker"))
        s.text(col(0), 5.25, span(6), 1.0, c.get("text"), "lead")
        s.chrome(c.get("section"))
        return
    left = c.get("image_side", "left") == "left"
    art = s.d.art
    vw = span(8)
    vx = col(0) if left else W - M - vw
    with Image.open(v["poster"] if v else s.d.path(c["image"])) as im:
        ar = im.width / im.height
    vh = min(vw / ar, BOTTOM - TOP - 0.3)
    vw_ = vh * ar
    vx = vx if left else W - M - vw_
    vy = TOP + (BOTTOM - TOP - vh) / 2 - 0.05
    if art == "bold":
        s.rect(vx + 0.3, vy + 0.3, vw_, vh, s.d.c["accent"], radius=s.d.radius * 0.5)
    if c.get("image"):
        s.image(vx, vy, vw_, vh, c["image"], radius=s.d.radius * (1 if art == "soft" else 0.5), flip=False)
    tx = vx + vw_ + 0.7 if left else col(0)
    tw = (W - M - tx) if left else (vx - 0.7 - M)
    s.label(tx, vy + 0.05, tw, c.get("kicker"))
    s.text(tx, vy + 0.4, tw, 2.2, c["title"], "h1", max_pt=c.get("title_pt", 40))
    y = vy + 0.4 + s.last_h + 0.3
    s.text(tx, y, tw, vy + vh - y - 0.8, c.get("text"), "body", color=s.body, max_pt=17)
    if v:
        secs = round(v["dur_ms"] / 1000)
        chip = f"{secs // 60}:{secs % 60:02d}" + (f"  ·  {c['meta']}" if c.get("meta") else "")
        s.hline(tx, vy + vh - 0.45, tw, s.line, 0.75)
        play = s.rect(tx + 0.02, vy + vh - 0.27, 0.17, 0.15, s.accent, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
        play.rotation = 90   # a drawn play mark (the ▶ character turns into an emoji on some systems)
        s.text(tx + 0.32, vy + vh - 0.33, tw - 0.32, 0.3, chip, "label", color=s.muted, caps=True, anchor="m", warn=False)
    s.chrome(c.get("section"))


VARIANTS = {
    ("cover", "block"): v_cover_block, ("statement", "block"): v_statement_block, ("section", "ghost"): v_section_ghost,
    ("kpi", "hero"): v_kpi_hero, ("kpi", "bars"): v_kpi_bars, ("cards", "rows"): v_cards_rows,
    ("image-text", "framed"): v_image_framed, ("image-text", "full"): v_image_full, ("comparison", "split"): v_comparison_split, ("closing", "block"): v_closing_block,
    ("process", "stairs"): v_process_stairs, ("timeline", "band"): v_timeline_band, ("agenda", "split"): v_agenda_split,
    ("team", "circles"): v_team_circles, ("gallery", "mosaic"): v_gallery_mosaic, ("quote", "bleed"): v_quote_bleed,
}

# Per art direction: the layouts each archetype rotates through (None = the base layout).
# A type that appears several times in one deck takes the next layout each time, so nothing repeats.
ROTATION = {
    "bold": {
        "cover": [None], "statement": ["block", None], "section": ["ghost", None],
        "kpi": ["hero", None], "cards": ["rows", None], "image-text": ["framed", None],
        "comparison": ["split", None], "closing": ["block"], "process": ["stairs", None],
        "timeline": [None, "band"], "agenda": ["split", None], "team": ["circles", None],
        "gallery": ["mosaic", None], "quote": ["bleed", None],
    },
    "editorial": {
        "section": ["ghost", None], "cards": ["rows", None], "image-text": ["framed", None], "kpi": ["hero", None],
        "agenda": [None, "split"], "process": [None, "stairs"], "quote": ["bleed", None], "team": ["circles", None],
    },
}


def default_variant(art, sd, seen):
    """Choose a layout for this slide: respect content limits, then rotate per type across the deck."""
    t = sd["type"]
    opts = list(ROTATION.get(art, {}).get(t, [None]))
    if t == "cover":
        opts = ["block" if sd.get("image") else None] if art == "bold" else [None]
    if t == "statement" and sd.get("image"):
        opts = [None]
    if t == "closing" and sd.get("image"):
        opts = [None]
    if t == "kpi" and len(sd.get("metrics", [])) > 3:
        opts = [o for o in opts if o != "hero"] or [None]
    if t == "kpi" and sd.get("weights"):  # bands only when their lengths mean something
        opts = ["bars"]
    if t == "cards" and len(sd.get("cards", [])) > 5:
        opts = [o for o in opts if o != "rows"] or [None]
    words = len((sd.get("title", "") + " " + (sd.get("text") or "")).split())
    if t == "image-text" and not sd.get("bullets") and sd.get("fit") != "contain" and words < 14             and seen.get("_full", 0) < 2:
        opts = ["full"]
    if (t == "image-text" and opts == ["full"]) or (t in ("statement", "closing") and sd.get("image"))             or (t == "cover" and sd.get("variant") == "image"):
        seen["_full"] = seen.get("_full", 0) + 1
    if t == "quote" and not sd.get("image"):
        opts = [o for o in opts if o != "bleed"] or [None]
    if t == "timeline" and len(sd.get("events", [])) > 6:
        opts = [o for o in opts if o != "band"] or [None]
    if t == "agenda" and len(sd.get("items", [])) > 7:
        opts = [o for o in opts if o != "split"] or [None]
    k = seen.get(t, 0)
    seen[t] = k + 1
    return opts[k % len(opts)]


ARCHETYPES = {
    "cover": (a_cover, "title, subtitle?, kicker?, image?, variant: type|image|split|block, dim? (0.3-0.7), max_pt?, width? (cols), meta? (type variant only), focus? [fx,fy]"),
    "section": (a_section, "title, number?, kicker?, text?, cutout? (png), cutout_x/y/w? (inches; ghost default x=col 7, y=0.95, w=3.3), variant ghost|base"),
    "statement": (a_statement, "title (the statement), kicker?, text?, image? (full-bleed), dim?, width? (cols)"),
    "agenda": (a_agenda, "title, kicker?, text?, items: [str | {title, text?, meta?}] (3-8)"),
    "kpi": (a_kpi, "title, kicker?, text? (source/context note), metrics: [{value, unit?, label}] (2-4), highlight: [idx] (hero = first highlighted), variant hero|bars|base, weights? (bars only)"),
    "process": (a_process, "title, kicker?, text?, steps: [{title, text, n?}] (3-5), highlight: [idx], note?"),
    "timeline": (a_timeline, "title, kicker?, text?, events: [{when, title, text?}] (3-8), highlight: [idx]"),
    "comparison": (a_comparison, "title, kicker?, text?, left: {title, items[]}, right: {title, items[]} (right = the good side)"),
    "cards": (a_cards, "title, kicker?, text?, cards: [{title, text, n?}] (2-6), highlight: [idx]"),
    "image-text": (a_image_text, "title, image, kicker?, text?, bullets?: [str] (<=4), image_side: left|right, image_share?, variant framed|base, mask? arch|circle|round, fit? cover|contain (contain for screenshots/UI), focus? [fx,fy]"),
    "quote": (a_quote, "quote, name?, role?, image? (portrait) | cutout? (transparent png)"),
    "team": (a_team, "title, kicker?, text?, people: [{name, role, photo}] (2-4)"),
    "chart": (a_chart, "title, kicker?, text?, chart: column|bar|line|doughnut, categories[], series: [{name, values[]}], highlight?: idx, takeaway? (a figure or 2-6 words), takeaway_label?, format?, clean?"),
    "pricing": (a_pricing, "title, kicker?, text?, plans: [{name, price, unit?, period?, features[], cta?, badge?}] (1-3), highlight: [idx], points? (single plan: up to 4 lines beside the card)"),
    "gallery": (a_gallery, "title, kicker?, text?, images: [path | {path, caption}] (2-3), weights?"),
    "closing": (a_closing, "title, title2? (accent line), kicker?, contacts: [[label, value]], note?, image?"),
    "text": (a_text, "title, body (str with blank lines | [paragraphs]) 80-220 words, kicker?, lead?, quote?, quote_by?, variant essay|columns"),
    "case": (a_case, "title, challenge, approach, result (30-70 words each), metric?: {value, label}, kicker?, *_label?"),
    "bullets": (a_bullets, "title, points: [{lead, text}] (3-6, text 12-35 words), kicker?, text?"),
    "faq": (a_faq, "title, items: [{q, a}] (3-6, answers 15-45 words), kicker?, text?"),
    "table": (a_table, "title, columns: [str] (2-7), rows: [[cell]] (1-8), kicker?, text?, highlight_col?: idx, highlight_row?: idx, total? (last row = total), source?, takeaway?, takeaway_label? (<=4 columns: title column beside), wide? (force full width)"),
    "funnel": (a_funnel, "title, stages: [{label, value, text?, conversion?}] (3-6, widest first), kicker?, text?, highlight: [idx] (default last), conversion? (false hides step rates)"),
    "matrix": (a_matrix, "title, quadrants: [{title, text? | items[], tag?}] x4 (top-left, top-right, bottom-left, bottom-right), x_axis?, y_axis?, highlight?: idx, kicker?, text?, variant base|swot"),
    "org": (a_org, "title, root: {name, role, photo?}, reports: [{name, role, photo?, children: [{name, role}] (<=3)}] (1-4), highlight?: [report] | [report, child], kicker?, text?"),
    "roadmap": (a_roadmap, "title, periods: [str] (3-6), lanes: [{name, text?, items: [{title, from, to, highlight?}]}] (2-4), now? (period name or number, e.g. 1.5), now_label?, kicker?, text?"),
    "device": (a_device, "title, image (screenshot) | video, kicker?, text?, bullets?: [str] (<=4), variant laptop|phone, image_side left|right"),
    "testimonials": (a_testimonials, "title, quotes: [{quote, name, role?, photo?}] (2-4), featured?: idx | null (default 0; null = equal columns), kicker?"),
    "logos": (a_logos, "title, logos: [path | {path}] (4-12), kicker?, text?, stat?: {value, unit?, label}, mono? (default true: one ink; false keeps colours; or a hex)"),
    "beforeafter": (a_beforeafter, "title, before: {image | title, text?, items[]?, caption?}, after: {same}, labels?: [before, after], caption?, kicker?, text?"),
    "pyramid": (a_pyramid, "title, tiers: [{title, text?, n?}] (3-5, apex first), highlight: [idx] (default apex), kicker?, text?"),
    "video": (a_video, "title, video (mp4 path | {path, poster?, loop?, autoplay?, muted?, audio?}), poster?, kicker?, text?, dim?, variant full|framed, image_side?, meta?"),
}


def build(spec, out, base_dir):
    d = Deck(spec, base_dir)
    seen = {}
    # type-only decks need colour moments: alternate statements / hero numbers onto the dark ground
    imageless = not any(k in sd for sd in spec["slides"] for k in ("image", "images", "people", "cutout", "overlays", "video"))
    rhythm = 0
    for sd in spec["slides"]:
        if imageless and d.art != "soft" and "tone" not in sd and sd["type"] in ("statement", "kpi"):
            rhythm += 1
            if rhythm % 2 == 1:
                sd = dict(sd, tone="inverse")
        if sd.get("video"):  # the video takes the slide's main image slot; its poster stands in for the photo
            vv = d.video(sd)
            if vv:
                sd = dict(sd, _video=vv, image=vv["poster"], flip=False)
        v = sd.get("variant") or default_variant(d.art, sd, seen)
        fn = VARIANTS.get((sd["type"], v)) or ARCHETYPES[sd["type"]][0]
        if v and (sd["type"], v) not in VARIANTS:
            sd = dict(sd, variant=v)  # built-in variants (cover type/image/split, closing split)
        s = S(d, sd)
        fn(s, sd)
        for ov in sd.get("overlays", []):  # free-placed transparent cut-outs (LTR coords, mirrored in RTL)
            s.cutout(ov["x"], ov["y"], ov["w"], ov["path"], flip=ov.get("flip"))
        if sd.get("notes"):
            s.slide.notes_slide.notes_text_frame.text = sd["notes"]
        tr = sd.get("transition", spec.get("transition", "fade"))
        if s.slide._element.find(qn("p:clrMapOvr")) is None:  # anchor for transition/timing order
            etree.SubElement(s.slide._element, qn("p:clrMapOvr")).append(etree.Element(qn("a:masterClrMapping")))
        motion.transition(s.slide, tr)
        motion.animate(s.slide, W, H, d.rtl, sd.get("motion", spec.get("motion", "subtle")), media=s.media)
    n_sections = sum(1 for sd in spec["slides"] if sd["type"] == "section")
    if n_sections == 1:
        d.warnings.append("deck: only one section divider - use at least two (01, 02 ...) or none; "
                          "a lone divider reads as a second cover")
    d.prs.save(out)
    print(f"saved {out} ({d.n} slides)")
    for w in d.warnings:
        print("WARN", w)
    return d.warnings


def check_fonts(brand):
    """Report which file each brand font resolves to and the family name PowerPoint will see."""
    ok = True
    for role, fam in brand["fonts"].items():
        f = font_for(fam, True, 20)
        if f is None:
            ok = False
            print(f"MISSING  {role:11} '{fam}' - not installed: PowerPoint will substitute and fitting uses estimates. "
                  f"Install it through an authorized font source or pick an available font.")
            continue
        fam_name, style = f.getname()
        fname = Path(f.path_).name
        note = ""
        if fam_name.lower().replace(" ", "") != fam.lower().replace(" ", ""):
            note = f"  <- PowerPoint will list this as '{fam_name}'. Prefer a static install so the name is exactly '{fam}'."
        elif any(t in fname.lower() for t in ("variable", "wght", "[")):
            note = ("  <- variable font file: fine on this machine, but other machines may name it differently. "
                    f"For decks you send out, use an authorized static font installation.")
        print(f"ok       {role:11} '{fam}' -> {fname} ({fam_name} {style}){note}")
    return ok


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2 or sys.argv[1] not in ("fonts", "schema", "build"):
        sys.exit(__doc__)
    if sys.argv[1] == "fonts":
        if len(sys.argv) < 3:
            sys.exit("usage: python engine/deck.py fonts <brand.json>")
        bp = Path(sys.argv[2])
        check_fonts(json.loads(bp.read_text(encoding="utf8")))
        sys.exit(0)
    if sys.argv[1] == "schema":
        for k, (_, sch) in ARCHETYPES.items():
            print(f"{k:12} {sch}")
    else:
        spec_path = Path(sys.argv[2])
        build(json.loads(spec_path.read_text(encoding="utf8")), sys.argv[3], spec_path.parent)

